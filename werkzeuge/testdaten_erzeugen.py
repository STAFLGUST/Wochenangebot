"""Erzeugt die anonymisierten Testdaten unter tests/fixtures aus echten Originaldateien.

Die Originale (Excel-Angebote, WinAB-Bestandsliste, Datenstände) liegen NIE im Repository.
Dieses Skript wird nur lokal ausgeführt, wenn sich das Format der Originale ändert:

    python3 werkzeuge/testdaten_erzeugen.py --angebote <Ordner mit Wochenangebot-*.xlsx> \
        --bestand <bestand.xls> --seed <Ordner mit angebot.php> --seed-clean <Ordner> --out tests/fixtures

Anonymisierung (deterministisch, damit Tests stabil bleiben):
- Verkaufspreise: je Artikelnummer mit festem Faktor 0,88–1,12 verändert.
- Bestand: Mengen je Zeile 0,6–1,4-fach, Kolli angepasst, Einkaufspreise frei erfunden,
  Partie- und Gruppennummern ersetzt. Artikelnummern, Bezeichnungen, Status und MHD bleiben
  (sie stehen ohnehin im öffentlichen Angebot bzw. sind ohne Mengen nicht vertraulich).
Benötigt: openpyxl, xlrd==2.0.1, xlwt==1.3.0
"""
import argparse, datetime, glob, hashlib, json, os

import openpyxl, xlrd, xlwt

GUARD = "<?php http_response_code(404); exit; ?>\n"


def h(*parts):
    return int(hashlib.sha256('|'.join(str(p) for p in parts).encode()).hexdigest()[:12], 16)


def factor(code, lo=0.88, hi=1.12):
    return lo + (h('preis', code) % 10000) / 10000 * (hi - lo)


def price(v, code):
    return round(v * factor(code), 2)


def anon_offer_xlsx(src, dst):
    wb = openpyxl.load_workbook(src)
    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for i, c in enumerate(row):
                v = c.value
                if isinstance(v, bool) or not isinstance(v, (int, float)) or c.data_type == 'f':
                    continue
                if v >= 1000:      # Artikelnummer
                    continue
                code = None
                for back in range(2, 7):  # Muster: Art.-Nr. | Bezeichnung | (verbundene Zellen) | Preis
                    v2 = row[i - back].value if i - back >= 0 else None
                    if isinstance(v2, (int, float)) and not isinstance(v2, bool) and v2 >= 1000:
                        code = int(v2); break
                    if isinstance(v2, (int, float)) and not isinstance(v2, bool):
                        break
                c.value = price(v, code if code is not None else f'{ws.title}!{c.coordinate}')
    # Metadaten ohne Personennamen
    p = wb.properties
    p.creator = p.lastModifiedBy = 'Testdaten'
    p.title = p.subject = p.description = p.keywords = None
    p.created = p.modified = datetime.datetime(2026, 1, 1)
    p.lastPrinted = None
    wb.save(dst)


def anon_store(src, dst):
    raw = open(src, encoding='utf-8').read()
    d = json.loads(raw[raw.index('\n') + 1:])
    for it in d.get('items', []):
        for p in [it.get('price')] + list(it.get('altPrices') or []):
            if p and p.get('status') == 'value' and isinstance(p.get('value'), (int, float)):
                p['value'] = price(p['value'], it.get('code') or it.get('id'))
                p['raw'] = str(p['value'])
    with open(dst, 'w', encoding='utf-8') as fh:
        fh.write(GUARD + json.dumps(d, ensure_ascii=False, separators=(',', ':')))


def ddmmyyyy(v):
    if v in ('', None):
        return None
    s = str(int(v)).zfill(8)
    return f'{s[4:]}-{s[2:4]}-{s[:2]}'


def anon_stock(src, dst_xls):
    """WinAB-Partieliste: Spalten 1 Art.-Nr., 4 Bezeichnung, 6 EK-Gruppe, 8 VK-Gruppe, 9 Status,
    12 Ankunft, 14 Partie, 16 MHD, 18 Kolli, 21 netto kg, 22 EK-Preis."""
    sh = xlrd.open_workbook(src, logfile=open(os.devnull, 'w')).sheet_by_index(0)
    out = xlwt.Workbook(encoding='utf-8')
    ws = out.add_sheet('Sheet 1')
    for r in range(sh.nrows):
        for c in range(sh.ncols):
            v = sh.cell_value(r, c)
            if v in ('', None):
                continue
            if r >= 3 and isinstance(v, float):
                code = sh.cell_value(r, 1)
                if c == 21:      # netto kg
                    v = round(v * (0.6 + (h('kg', r, code) % 800) / 1000), 3)
                elif c == 18:    # Kolli
                    v = float(max(1, round(v * (0.6 + (h('kg', r, code) % 800) / 1000))))
                elif c == 22:    # EK-Preis: frei erfunden
                    v = round(3 + (h('ek', r) % 3700) / 100, 2)
                elif c == 14:    # Partie
                    v = float(500000 + r)
                elif c in (6, 8):  # Gruppen
                    v = float(100 + h('grp', c, v) % 900)
            ws.write(r, c, v)
    out.save(dst_xls)
    return aggregate_stock(dst_xls)


def aggregate_stock(path):
    """Referenzauswertung (unabhängig von der App): je Artikel Summe kg der eingelagerten Zeilen,
    frühestes MHD und kg am frühesten MHD."""
    sh = xlrd.open_workbook(path, logfile=open(os.devnull, 'w')).sheet_by_index(0)
    arts, order = {}, []
    for r in range(3, sh.nrows):
        code = sh.cell_value(r, 1)
        if not isinstance(code, float):
            continue
        code = str(int(code)).zfill(5)
        if code not in arts:
            arts[code] = {'name': sh.cell_value(r, 4), 'rows': []}; order.append(code)
        if sh.cell_value(r, 9) != 'eingelagert':
            continue
        arts[code]['rows'].append((sh.cell_value(r, 21) or 0, ddmmyyyy(sh.cell_value(r, 16))))
    res = []
    for c in order:
        rows = arts[c]['rows']
        kg = round(sum(k for k, _ in rows), 6)
        if not rows or kg <= 0:      # App zeigt nur positiven Bestand
            continue
        m = [d for _, d in rows if d]
        mhd = min(m) if m else None
        mk = round(sum(k for k, d in rows if d == mhd), 6) if mhd else None
        res.append([c, arts[c]['name'], kg, mhd, mk])
    return res


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--angebote', required=True)
    ap.add_argument('--bestand', required=True)
    ap.add_argument('--seed', required=True)
    ap.add_argument('--seed-clean', required=True)
    ap.add_argument('--out', default='tests/fixtures')
    a = ap.parse_args()
    os.makedirs(f'{a.out}/angebote', exist_ok=True)
    os.makedirs(f'{a.out}/bestand', exist_ok=True)
    for f in sorted(glob.glob(os.path.join(a.angebote, '*.xlsx'))):
        anon_offer_xlsx(f, os.path.join(a.out, 'angebote', os.path.basename(f)))
    master = anon_stock(a.bestand, f'{a.out}/bestand/bestand.xls')
    json.dump(master, open(f'{a.out}/bestand/master_stock.json', 'w'), ensure_ascii=False)
    for name, src in (('data-seed', a.seed), ('data-seed-clean', a.seed_clean)):
        d = f'{a.out}/{name}'
        os.makedirs(d, exist_ok=True)
        anon_store(os.path.join(src, 'angebot.php'), os.path.join(d, 'angebot.php'))
        with open(os.path.join(d, 'index.php'), 'w') as fh:
            fh.write(GUARD)
    print('Testdaten erzeugt in', a.out)


if __name__ == '__main__':
    main()
