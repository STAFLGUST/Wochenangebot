"""REWE (Modul rewe, Oberfläche rewe.intern unter /module/rewe/) – Vordisposition Dry Aged für REWE-Märkte.

Geprüft: Register (in Prüfung, ohne Freigabe kein Zugang, Testfreigabe nur für eingetragene Tester); Rechte serverseitig (GR-103):
ohne Anmeldung, ohne REWE-Recht, nur ansehen, bearbeiten (nur eigene Märkte, fremde wie nicht vorhanden), alle Listen, verwalten,
Wochenabschluss (Recht oder Team des Moduls); ohne CSRF, GET auf schreibende Aktion; Grundwerte (Direktbelieferung mit Bestellschluss
Freitag früh, Artikel Roastbeef/Filetkopf/Vac. Set); Lieferwege und Artikel (Pflichtfelder, doppelt, nicht entfernbar, WinAB-Nummer);
Märkte (anlegen nur verwalten, Pflichtfelder, Kd.-Nr. eindeutig, Name/Anschrift aus dem Kundenstamm der Verwaltung, Stammfelder nur
verwalten, Stand 409, Ansprechpartner, Kaufbereitschaft und Dauerauftrag mit Verlauf); Wochenabfrage (Standard = nächste Lieferwoche,
eigene Liste, gruppiert Zentrallager vor Direkt, alle Listen nur mit Recht/Team, ruhende Märkte als eigene Liste); Antworten
(Fehlerfälle, Mengen je Artikel, Stand 409, zurücknehmen, höchstens zwölf Wochen voraus, Schnappschuss von Lieferweg und
Zuständigkeit); Daueraufträge (fällig nach Rhythmus und ab-KW, Pause, übernehmen nur ohne Antwort, mit Filetkopf); Zuweisung und
Lieferweg ab KW (frühere Wochen unverändert, Verlauf mit Namen); Notizen und Wiedervorlagen; Wochenabschluss (Einzelaufträge je Markt,
Summen je Lieferweg und Mitarbeiter, offene Märkte, Artikel ohne Nummer, Nachträge nach dem Export, in WinAB erfasst) und Excel
(Blätter, Kopf, Auftragsnummern, Texte bleiben Text, Protokoll); Datenpaket (Rechte, Fehler je Markt, Vorschau schreibt nichts, nur in
leeres Modul oder mit „ersetzen“, Listen den Mitarbeitern zugeordnet, Verlauf über mehrere Jahre); Oberfläche (Meine Woche mit Antwort
per Klick und Mengen, Marktakte mit Notiz, Zuweisung, Wochenabschluss mit Excel, Einstellungen mit Datenpaket; 1280 und 390 px, hell
und dunkel, kein waagerechtes Scrollen, keine JavaScript-Fehler). Nur erfundene Testdaten.
"""
import datetime
import io
import json
import os

from openpyxl import load_workbook
from playwright.sync_api import sync_playwright

from kunden_hilfe import Client, B
from umgebung import SH

PW = 'Sicheres-Passwort-1'
HIER = os.path.dirname(os.path.abspath(__file__))
PAKET = json.load(open(os.path.join(HIER, '..', 'fixtures', 'rewe', 'datenpaket.json'), encoding='utf-8'))
res = []
fehler = []


def check(n, c, i=''):
    res.append(bool(c)); print(('PASS ' if c else 'FAIL ') + n + (('  ' + str(i)[:400]) if i else ''))


def kw_von(d):
    j, w, _ = d.isocalendar(); return f'{j}-{w:02d}'


def kw_plus(kw, n):
    return kw_von(datetime.date.fromisocalendar(int(kw[:4]), int(kw[5:]), 1) + datetime.timedelta(weeks=n))


HEUTE = datetime.date.today()
ABSCHLUSS = kw_von(HEUTE + datetime.timedelta(days=3))   # Woche des zuletzt erreichten Bestellschlusses (Freitag früh)
ABFRAGE = kw_plus(ABSCHLUSS, 1)
kw_text = f'KW {int(ABFRAGE[5:])}/{ABFRAGE[:4]}'

A = Client('198.51.100.2')
check('Administrator angemeldet', A.call('login', {'username': 'admin', 'password': PW})[0] == 200)
NUTZER = (('lena', 'Lena Leitung', 'editor', 'Leitung-Konto-11', {'grant': ['rewe.verwalten', 'rewe.exportieren']}),
          ('mark', 'Mark Markt', 'editor', 'Laden-Konto-22', {'grant': ['rewe.bearbeiten']}),
          ('evi', 'Evi Eins', 'editor', 'Eins-Konto-33', {'grant': ['rewe.bearbeiten']}),
          ('lea', 'Lea Lesen', 'viewer', 'Leserin-Konto-44', {'grant': ['rewe.sehen']}),
          ('sabi', 'Sabi Summe', 'editor', 'Summe-Konto-55', {'grant': ['rewe.bearbeiten']}),
          ('otto', 'Otto Ohne', 'editor', 'Ohne-Konto-66', {}),
          ('nina', 'Nina Nichttester', 'editor', 'Nicht-Konto-77', {'grant': ['rewe.bearbeiten']}))
for u, n, r, p, extra in NUTZER:
    st, r_ = A.call('user_save', {'id': '', 'name': n, 'username': u, 'role': r, 'active': True, 'password': p, 'mustChange': False, **extra})
    assert st == 200, r_
ids = {u['username']: u['id'] for u in A.call('users')[1]['users']}
K = {}
for i, (u, _, _, p, _) in enumerate(NUTZER):
    K[u] = Client(f'198.51.100.{10 + i}'); K[u].call('login', {'username': u, 'password': p})
LE, M, E, L, SA, O, NI, N = K['lena'], K['mark'], K['evi'], K['lea'], K['sabi'], K['otto'], K['nina'], Client('198.51.100.30')

# ── Register und Testfreigabe ──
mod = {m['id']: m for m in A.call('module_status')[1]['module']}['rewe']
ob = mod['oberflaechen'][0]
check('Register: REWE in Prüfung, nicht freigegeben, Adresse rewe.dashboard-stafl.de/module/rewe/',
      ob['entwicklung'] == 'in_pruefung' and ob['freigabe'] == 'keine' and ob['adresse'] == 'https://rewe.dashboard-stafl.de/module/rewe/', ob)
check('ohne Freigabe: auch der Administrator ohne Zugang (404)', A.call('rewe_start')[0] == 404)
st, r = A.call('modul_freigabe_save', {'bereich': 'rewe.intern', 'freigabe': 'test', 'grund': 'Test REWE'})
check('Testfreigabe gesetzt', st == 200, (st, r))
tester = [ids[u] for u in ('admin', 'lena', 'mark', 'evi', 'lea', 'sabi', 'otto')]
st, r = A.call('modul_tester_save', {'bereich': 'rewe.intern', 'tester': tester, 'grund': 'Test REWE'})
check('Tester eingetragen', st == 200, (st, r))
check('nicht eingetragen: kein Zugang (404), auch mit REWE-Recht', NI.call('rewe_start')[0] == 404)

# ── Rechte ──
check('ohne Anmeldung: kein Zugang (401/404)', N.call('rewe_start')[0] in (401, 404))
check('ohne REWE-Recht (nur Modulzugang der Rolle): 403', O.call('rewe_start')[0] == 403)
st, d = L.call('rewe_start')
check('Ansehen: Start ok, keine weiteren Rechte, keine Mitarbeiterliste', st == 200 and d['rechte'] == {'bearbeiten': False, 'alle': False, 'verwalten': False, 'exportieren': False}
      and d['personen'] == [], d.get('rechte'))
check('Kalenderwochen: Abfrage = Woche nach dem nächsten Bestellschluss, Abschluss = Woche davor', d['abfrage'] == ABFRAGE and d['abschlussKw'] == ABSCHLUSS, (d['abfrage'], d['abschlussKw']))
check('Grundwerte: nur Direktbelieferung (Bestellschluss Freitag früh der Vorwoche), Artikel Roastbeef, Filetkopf, Vac. Set (Stück)',
      [(x['id'], x['art'], x['bestellschluss']) for x in d['lieferwege']] == [('direkt', 'direkt', 'Freitag früh der Vorwoche')]
      and [x['id'] for x in d['artikel']] == ['rb', 'fk', 'set'] and all(x['einheit'] == 'Stück' and x['nr'] == '' for x in d['artikel']), d['lieferwege'])
check('Bearbeiten: nur bearbeiten', M.call('rewe_start')[1]['rechte'] == {'bearbeiten': True, 'alle': False, 'verwalten': False, 'exportieren': False})
st, d = LE.call('rewe_start')
check('Verwalten + Wochenabschluss: alle Rechte, Mitarbeiter mit „REWE ansehen“ (ohne Otto)', d['rechte'] == {'bearbeiten': True, 'alle': True, 'verwalten': True, 'exportieren': True}
      and {p['name'] for p in d['personen']} >= {'Mark Markt', 'Evi Eins', 'Lea Lesen'} and 'Otto Ohne' not in {p['name'] for p in d['personen']} and d['leer'] is True, d)

# ── Lieferwege und Artikel ──
check('Stammdaten nur mit „verwalten“ (Bearbeiten 403)', M.call('rewe_stamm_save', {'lieferwege': [{'name': 'X', 'art': 'lager'}]})[0] == 403)
LW = [{'id': 'direkt', 'name': 'Direktbelieferung', 'art': 'direkt', 'bestellschluss': 'Freitag früh der Vorwoche', 'aktiv': True},
      {'name': 'Zentrallager Nord', 'art': 'lager'}, {'name': 'Zentrallager Süd', 'art': 'lager', 'bestellschluss': 'Donnerstag 12 Uhr'}]
for name, b in (('ohne Name', {'lieferwege': [LW[0], {'name': '', 'art': 'lager'}]}), ('Name doppelt', {'lieferwege': [LW[0], LW[1], {'name': 'zentrallager nord', 'art': 'lager'}]}),
                ('Art unbekannt', {'lieferwege': [LW[0], {'name': 'A', 'art': 'rakete'}]}), ('Direkt entfernt', {'lieferwege': [LW[1]]}),
                ('Artikelnummer mit Leerzeichen', {'artikel': [{'id': 'rb', 'name': 'Roastbeef Dry Aged', 'nr': '10 01'}, {'id': 'fk', 'name': 'Filetkopf'}, {'id': 'set', 'name': 'Vac. Set'}]}),
                ('kein Artikel aktiv', {'artikel': [{'id': 'rb', 'name': 'R', 'aktiv': False}, {'id': 'fk', 'name': 'F', 'aktiv': False}, {'id': 'set', 'name': 'S', 'aktiv': False}]})):
    st, r = LE.call('rewe_stamm_save', b)
    check(f'Stammdaten: {name} → 422', st == 422 and r.get('error'), (st, r))
st, r = LE.call('rewe_stamm_save', {'lieferwege': LW, 'artikel': [{'id': 'rb', 'name': 'Roastbeef Dry Aged', 'einheit': 'Stück', 'nr': '10001'},
                                                                    {'id': 'fk', 'name': 'Filetkopf', 'einheit': 'Stück'}, {'id': 'set', 'name': 'Vac. Set', 'einheit': 'Stück'}]})
lw = {x['name']: x['id'] for x in r.get('lieferwege', [])}
check('Lieferwege und Artikel gespeichert (neue Kennungen, Bestellschluss Standard, WinAB-Nummer)', st == 200 and len(lw) == 3
      and r['lieferwege'][1]['bestellschluss'] == 'Freitag früh der Vorwoche' and r['lieferwege'][2]['bestellschluss'] == 'Donnerstag 12 Uhr' and r['artikel'][0]['nr'] == '10001', r)
NORD, SUED = lw['Zentrallager Nord'], lw['Zentrallager Süd']

# ── Kundenstamm der Verwaltung ──
rev = A.call('kunden')[1]['rev']
st, r = A.call('kunde_save', {'id': '', 'rev': rev, 'firma': 'REWE Markt GmbH', 'nr': '1799003', 'strasse': 'Stammstraße 3', 'plz': '20095', 'ort': 'Stammstadt',
                              'telefon': '040 333', 'email': 'markt3@example.org', 'active': True})
check('Kunde im Kundenstamm angelegt (WinAB-Kd.-Nr. 1799003)', st == 200, (st, r))

# ── Märkte ──
M1 = {'name': 'REWE Musterstadt', 'kdNr': '1799001', 'reweNr': '40001001', 'strasse': 'Marktplatz 1', 'plz': '20001', 'ort': 'Musterstadt', 'lieferweg': NORD, 'zustaendig': ids['mark']}
check('Markt anlegen nur mit „verwalten“ (Bearbeiten 403)', M.call('rewe_markt_save', M1)[0] == 403)
for name, b in (('ohne Ort', {'ort': ''}), ('ohne Name', {'name': ''}), ('ohne Lieferweg', {'lieferweg': ''}), ('Lieferweg unbekannt', {'lieferweg': 'lzzz'}),
                ('Kd.-Nr. mit Buchstaben', {'kdNr': '17A9'}), ('PLZ zu lang', {'plz': '123456'}), ('Kaufbereitschaft unbekannt', {'kauf': 'vielleicht'}),
                ('Zuständig unbekannt', {'zustaendig': 'uzzz'}), ('Abfragetag Sonntag', {'abfragetag': 'so'}), ('Dauerauftrag ohne Menge', {'dauer': {'rhythmus': 2, 'mengen': {}, 'ab': ABFRAGE}}),
                ('Dauerauftrag ohne ab', {'dauer': {'rhythmus': 1, 'mengen': {'rb': 1}}}), ('E-Mail des Ansprechpartners', {'ansprechpartner': [{'name': 'Frau X', 'email': 'kein-at'}]}),
                ('Ansprechpartner ohne Name', {'ansprechpartner': [{'telefon': '040 1'}]}), ('13 Ansprechpartner', {'ansprechpartner': [{'name': f'P{i}'} for i in range(13)]})):
    st, r = LE.call('rewe_markt_save', {**M1, **b})
    check(f'Markt: {name} → 422', st == 422 and r.get('error'), (st, r))
st, r = LE.call('rewe_markt_save', M1)
m1 = r.get('markt', {})
check('Markt 1 angelegt (Nord, Mark, ab Abfragewoche, nicht im Kundenstamm)', st == 200 and m1['lieferwegName'] == 'Zentrallager Nord' and m1['zustaendigName'] == 'Mark Markt'
      and m1['zustaendigHist'][0]['ab'] == ABFRAGE and m1['winab'] is False and m1['kauf'] == 'regelmaessig', r)
check('Kd.-Nr. doppelt → 422', LE.call('rewe_markt_save', {**M1, 'ort': 'Anderswo'})[0] == 422)
st, r = LE.call('rewe_markt_save', {'kdNr': '1799002', 'name': 'REWE Evistadt', 'ort': 'Evistadt', 'strasse': 'Evi-Weg 2', 'lieferweg': SUED, 'zustaendig': ids['evi'], 'kauf': 'interessiert'})
m2 = r.get('markt', {})
check('Markt 2 angelegt (Süd, Evi, Kaufbereitschaft beim Anlegen „interessiert“)', st == 200 and m2['zustaendigName'] == 'Evi Eins' and m2['kauf'] == 'interessiert', r)
st, r = LE.call('rewe_markt_save', {'kdNr': '1799003', 'lieferweg': NORD, 'zustaendig': ids['mark']})
m3 = r.get('markt', {})
check('Markt 3 nur mit Kd.-Nr.: Name, Anschrift und Telefon aus dem Kundenstamm', st == 200 and m3['winab'] is True and m3['name'] == 'REWE Markt GmbH'
      and m3['ort'] == 'Stammstadt' and m3['strasse'] == 'Stammstraße 3' and m3['telefon'] == '040 333', r)
st, r = LE.call('rewe_markt_save', {'kdNr': '1799004', 'name': 'REWE Dauerdorf', 'ort': 'Dauerdorf', 'lieferweg': 'direkt', 'zustaendig': ids['mark'],
                                    'dauer': {'rhythmus': 2, 'mengen': {'rb': 1, 'fk': 1}, 'ab': ABFRAGE}, 'abfragetag': 'do',
                                    'ansprechpartner': [{'name': 'Frau Dauer', 'rolle': 'Metzgerei', 'telefon': '040 4', 'email': 'dauer@example.org'}]})
m4 = r.get('markt', {})
check('Markt 4 (Direkt) mit Dauerauftrag 14-tägig 1 Roastbeef + 1 Filetkopf, Abfragetag, Ansprechpartner', st == 200 and m4['dauer'] == {'rhythmus': 2, 'mengen': {'rb': 1, 'fk': 1}, 'ab': ABFRAGE, 'pauseBis': ''}
      and m4['ansprechpartner'][0]['email'] == 'dauer@example.org' and m4['abfragetag'] == 'do', r)

# ── Wochenabfrage ──
st, r = M.call('rewe_woche')
w = r.get('woche', {})
check('Woche: Standard = Abfragewoche', st == 200 and w['kw'] == ABFRAGE, (st, w.get('kw')))
gr = [(g['name'], [x['id'] for x in g['maerkte']]) for g in w.get('gruppen', [])]
check('Woche: nur eigene Märkte, Zentrallager vor Direktbelieferung', gr == [('Zentrallager Nord', sorted([m1['id'], m3['id']], key=lambda i: {m1['id']: 'Musterstadt', m3['id']: 'Stammstadt'}[i])),
                                                                              ('Direktbelieferung', [m4['id']])], gr)
x4 = w['gruppen'][1]['maerkte'][0]
check('Woche: Dauerauftrag fällig, Text, Verlauf acht Wochen, darf antworten', x4['dauerFaellig'] is True and 'Dauerauftrag 14-tägig' in x4['dauerText']
      and len(x4['verlauf']) == 8 and x4['verlauf'][-1]['kw'] == kw_plus(ABFRAGE, -1) and x4['darfAntworten'] is True, x4)
check('Woche: Evi sieht nur Markt 2', [x['id'] for g in E.call('rewe_woche')[1]['woche']['gruppen'] for x in g['maerkte']] == [m2['id']])
check('Woche: Lesen ohne eigene Märkte → leer', L.call('rewe_woche')[1]['woche']['gruppen'] == [])
check('andere Liste ohne „alle“ → 403', M.call('rewe_woche&wer=*')[0] == 403 and M.call('rewe_woche&wer=' + ids['evi'])[0] == 403)
st, r = LE.call('rewe_woche&wer=*')
check('Leitung: alle Listen (4 Märkte), Liste von Evi (1 Markt)', st == 200 and sum(len(g['maerkte']) for g in r['woche']['gruppen']) == 4
      and [x['id'] for g in LE.call('rewe_woche&wer=' + ids['evi'])[1]['woche']['gruppen'] for x in g['maerkte']] == [m2['id']], r)
check('Woche: ungültige KW → 422', M.call('rewe_woche&kw=2026-60')[0] == 422 and M.call('rewe_woche&kw=morgen')[0] == 422)

# ── Antworten ──
AW = {'markt': m1['id'], 'kw': ABFRAGE, 'status': 'bestellt', 'mengen': {'rb': 2}}
check('Antwort: nur ansehen → 403', L.call('rewe_antwort', AW)[0] == 403)
csrf = M.csrf; M.csrf = ''
check('Antwort ohne CSRF: 403', M.call('rewe_antwort', AW)[0] == 403)
M.csrf = csrf
check('Antwort per GET: 405', M.call('rewe_antwort')[0] == 405)
check('fremder Markt (Evi) für Mark: 404', M.call('rewe_antwort', {**AW, 'markt': m2['id']})[0] == 404)
check('unbekannter Markt: 404', M.call('rewe_antwort', {**AW, 'markt': 'mzzz'})[0] == 404)
for name, b in (('bestellt ohne Menge', {'mengen': {}}), ('Menge 0', {'mengen': {'rb': 0}}), ('negative Menge', {'mengen': {'rb': -1}}), ('Kommazahl', {'mengen': {'rb': 1.5}}),
                ('Menge 1000', {'mengen': {'rb': 1000}}), ('unbekannter Artikel', {'mengen': {'lachs': 1}}), ('Antwort unbekannt', {'status': 'vielleicht'}),
                ('KW ungültig', {'kw': '2026-54'}), ('13 Wochen voraus', {'kw': kw_plus(kw_von(HEUTE), 13)}), ('Notiz zu lang', {'notiz': 'x' * 301})):
    st, r = M.call('rewe_antwort', {**AW, **b})
    check(f'Antwort: {name} → 422', st == 422 and r.get('error'), (st, r))
st, r = M.call('rewe_antwort', {**AW, 'r': 0, 'notiz': '=SUMME(A1)'})
a1 = r.get('antwort', {})
check('Antwort bestellt (2 Roastbeef, Stand 1, Schnappschuss Lieferweg und Zuständigkeit, Name)', st == 200 and a1['status'] == 'bestellt' and a1['mengen'] == {'rb': 2}
      and a1['r'] == 1 and a1['lieferweg'] == NORD and a1['zustaendig'] == 'Mark Markt' and a1['von'] == 'Mark Markt' and a1['quelle'] == 'hand', r)
check('Antwort: veralteter Stand → 409 mit aktueller Antwort', (lambda x: x[0] == 409 and x[1]['antwort']['r'] == 1)(M.call('rewe_antwort', {**AW, 'r': 0})))
st, r = M.call('rewe_antwort', {**AW, 'status': 'kein', 'mengen': {'rb': 5}, 'r': 1})
check('Kein Bedarf: Mengen entfallen', st == 200 and r['antwort']['status'] == 'kein' and r['antwort']['mengen'] == [] or r['antwort']['mengen'] == {}, r)
st, r = M.call('rewe_antwort', {**AW, 'status': '', 'r': 2})
check('Antwort zurückgenommen (offen)', st == 200 and r['antwort'] is None, r)
st, r = M.call('rewe_antwort', {**AW, 'mengen': {'rb': 2, 'set': 1}, 'ohneBerechnung': True, 'notiz': '=SUMME(A1)'})
check('wieder bestellt: 2 Roastbeef + 1 Vac. Set, ohne Berechnung', st == 200 and r['antwort']['mengen'] == {'rb': 2, 'set': 1} and r['antwort']['ohneBerechnung'] is True, r)
check('Mark beantwortet Markt 3 (kein Bedarf) und vergangene Woche bei Markt 1', M.call('rewe_antwort', {'markt': m3['id'], 'kw': ABFRAGE, 'status': 'kein'})[0] == 200
      and M.call('rewe_antwort', {'markt': m1['id'], 'kw': kw_plus(ABFRAGE, -1), 'status': 'bestellt', 'mengen': {'rb': 1}})[0] == 200)
x1 = next(x for g in M.call('rewe_woche')[1]['woche']['gruppen'] for x in g['maerkte'] if x['id'] == m1['id'])
check('Verlauf zeigt die Vorwoche (1 bestellt)', x1['verlauf'][-1] == {'kw': kw_plus(ABFRAGE, -1), 'status': 'bestellt', 'menge': 1}, x1['verlauf'][-1])
st, r = LE.call('rewe_antwort', {'markt': m2['id'], 'kw': ABFRAGE, 'status': 'ne'})
check('Leitung (alle + bearbeiten) antwortet für Evis Markt (Vertretung), Schnappschuss Zuständig Evi', st == 200 and r['antwort']['zustaendig'] == 'Evi Eins' and r['antwort']['von'] == 'Lena Leitung', r)

# ── Daueraufträge ──
check('Dauerauftrag übernehmen: nur ansehen 403', L.call('rewe_dauer_uebernehmen', {'kw': ABFRAGE})[0] == 403)
st, r = M.call('rewe_dauer_uebernehmen', {'kw': ABFRAGE})
check('fällige Daueraufträge übernommen (1)', st == 200 and r['uebernommen'] == 1, r)
x4 = next(x for g in M.call('rewe_woche')[1]['woche']['gruppen'] for x in g['maerkte'] if x['id'] == m4['id'])
check('Dauerauftrag als Bestellung: 1 Roastbeef + 1 Filetkopf, Herkunft Dauerauftrag', x4['antwort']['mengen'] == {'rb': 1, 'fk': 1} and x4['antwort']['quelle'] == 'dauer', x4['antwort'])
check('zweites Übernehmen ändert nichts (nur ohne Antwort)', M.call('rewe_dauer_uebernehmen', {'kw': ABFRAGE})[1]['uebernommen'] == 0)
x4n = next(x for g in M.call(f'rewe_woche&kw={kw_plus(ABFRAGE, 1)}')[1]['woche']['gruppen'] for x in g['maerkte'] if x['id'] == m4['id'])
check('Folgewoche: nicht fällig, nächste Lieferung in zwei Wochen', x4n['dauerFaellig'] is False and x4n['dauerNaechste'] == kw_plus(ABFRAGE, 2), x4n)
st, r = M.call('rewe_markt_save', {'id': m4['id'], 'dauer': {'rhythmus': 2, 'mengen': {'rb': 1, 'fk': 1}, 'ab': ABFRAGE, 'pauseBis': kw_plus(ABFRAGE, 2)}})
x4p = next(x for g in M.call(f'rewe_woche&kw={kw_plus(ABFRAGE, 2)}')[1]['woche']['gruppen'] for x in g['maerkte'] if x['id'] == m4['id'])
check('Pause bis zur fälligen Woche: nächste Lieferung danach im Rhythmus', st == 200 and x4p['dauerFaellig'] is False and x4p['dauerNaechste'] == kw_plus(ABFRAGE, 4), x4p.get('dauerNaechste'))

# ── Markt pflegen ──
st, r = M.call('rewe_markt_save', {'id': m1['id'], 'rev': m1['rev'], 'kauf': 'gelegentlich', 'hinweis': 'nur gereifte Ware'})
check('Mark ändert Kaufbereitschaft und Hinweis seines Marktes', st == 200 and r['markt']['kauf'] == 'gelegentlich' and r['markt']['rev'] == m1['rev'] + 1, r)
check('veralteter Stand → 409', M.call('rewe_markt_save', {'id': m1['id'], 'rev': m1['rev'], 'kauf': 'regelmaessig'})[0] == 409)
check('Stammfelder nur mit „verwalten“ (403)', M.call('rewe_markt_save', {'id': m1['id'], 'ort': 'Woanders'})[0] == 403)
check('fremden Markt pflegen → 404', M.call('rewe_markt_save', {'id': m2['id'], 'kauf': 'ruhend'})[0] == 404)
check('Leitung ändert Stammfelder', LE.call('rewe_markt_save', {'id': m1['id'], 'strasse': 'Marktplatz 1a'})[1]['markt']['strasse'] == 'Marktplatz 1a')

# ── Notizen und Wiedervorlagen ──
NZ = {'markt': m1['id'], 'art': 'telefonat', 'text': 'Möchte ab nächster Woche wieder wöchentlich', 'wiedervorlage': HEUTE.isoformat()}
check('Notiz: nur ansehen 403', L.call('rewe_notiz_save', NZ)[0] == 403)
for name, b in (('ohne Text', {'text': ' '}), ('Art Änderung', {'art': 'system'}), ('Datum ungültig', {'wiedervorlage': '2026-02-30'}), ('Datum in drei Jahren', {'wiedervorlage': (HEUTE + datetime.timedelta(days=1100)).isoformat()})):
    st, r = M.call('rewe_notiz_save', {**NZ, **b})
    check(f'Notiz: {name} → 422', st == 422, (st, r))
check('Notiz an fremdem Markt → 404', M.call('rewe_notiz_save', {**NZ, 'markt': m2['id']})[0] == 404)
st, r = M.call('rewe_notiz_save', NZ)
n1 = r.get('notiz', {})
check('Notiz mit Wiedervorlage heute', st == 200 and n1['wv'] == HEUTE.isoformat() and n1['von'] == 'Mark Markt', r)
x1 = next(x for g in M.call('rewe_woche')[1]['woche']['gruppen'] for x in g['maerkte'] if x['id'] == m1['id'])
check('Woche zeigt die fällige Wiedervorlage', [v['id'] for v in x1['wiedervorlagen']] == [n1['id']], x1['wiedervorlagen'])
st, r = M.call('rewe_notiz_save', {'id': n1['id'], 'erledigt': True})
check('Wiedervorlage erledigt (mit Namen)', st == 200 and r['notiz']['erledigt']['von'] == 'Mark Markt', r)
st, r = M.call('rewe_markt&id=' + m1['id'])
arten = [n['art'] for n in r.get('notizen', [])]
check('Marktakte: Wochen, Notizen und Änderungen (Kaufbereitschaft, Anlage) im Verlauf, Kennzahlen', st == 200 and r['wochen'][ABFRAGE]['status'] == 'bestellt'
      and 'telefonat' in arten and arten.count('system') >= 2 and any('Kaufbereitschaft: Kauft regelmäßig → Gelegentlich' in n['text'] for n in r['notizen']), r.get('notizen'))
check('Marktakte eines fremden Marktes → 404, mit „alle“ sichtbar', M.call('rewe_markt&id=' + m2['id'])[0] == 404 and LE.call('rewe_markt&id=' + m2['id'])[0] == 200)

# ── Zuweisung und Lieferweg ab KW ──
ZW = {'maerkte': [m1['id']], 'ab': kw_plus(ABFRAGE, 1), 'nutzer': ids['evi'], 'lieferweg': SUED}
check('Zuweisen nur mit „verwalten“ (403)', M.call('rewe_zuweisen', ZW)[0] == 403)
for name, b in (('ohne Märkte', {'maerkte': []}), ('KW ungültig', {'ab': 'bald'}), ('Mitarbeiter unbekannt', {'nutzer': 'uzzz'}), ('Lieferweg unbekannt', {'lieferweg': 'lzzz', 'nutzer': None}),
                ('nichts gewählt', {'nutzer': None, 'lieferweg': None})):
    body = {**ZW, **b}
    body = {k: v for k, v in body.items() if v is not None}
    st, r = LE.call('rewe_zuweisen', body)
    check(f'Zuweisen: {name} → 422', st == 422, (st, r))
st, r = LE.call('rewe_zuweisen', ZW)
check('Markt 1 ab nächster Woche zu Evi und Lager Süd', st == 200 and r['geaendert'] == 1, r)
check('Abfragewoche bleibt bei Mark (Lager Nord), Folgewoche bei Evi (Lager Süd)', m1['id'] in [x['id'] for g in M.call('rewe_woche')[1]['woche']['gruppen'] for x in g['maerkte']]
      and m1['id'] not in [x['id'] for g in M.call(f'rewe_woche&kw={kw_plus(ABFRAGE, 1)}')[1]['woche']['gruppen'] for x in g['maerkte']]
      and [(g['name'], [x['id'] for x in g['maerkte']]) for g in E.call(f'rewe_woche&kw={kw_plus(ABFRAGE, 1)}')[1]['woche']['gruppen']] == [('Zentrallager Süd', sorted([m1['id'], m2['id']], key=lambda i: {m1['id']: 'Musterstadt', m2['id']: 'Evistadt'}[i]))])
akte = LE.call('rewe_markt&id=' + m1['id'])[1]
check('Akte: Historie Zuständig und Lieferweg, Verlauf mit Namen, alte Antwort mit altem Schnappschuss', [h['name'] for h in akte['markt']['zustaendigHist']] == ['Mark Markt', 'Evi Eins']
      and [h['name'] for h in akte['markt']['lieferwegHist']] == ['Zentrallager Nord', 'Zentrallager Süd']
      and any('Zuständig Mark Markt → Evi Eins' in n['text'] and 'Lieferweg Zentrallager Nord → Zentrallager Süd' in n['text'] for n in akte['notizen']), akte['markt'])
check('dieselbe Zuweisung noch einmal ändert nichts', LE.call('rewe_zuweisen', ZW)[1]['geaendert'] == 0)

# ── Ruhende Märkte ──
check('Evi setzt Markt 2 auf ruhend', E.call('rewe_markt_save', {'id': m2['id'], 'kauf': 'ruhend'})[0] == 200)
w2 = E.call(f'rewe_woche&kw={kw_plus(ABFRAGE, 2)}')[1]['woche']
check('ruhender Markt nicht in der Abfrage, aber in der eigenen Liste „ruhend“', m2['id'] not in [x['id'] for g in w2['gruppen'] for x in g['maerkte']]
      and [x['id'] for x in w2['ruhend']] == [m2['id']] and w2['ruhend'][0]['darfPflegen'] is True, w2.get('ruhend'))
w2b = E.call('rewe_woche')[1]['woche']
check('mit Antwort in der Woche bleibt er dort in der Abfrage', m2['id'] in [x['id'] for g in w2b['gruppen'] for x in g['maerkte']] and not w2b['ruhend'])

# ── Team und Wochenabschluss ──
check('Wochenabschluss ohne Recht/Team → 403', SA.call('rewe_abschluss')[0] == 403 and M.call('rewe_abschluss')[0] == 403)
check('Team nur mit „verwalten“ (403)', M.call('rewe_team_save', {'team': {}})[0] == 403)
check('Team: unbekannter Mitarbeiter / ohne REWE-Recht → 422', LE.call('rewe_team_save', {'team': {'uzzz': {'alle': True}}})[0] == 422
      and LE.call('rewe_team_save', {'team': {ids['otto']: {'alle': True}}})[0] == 422)
st, r = LE.call('rewe_team_save', {'team': {ids['sabi']: {'export': True}, ids['lea']: {'alle': True}}})
check('Team gespeichert: Sabi Wochenabschluss, Lea alle Listen', st == 200, r)
check('Team wirkt: Sabi sieht den Abschluss, Lea alle Listen (ohne bearbeiten)', SA.call('rewe_abschluss')[0] == 200 and SA.call('rewe_start')[1]['rechte']['exportieren'] is True
      and L.call('rewe_woche&wer=*')[0] == 200 and L.call('rewe_start')[1]['rechte'] == {'bearbeiten': False, 'alle': True, 'verwalten': False, 'exportieren': False}
      and L.call('rewe_antwort', AW)[0] == 403)
pers = {p['name']: p for p in LE.call('rewe_start')[1]['personen']}
check('Mitarbeiterliste zeigt Team und Rechte getrennt', pers['Sabi Summe']['teamExport'] is True and pers['Sabi Summe']['rechtExport'] is False and pers['Lena Leitung']['rechtExport'] is True, pers.get('Sabi Summe'))

st, r = SA.call(f'rewe_abschluss&kw={ABFRAGE}')
ab = r.get('abschluss', {})
pos = [(z['auftrag'], z['lieferweg'], z['kdNr'], z['artikel'], z['menge']) for z in ab.get('zeilen', [])]
check('Abschluss: Einzelaufträge je Markt, Zentrallager vor Direkt, Artikel in Reihenfolge', pos == [(1, 'Zentrallager Nord', '1799001', 'Roastbeef Dry Aged', 2), (1, 'Zentrallager Nord', '1799001', 'Vac. Set', 1),
                                                                                              (2, 'Direktbelieferung', '1799004', 'Roastbeef Dry Aged', 1), (2, 'Direktbelieferung', '1799004', 'Filetkopf', 1)], pos)
check('Abschluss: Summen je Lieferweg und Mitarbeiter, Artikel ohne WinAB-Nummer, ohne Kundenstamm', ab['maerkte'] == 2 and [x['name'] for x in ab['lieferwege']] == ['Zentrallager Nord', 'Direktbelieferung']
      and [(x['name'], x['maerkte']) for x in ab['personen']] == [('Mark Markt', 2)] and ab['ohneNr'] == ['Filetkopf', 'Vac. Set'] and ab['ohneStamm'] == 2
      and ab['zeilen'][0]['ohneBerechnung'] is True and ab['exportiert'] is None, ab)
check('Abschluss: „in WinAB erfasst“ erst nach dem Excel (422)', SA.call('rewe_erfasst', {'kw': ABFRAGE, 'erfasst': True})[0] == 422)
check('Excel: Woche ohne Bestellung → 422', SA.call('rewe_export', {'kw': kw_plus(ABFRAGE, 8)})[0] == 422)
st, r, roh = SA.call('rewe_export', {'kw': ABFRAGE}, raw=True)
check('Excel erstellt (200, keine JSON-Antwort)', st == 200 and not r, st)
import urllib.request  # noqa: E402  (für die Binärantwort)
req = urllib.request.Request(B + '/api/?action=rewe_export', data=json.dumps({'kw': ABFRAGE}).encode(),
                             headers={'X-Forwarded-For': SA.ip, 'Content-Type': 'application/json', 'X-CSRF-Token': SA.csrf})
antwort = SA.http.open(req)
daten = antwort.read()
check('Excel: Dateiname und Typ', 'REWE_WinAB_' + ABFRAGE.replace('-', '_KW') + '.xlsx' in antwort.headers.get('Content-Disposition', '')
      and antwort.headers.get('Content-Type', '').startswith('application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'), antwort.headers.get('Content-Disposition'))
wb = load_workbook(io.BytesIO(daten))
ws = wb['WinAB-Erfassung']
kopf = [c.value for c in ws[4]]
zeilen = [[c.value for c in row] for row in ws.iter_rows(min_row=5, max_row=8)]
check('Excel: Blätter und Kopfzeile', wb.sheetnames == ['WinAB-Erfassung', 'Je Lieferweg', 'Je Mitarbeiter'] and kopf[:6] == ['Auftrag', 'Lieferweg', 'Liefer-KW', 'Kd.-Nr. (WinAB)', 'REWE-Markt-Nr.', 'Markt'], kopf)
check('Excel: vier Positionen in zwei Aufträgen, Kd.-Nr. als Zahl, WinAB-Nummer, Menge, ohne Berechnung', [z[0] for z in zeilen] == [1, 1, 2, 2] and zeilen[0][3] == 1799001
      and zeilen[0][9] == '10001' and zeilen[0][11] == 2 and zeilen[0][13] == 'ja' and zeilen[3][10] == 'Filetkopf', zeilen)
notiz = ws.cell(row=5, column=16)
check('Excel: Notiz „=SUMME(A1)“ bleibt Text (keine Formel)', notiz.value == '=SUMME(A1)' and notiz.data_type == 's', (notiz.value, notiz.data_type))
summen = {row[0].value: row[1].value for row in wb['Je Lieferweg'].iter_rows(min_row=4) if row[0].value}
check('Excel: Je Lieferweg mit Summe', summen.get('Zentrallager Nord') == 1 and summen.get('Direktbelieferung') == 1 and summen.get('Summe') == 2, summen)
ab = SA.call(f'rewe_abschluss&kw={ABFRAGE}')[1]['abschluss']
check('nach dem Export: exportiert mit Name, Zeilen und Stück', ab['exportiert']['von'] == 'Sabi Summe' and ab['exportiert']['zeilen'] == 4 and ab['exportiert']['stueck'] == 5, ab['exportiert'])
st, r = M.call('rewe_antwort', {'markt': m3['id'], 'kw': ABFRAGE, 'status': 'bestellt', 'mengen': {'rb': 1}})
check('Änderung nach dem Export ist ein Nachtrag', st == 200 and r['antwort']['nachtrag'] is True, r)
check('Abschluss zählt den Nachtrag', SA.call(f'rewe_abschluss&kw={ABFRAGE}')[1]['abschluss']['nachtraege'] == 1)
st, r = SA.call('rewe_erfasst', {'kw': ABFRAGE, 'erfasst': True})
check('in WinAB erfasst (mit Name), zurücknehmbar', st == 200 and r['erfasst']['von'] == 'Sabi Summe' and SA.call('rewe_erfasst', {'kw': ABFRAGE, 'erfasst': False})[1]['erfasst'] is None, r)
prot = json.dumps(A.call('log')[1], ensure_ascii=False)
check('Protokoll: Antworten, Dauerauftrag, Zuweisung, Excel, erfasst', all(t in prot for t in ('REWE: REWE Musterstadt, ' + kw_text, 'Daueraufträge für', 'Märkte ab', 'neu zugewiesen',
      'REWE: Excel für WinAB erstellt – ' + kw_text + ', 2 Einzelaufträge, 4 Positionen', 'als in WinAB erfasst markiert')), prot[:400])
check('Protokoll ohne Notizen', 'SUMME' not in prot and 'Möchte ab nächster Woche' not in prot)

# ── Datenpaket ──
check('Datenpaket nur mit „verwalten“ (403)', M.call('rewe_import', {'paket': PAKET, 'vorschau': True})[0] == 403)
st, r = LE.call('rewe_import', {'paket': {'format': 'falsch'}, 'vorschau': True})
check('falsches Format: Fehler', st == 200 and r['fehler'] and not r['zahlen'], r)
kaputt = json.loads(json.dumps(PAKET))
kaputt['maerkte'][1]['lieferweg'][0]['id'] = 'gibtsnicht'
kaputt['maerkte'][2]['kdNr'] = kaputt['maerkte'][0]['kdNr']
kaputt['maerkte'][3]['wochen'] = {'2026-05': {'s': 'bestellt', 'm': {}}}
st, r = LE.call('rewe_import', {'paket': kaputt, 'vorschau': True})
check('Vorschau mit Fehlern je Markt (Lieferweg, Kd.-Nr. doppelt, bestellt ohne Menge)', st == 200 and len(r['fehler']) == 3 and 'Musterdorf' in r['fehler'][0]
      and 'doppelt' in r['fehler'][1] and 'ohne Menge' in r['fehler'][2], r['fehler'])
check('mit Fehlern nichts übernommen (422)', LE.call('rewe_import', {'paket': kaputt, 'ersetzen': True})[0] == 422)
zu = {'Liste A': ids['mark'], 'Liste B': ids['evi']}
st, r = LE.call('rewe_import', {'paket': PAKET, 'zuordnung': zu, 'vorschau': True})
check('Vorschau: Zahlen, Listen, nicht leer', st == 200 and r['zahlen']['maerkte'] == 4 and r['zahlen']['antworten'] == 11 and r['zahlen']['bestellungen'] == 6
      and r['zahlen']['stueck'] == 8 and r['zahlen']['dauerauftraege'] == 1 and r['listen'] == {'Liste A': 2, 'Liste B': 2, 'Liste C': 1} and r['leer'] is False and not r['fehler'], r)
check('Vorschau schreibt nichts', len(LE.call('rewe_maerkte')[1]['maerkte']) == 4)
check('Liste einem Mitarbeiter ohne „REWE ansehen“ zuordnen → Fehler', LE.call('rewe_import', {'paket': PAKET, 'zuordnung': {'Liste A': ids['otto']}, 'vorschau': True})[1]['fehler'])
check('nicht leer ohne „ersetzen“ → 409', LE.call('rewe_import', {'paket': PAKET, 'zuordnung': zu})[0] == 409)
st, r = LE.call('rewe_import', {'paket': PAKET, 'zuordnung': zu, 'ersetzen': True})
check('Datenpaket übernommen (ersetzt)', st == 200 and r['zahlen']['maerkte'] == 4, r)
d = LE.call('rewe_start')[1]
check('Lieferwege aus dem Paket (mit Direkt), Team bleibt', [x['id'] for x in d['lieferwege']] == ['nordlager', 'suedlager', 'direkt']
      and next(p for p in d['personen'] if p['name'] == 'Sabi Summe')['teamExport'] is True, d['lieferwege'])
maerkte = {m['ort']: m for m in LE.call('rewe_maerkte')[1]['maerkte']}
check('Märkte: Zuständigkeit aus den Listen (Liste C → niemand), Kaufbereitschaft, Name fehlt → „REWE <Ort>“', maerkte['Beispielstadt']['zustaendigName'] == 'Mark Markt'
      and maerkte['Neustadt']['zustaendigName'] == '' and maerkte['Stillhausen']['kauf'] == 'ruhend' and maerkte['Musterdorf']['name'] == 'REWE Musterdorf',
      {o: (x['zustaendigName'], x['kauf'], x['name']) for o, x in maerkte.items()})
wa = M.call(f'rewe_woche&kw=2026-43')[1]['woche']
check('Lagerwechsel KW 42: Musterdorf ab KW 42 bei Mark über Lager Nord', [(g['id'], [x['ort'] for x in g['maerkte']]) for g in wa['gruppen']] == [('nordlager', ['Beispielstadt', 'Musterdorf'])], wa['gruppen'])
ev = E.call('rewe_woche&kw=2026-41')
check('vor KW 42 bei Evi über Lager Süd', [(g['id'], [x['ort'] for x in g['maerkte']]) for g in ev[1]['woche']['gruppen']] == [('suedlager', ['Musterdorf'])], ev)
akte = M.call('rewe_markt&id=' + maerkte['Beispielstadt']['id'])[1]
check('Akte: Verlauf über zwei Jahre, Kennzahlen, Notiz aus Excel, Herkunft Excel', akte['wochen']['2025-50']['status'] == 'bestellt' and akte['wochen']['2026-06']['status'] == 'pause'
      and akte['zahlen']['2026']['lieferungen'] == 2 and akte['zahlen']['2026']['mengen'] == {'rb': 3, 'fk': 1} and akte['notizen'][0]['q'] == 'datenpaket'
      and akte['wochen']['2026-04']['quelle'] == 'excel', akte['zahlen'])
ev = E.call('rewe_woche')[1]['woche']
check('Ruhend: Stillhausen in Evis Liste „ruhend“ mit letzter Bestellung', [(x['ort'], x['letzte']) for x in ev['ruhend']] == [('Stillhausen', '2024-10')], (ev['ruhend'], [[x['ort'] for x in g['maerkte']] for g in ev['gruppen']]))
a43, a44 = (SA.call(f'rewe_abschluss&kw={k}')[1]['abschluss'] for k in ('2026-43', '2026-44'))
check('Abschluss: ohne Antwort nur abzufragende Märkte (ruhend und nicht fälliger Dauerauftrag zählen nicht), fälliger Dauerauftrag als Hinweis',
      [x['ort'] for x in a43['offen']] == ['Musterdorf', 'Neustadt'] and a43['dauerOffen'] == [] and [x['ort'] for x in a44['dauerOffen']] == ['Beispielstadt']
      and 'Beispielstadt' not in [x['ort'] for x in a44['offen']], ([x['ort'] for x in a43['offen']], a43['dauerOffen'], a44['dauerOffen']))
check('ohne Zuständigkeit: niemandem außer „alle“ sichtbar', LE.call('rewe_markt&id=' + maerkte['Neustadt']['id'])[0] == 200 and M.call('rewe_markt&id=' + maerkte['Neustadt']['id'])[0] == 404)

# ── Oberfläche ──
with sync_playwright() as p:
    br = p.chromium.launch()

    def text(pg, sel):
        return pg.inner_text(sel).replace('\xa0', ' ')

    def anmelden(user, pw, ip, breite=1280, thema='light', hoehe=900, ziel='#rw-app:not([hidden])', pfad='/module/rewe/'):
        ctx = br.new_context(viewport={'width': breite, 'height': hoehe}, extra_http_headers={'X-Forwarded-For': ip}, locale='de-DE', timezone_id='Europe/Berlin',
                             color_scheme=thema, accept_downloads=True)
        pg = ctx.new_page()
        pg.on('pageerror', lambda e: fehler.append(str(e)))
        pg.on('console', lambda m: fehler.append(m.text) if m.type == 'error' and 'Failed to load resource' not in m.text else None)
        pg.goto(B + '/admin/'); pg.wait_for_selector('#screen-auth:not([hidden])')
        pg.fill('#au-user', user); pg.fill('#pw1', pw); pg.click('#auth-submit')
        pg.wait_for_selector('#screen-auth', state='hidden')
        pg.goto(B + pfad); pg.wait_for_selector(ziel)
        return ctx, pg

    WOCHE = '2026-43'
    ctx, pg = anmelden('mark', 'Laden-Konto-22', '198.51.100.11', pfad=f'/module/rewe/#woche/{WOCHE}')
    pg.wait_for_selector('.rw-gruppe')
    check('Meine Woche: Adresse #woche/<KW>, Titel, Gruppe Lager Nord, nur Mark-Reiter', pg.url.endswith('#woche/' + WOCHE) and 'KW 43' in text(pg, '#rw-w-titel')
          and 'Zentrallager Nordheim' in text(pg, '#rw-w-gruppen') and pg.is_hidden('#rw-tabs [data-tab="zuweisung"]') and pg.is_hidden('#rw-tabs [data-tab="abschluss"]'))
    pg.click('#rw-w-filter [data-f="alle"]')
    zeile = '.rw-zeile:has-text("Musterdorf")'
    pg.click(f'{zeile} [data-s="kein"]'); pg.wait_for_selector(f'{zeile} [data-s="kein"][aria-pressed="true"]')
    pg.click(f'{zeile} [data-s="bestellt"]'); pg.wait_for_selector(f'{zeile} [data-menge="rb"]')
    pg.click(f'{zeile} [data-menge="rb"][data-d="1"]'); pg.wait_for_function("document.querySelector('.rw-zeile[data-id] .rw-stepper b') && [...document.querySelectorAll('.rw-zeile')].some(z => z.textContent.includes('Musterdorf') && z.querySelector('.rw-stepper b') && z.querySelector('.rw-stepper b').textContent === '2')")
    pg.fill(f'{zeile} [data-notiz]', 'Kunde ruft zurück'); pg.press(f'{zeile} [data-notiz]', 'Tab')
    pg.wait_for_function("[...document.querySelectorAll('.rw-zeile')].some(z => z.textContent.includes('Musterdorf') && z.querySelector('[data-notiz]') && z.querySelector('[data-notiz]').value === 'Kunde ruft zurück' && !z.querySelector('[data-s][disabled]'))")
    aw = next(x for g in M.call(f'rewe_woche&kw={WOCHE}')[1]['woche']['gruppen'] for x in g['maerkte'] if x['ort'] == 'Musterdorf')['antwort']
    check('Antwort per Klick: bestellt, 2 Roastbeef, Notiz gespeichert', aw['status'] == 'bestellt' and aw['mengen'] == {'rb': 2} and aw['notiz'] == 'Kunde ruft zurück', aw)
    check('Verlauf mit acht Zellen und Legende', pg.locator(f'{zeile} .rw-z').count() == 8 and pg.is_visible('.rw-legende'))
    pg.screenshot(path=os.path.join(SH, 'rewe_woche_1280.png'), full_page=True)
    pg.click('#rw-w-heute'); pg.wait_for_function(f"location.hash === '#woche/{ABFRAGE}' && document.querySelector('#rw-w-titel').textContent.includes('KW {int(ABFRAGE[5:])}')")
    pg.wait_for_selector('.rw-zeile:has-text("Beispielstadt")')
    faellig = next(x for g in M.call(f'rewe_woche&kw={ABFRAGE}')[1]['woche']['gruppen'] for x in g['maerkte'] if x['ort'] == 'Beispielstadt')['dauerFaellig']
    box = pg.locator('.rw-zeile:has-text("Beispielstadt")')
    if faellig:   # Dauerauftrag 14-tägig ab KW 2: in geraden Wochen fällig
        box.locator('[data-dauer-ok]').click(); pg.wait_for_selector('.rw-zeile:has-text("Beispielstadt") [data-s="bestellt"][aria-pressed="true"]')
        check('Abfragewoche: Dauerauftrag fällig → bestätigt (Roastbeef + Filetkopf)', pg.locator('.rw-zeile:has-text("Beispielstadt") .rw-stepper b').all_inner_texts()[:2] == ['1', '1'])
    else:
        check('Abfragewoche: Dauerauftrag nicht fällig → „Trotzdem abfragen“', box.locator('.rw-nichtfaellig').count() == 1)
        box.locator('[data-offen]').click(); pg.wait_for_selector('.rw-zeile:has-text("Beispielstadt") [data-s="bestellt"]')
    # Marktakte mit Notiz
    pg.click('#rw-tabs [data-tab="maerkte"]'); pg.wait_for_selector('#rw-m-tabelle tbody tr')
    check('Märkte: nur eigene (2)', pg.locator('#rw-m-tabelle tbody tr').count() == 2 and pg.is_hidden('#rw-m-neu'))
    pg.click('#rw-m-tabelle tr:has-text("Beispielstadt")'); pg.wait_for_selector('.rw-heat')
    check('Marktakte: Adresse #markt/<id>, Jahre im Verlauf, Kennzahlen, Ansprechpartner', '#markt/' in pg.url and pg.locator('.rw-heat .jahr').count() >= 2
          and 'Frau Test' in text(pg, '#rw-akte') and 'Nur gereifte Ware' in text(pg, '#rw-akte'))
    pg.click('#rw-a-notiz'); pg.wait_for_selector('#dlg-notiz[open]')
    pg.select_option('#nz-art', 'telefonat'); pg.fill('#nz-text', 'Vac. Set angeboten'); pg.click('#nz-form button[type="submit"]')
    pg.wait_for_selector('#dlg-notiz[open]', state='hidden'); pg.wait_for_selector('.rw-verlaufsliste li:has-text("Vac. Set angeboten")')
    check('Notiz über die Eingabemaske', 'Telefonat' in text(pg, '.rw-verlaufsliste'))
    pg.click('#rw-a-bearbeiten'); pg.wait_for_selector('#dlg-markt[open]')
    check('Bearbeiten ohne Stammdaten (nur verwalten)', pg.is_hidden('#mk-stamm'))
    pg.select_option('#mk-tag', 'mi'); pg.click('#mk-form button[type="submit"]'); pg.wait_for_selector('#dlg-markt[open]', state='hidden')
    pg.wait_for_selector('.rw-kv:has-text("Mittwoch")')
    check('Abfragetag geändert', 'Mittwoch' in text(pg, '.rw-kv'))
    pg.screenshot(path=os.path.join(SH, 'rewe_akte_1280.png'), full_page=True)
    ctx.close()

    ctx, pg = anmelden('lena', 'Leitung-Konto-11', '198.51.100.10')
    pg.wait_for_selector('.rw-gruppe, .rw-leer')
    check('Leitung: alle Reiter, Listenwahl', pg.is_visible('#rw-tabs [data-tab="zuweisung"]') and pg.is_visible('#rw-tabs [data-tab="abschluss"]') and pg.is_visible('#rw-w-wer'))
    pg.select_option('#rw-w-wer', '*'); pg.wait_for_function("document.querySelector('#rw-w-titel').textContent.startsWith('Alle Listen')")
    pg.click('#rw-tabs [data-tab="zuweisung"]'); pg.wait_for_selector('#rw-z-tabelle tbody tr [data-wahl]')
    check('Zuweisung: Team-Karten, alle Märkte', pg.locator('#rw-z-personen [data-person]').count() >= 4 and pg.locator('#rw-z-tabelle [data-wahl]').count() == 4)
    pg.check('#rw-z-tabelle tr:has-text("Neustadt") [data-wahl]'); pg.wait_for_selector('#rw-z-bulk:not([hidden])')
    pg.select_option('#rw-z-nutzer', ids['evi']); pg.fill('#rw-z-ab', ABFRAGE); pg.click('#rw-z-los')
    pg.wait_for_selector('#rw-z-tabelle tr:has-text("Neustadt"):has-text("Evi Eins")')
    check('Zuweisung über die Oberfläche: Neustadt → Evi', any(m['ort'] == 'Neustadt' and m['zustaendigName'] == 'Evi Eins' for m in LE.call('rewe_maerkte')[1]['maerkte']))
    pg.screenshot(path=os.path.join(SH, 'rewe_zuweisung_1280.png'), full_page=True)
    pg.goto(B + f'/module/rewe/#abschluss/{WOCHE}'); pg.wait_for_selector('.rw-schritte')
    check('Wochenabschluss: Schritte, Vorschau mit Aufträgen', pg.locator('.rw-schritte li').count() == 4 and 'Vorschau Excel für WinAB' in text(pg, '#rw-a-inhalt'))
    with pg.expect_download() as dl:
        pg.click('#rw-a-export')
    check('Excel über die Oberfläche (KW 43, Bestellung aus der Oberfläche)', dl.value.suggested_filename == 'REWE_WinAB_2026_KW43.xlsx', dl.value.suggested_filename)
    pg.wait_for_selector('#rw-a-erfasst:not([hidden])')
    pg.screenshot(path=os.path.join(SH, 'rewe_abschluss_1280.png'), full_page=True)
    pg.click('#rw-tabs [data-tab="einstellungen"]'); pg.wait_for_selector('#rw-e-lw tbody tr')
    namen = json.loads(json.dumps(PAKET).replace('"Liste A"', '"Liste Mark"').replace('"Liste B"', '"Evi"'))
    pg.set_input_files('#rw-e-datei', files=[{'name': 'paket.json', 'mimeType': 'application/json', 'buffer': json.dumps(namen).encode()}])
    pg.wait_for_selector('#rw-e-vorschau .rw-vorschau')
    wahl = {s.get_attribute('data-liste'): s.input_value() for s in pg.locator('#rw-e-vorschau [data-liste]').all()}
    check('Einstellungen: Lieferwege, Artikel, Datenpaket-Vorschau mit Listen-Zuordnung', pg.locator('#rw-e-lw tbody tr').count() == 3
          and len(wahl) == 3 and pg.is_visible('#rw-e-ersetzen'), wahl)
    check('Datenpaket: Listen nach Vornamen vorgeschlagen („Liste Mark“ → Mark, „Evi“ → Evi, „Liste C“ → niemand)',
          wahl == {'Evi': ids['evi'], 'Liste C': '', 'Liste Mark': ids['mark']} and '2 Listen wurden' in text(pg, '#rw-e-vorschlag'), wahl)
    pg.screenshot(path=os.path.join(SH, 'rewe_einstellungen_1280.png'), full_page=True)
    ctx.close()

    for breite, thema in ((390, 'light'), (390, 'dark'), (1280, 'dark')):
        ctx, pg = anmelden('mark', 'Laden-Konto-22', '198.51.100.11', breite, thema, 844 if breite < 500 else 900, pfad=f'/module/rewe/#woche/{WOCHE}')
        pg.wait_for_selector('.rw-gruppe')
        breit = []
        for ziel, warte in ((f'#woche/{WOCHE}', '.rw-gruppe'), ('#maerkte', '#rw-m-tabelle tbody tr'), ('#markt/' + maerkte['Beispielstadt']['id'], '.rw-heat')):
            pg.goto(B + '/module/rewe/' + ziel); pg.wait_for_selector(warte)
            if not pg.evaluate('document.documentElement.scrollWidth <= window.innerWidth + 1'):
                breit.append(ziel)
            if breite == 390:
                pg.screenshot(path=os.path.join(SH, f'rewe_{ziel.split("/")[0][1:]}_{breite}_{thema}.png'), full_page=True)
        check(f'{breite}px {thema}: kein waagerechtes Scrollen', not breit, breit)
        if thema == 'dark':
            check(f'{breite}px: dunkler Grund', pg.evaluate('getComputedStyle(document.body).backgroundColor') == 'rgb(18, 16, 17)')
        ctx.close()

    ctx, pg = anmelden('otto', 'Ohne-Konto-66', '198.51.100.15', ziel='#rw-err:not([hidden])')
    check('ohne REWE-Recht: Hinweis statt Inhalt', 'Berechtigung' in text(pg, '#rw-err') and pg.is_hidden('#rw-app'))
    ctx.close()
    br.close()
check('keine JavaScript-Fehler', not fehler, fehler)

print(f'{sum(res)}/{len(res)} passed')
