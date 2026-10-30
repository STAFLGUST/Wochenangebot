"""Qualitätssicherung – Audit-Bericht (app/api/module/qs/bericht.php, Aktionen qs_bericht und qs_bericht_excel, Reiter „Audit“).

Geprüft: Rechenkern mit festem Stichtag (php -r, erfundene Daten): Monatsrechnung rückwärts über den Jahreswechsel (Fehler bis
Oktober 2026: Januar–Juni ergaben Monat ≤ 0) und Fälligkeit im Januar; Kalender-Intervalle (Arbeitstage ohne Hamburger Feiertage,
monatlich, laufende Periode erst mit Nachweis, Nachweise nach dem Stichtag zählen nicht, angebrochene Periode am Stichtag,
stornierte Nachweise zählen nie, versäumte Bereiche, Zählung im Monat des Periodenendes, gleiche Versäumnisse wie die Fälligkeit);
längere Intervalle (eingehalten, verspätet mit Tagen, überfällig zum Stichtag, fällig am Stichtag, noch nie geprüft, vor dem
Zeitraum erledigt bleibt weg); „bei Bedarf“ und inaktive Prüfungen; Abweichungen mit Maßnahme; Maßnahmen (vor dem Zeitraum
abgeschlossen bzw. nach dem Stichtag erfasst bleiben weg, wirksam abgeschlossen, verworfen, offen, überfällig); Dokumentenlenkung
(Freigaben, Überprüfungen, Einziehungen im Zeitraum; gültig, in Arbeit, nie freigegeben, Überprüfung fällig); Handbuch-Änderungen;
Zeitraum (Pflicht, gültiges Datum, nicht in der Zukunft, Von ≤ Bis, höchstens zwei Jahre); Excel (Blätter, Werte, Texte mit „=“
bleiben Text, öffnet in openpyxl). Über die Schnittstelle: Rechte (ohne Anmeldung 401, ohne QM-Recht 403, „ansehen“ genügt),
Fehler 422, gleiche Versäumnisse wie die Fälligkeit im Prüfplan, Dokumente und Handbuch wie im QM-System, Abweichung mit Maßnahme,
storniert, Excel mit Kopfzeilen und Protokoll. Oberfläche: Reiter „Audit“ erstellt den Bericht der letzten 12 Monate, Zeitraum
wählen, Fehler bei Zukunft, Excel-Download, Druckansicht (nur der Bericht, PDF A4 quer), 1280/390 px hell und dunkel ohne
waagerechtes Scrollen, keine JavaScript-Fehler. Nur erfundene Testdaten.
"""
import datetime
import io
import json
import math
import os
import re
import subprocess
import urllib.error
import urllib.request
import zipfile

import openpyxl
from playwright.sync_api import sync_playwright

from kunden_hilfe import Client, B, API
from umgebung import SH

HIER = os.path.dirname(os.path.abspath(__file__))
REPO = os.path.abspath(os.path.join(HIER, '..', '..'))
PAKET_DATEI = os.path.join(HIER, '..', 'fixtures', 'qs', 'datenpaket.json')
HEUTE = datetime.date.today()
res = []
fehler = []


def check(n, c, i=''):
    res.append(bool(c)); print(('PASS ' if c else 'FAIL ') + n + (('  ' + str(i)[:400]) if i else ''))


def tag(n):
    return (HEUTE + datetime.timedelta(days=n)).isoformat()


# ── Rechenkern mit festem Stichtag (ohne Server) ──
PHP = r"""define('WA_API', 1);
require $argv[1] . '/app/api/module/hr/urlaub.php';
require $argv[1] . '/app/api/module/qs/qs.php';
require $argv[1] . '/app/api/module/qs/bericht.php';
$f = json_decode(file_get_contents('php://stdin'), true); $o = [];
if (isset($f['monate'])) $o['monate'] = array_map(function ($x) { return qs_plus_monate($x[0], $x[1]); }, $f['monate']);
if (isset($f['faellig'])) { $x = $f['faellig']; $o['faellig'] = qs_faelligkeit($x['p'], qs_nachweise_nach_pruefung($x['n'])[$x['p']['id']] ?? [], $x['heute']); }
if (isset($f['zeitraum'])) foreach ($f['zeitraum'] as $z) { try { qs_bericht_zeitraum($z[0], $z[1], $z[2]); $o['zeitraum'][] = 'ok'; } catch (InvalidArgumentException $e) { $o['zeitraum'][] = $e->getMessage(); } }
if (isset($f['bericht'])) { $b = $f['bericht'];
  $o['bericht'] = qs_bericht_daten($b['stamm'], $b['dokumente'], $b['handbuch'], $b['plan'], $b['nachweise'], $b['massnahmen'], $b['namen'], $b['von'], $b['bis'], $b['heute']);
  if (!empty($b['xlsx'])) file_put_contents($b['xlsx'], qs_bericht_xlsx($o['bericht'], 'Quinn Prüfer')); }
echo json_encode($o);"""


def kern(fall):
    p = subprocess.run(['php', '-r', PHP, REPO], input=json.dumps(fall), capture_output=True, text=True)
    try:
        return json.loads(p.stdout)
    except ValueError:
        print(p.stdout[-800:], p.stderr[-800:])
        return {}


STAMM = {'kapitel': [{'id': '1.1', 'titel': 'Managementsystem'}, {'id': '1.3', 'titel': 'HACCP'}, {'id': '1.4', 'titel': 'Qualitätssicherung'}, {'id': '2.1', 'titel': 'Wareneingang'}],
         'personen': [], 'einstellungen': {'ueberpruefungMonate': 36}}
_n = [0]


def nw(p, d, e='io', storno=None, massnahme=''):
    _n[0] += 1
    return {'id': f'n{_n[0]}', 'pruefung': p, 'datum': d, 'ergebnis': e, 'bemerkung': 'zu warm' if e == 'abweichung' else '', 'durch': 'Uwe Unterwegs',
            'massnahme': massnahme, 'erfasst': {'t': d + 'T08:00:00Z'}, 'storno': storno}


def bericht(plan, nachweise, von, bis, heute, massnahmen=(), dokumente=(), handbuch=(), xlsx=''):
    return kern({'bericht': {'stamm': STAMM, 'dokumente': list(dokumente), 'handbuch': list(handbuch), 'plan': plan, 'nachweise': nachweise, 'massnahmen': list(massnahmen),
                             'namen': {'u1': 'Uwe Unterwegs'}, 'von': von, 'bis': bis, 'heute': heute, 'xlsx': xlsx}}).get('bericht', {})


k = kern({'monate': [['2027-01-15', -6], ['2026-03-31', -1], ['2026-06-30', -6], ['2025-12-31', 2], ['2026-01-31', -13], ['2026-10-31', 4]]})
check('Monate rückwärts über den Jahreswechsel (Januar −6 = Juli des Vorjahres; Monatsende bleibt gültig)',
      k.get('monate') == ['2026-07-15', '2026-02-28', '2025-12-30', '2026-02-28', '2024-12-31', '2027-02-28'], k)

# Arbeitstäglich über Weihnachten und Neujahr (25./26.12. und 01.01. Feiertage in Hamburg), Stichtag = heute 15.01.2027
PA = {'id': 'pA', 'kapitel': '2.1', 'titel': 'Wareneingang', 'intervall': 'arbeitstaeglich', 'start': '2026-12-21', 'aktiv': True, 'zustaendig': 'QB', 'ausfuehrende': ['u1']}
NA = [nw('pA', d) for d in ('2026-12-21', '2026-12-22', '2026-12-23', '2026-12-28', '2027-01-15')]
b = bericht([PA], NA, '2026-12-01', '2027-01-15', '2027-01-15')
pa = (b.get('pruefungen') or [{}])[0]
check('arbeitstäglich: 18 fällige Arbeitstage (ohne Feiertage), 5 erfüllt, heute zählt mit Nachweis',
      (pa.get('faellig'), pa.get('erfuellt'), pa.get('versaeumt')) == (18, 5, 13) and pa.get('art') == 'kalender', pa)
check('arbeitstäglich: versäumte Bereiche zusammengefasst (24.12. und 29.12.–14.01.)', pa.get('bereiche') == [['2026-12-24', '2026-12-24', 1], ['2026-12-29', '2027-01-14', 12]], pa.get('bereiche'))
check('arbeitstäglich: je Monat erfüllt/fällig und Anzahl Nachweise', pa.get('monate') == {'2026-12': {'n': 4, 'f': 8, 'e': 4}, '2027-01': {'n': 1, 'f': 10, 'e': 1}}, pa.get('monate'))
check('Durchführung mit Namen, Kennzahlen mit Quote', pa.get('ausfuehrende') == ['Uwe Unterwegs'] and b['kennzahlen']['pruefplan']['quote'] == round(5 / 18, 4), b.get('kennzahlen'))
f = kern({'faellig': {'p': PA, 'n': NA, 'heute': '2027-01-15'}}).get('faellig', {})
check('gleiche Versäumnisse wie die Fälligkeit im Prüfplan – und der Jahresplan rechnet im Januar (7 Monate ab Juli)',
      f.get('versaeumtAnzahl') == 13 and [m['m'] for m in f.get('monate', [])] == ['2026-07', '2026-08', '2026-09', '2026-10', '2026-11', '2026-12', '2027-01'], f)
b = bericht([PA], NA[:-1], '2026-12-01', '2027-01-15', '2027-01-15')
check('laufende Periode ohne Nachweis zählt noch nicht (heute offen)', (b['pruefungen'][0]['faellig'], b['pruefungen'][0]['erfuellt']) == (17, 4), b['pruefungen'][0])

# Monatlich mit Stichtag mitten im Monat, Storno, Abweichung mit Maßnahme
M1 = {'id': 'pM1', 'kapitel': '1.4', 'titel': 'Monatlich eins', 'intervall': 'monatlich', 'start': '2026-06-01', 'aktiv': True}
M2 = {'id': 'pM2', 'kapitel': '1.4', 'titel': 'Monatlich zwei', 'intervall': 'monatlich', 'start': '2026-12-01', 'aktiv': True}
MA = {'id': 'm1', 'nr': '2026-001', 'quelle': 'pruefung', 'kapitel': '1.4', 'feststellung': 'zu warm', 'massnahme': 'Kühlung prüfen', 'zustaendig': 'QB', 'frist': '2026-09-01',
      'status': 'offen', 'erfasst': {'t': '2026-08-05T10:00:00Z'}}
NM = [nw('pM1', '2026-06-10'), nw('pM1', '2026-08-05', 'abweichung', massnahme='m1'), nw('pM1', '2026-09-30'), nw('pM1', '2026-10-03', storno={'t': '2026-10-04T08:00:00Z', 'grund': 'x'}),
      nw('pM1', '2026-12-28'), nw('pM2', '2026-12-05')]
b = bericht([M1, M2], NM, '2026-07-01', '2026-12-20', '2027-01-15', massnahmen=[MA])
pm = {x['id']: x for x in b.get('pruefungen', [])}
check('monatlich: Juli–November gezählt, angebrochener Dezember ohne Nachweis bis zum Stichtag nicht (Nachweis 28.12. zählt nicht)',
      (pm['pM1']['faellig'], pm['pM1']['erfuellt'], pm['pM1']['versaeumt']) == (5, 2, 3) and '2026-12' not in pm['pM1']['monate'], pm['pM1'])
check('monatlich: stornierter Nachweis zählt nicht (Oktober versäumt)', pm['pM1']['bereiche'] == [['2026-07', '2026-07', 1], ['2026-10', '2026-11', 2]], pm['pM1']['bereiche'])
check('monatlich: angebrochener Dezember mit Nachweis vor dem Stichtag zählt als erfüllt', pm['pM2']['monate'] == {'2026-12': {'n': 1, 'f': 1, 'e': 1}}, pm['pM2'])
kz = b['kennzahlen']['pruefplan']
check('Kennzahlen: fällig, erfüllt, Nachweise im Zeitraum (ohne Storno, ohne nach dem Stichtag), Abweichungen, storniert',
      (kz['faellig'], kz['erfuellt'], kz['versaeumt'], kz['nachweise'], kz['abweichungen'], kz['storniert']) == (6, 3, 3, 3, 1, 1), kz)
check('Abweichung mit Datum, Prüfung, Feststellung und Maßnahme', b['abweichungen'] == [{'datum': '2026-08-05', 'kapitel': '1.4', 'pruefung': 'Monatlich eins', 'bemerkung': 'zu warm',
                                                                             'durch': 'Uwe Unterwegs', 'massnahme': {'nr': '2026-001', 'status': 'offen'}}], b['abweichungen'])
check('Maßnahme der Abweichung: offen und über der Frist', [m['nr'] for m in b['massnahmen']] == ['2026-001'] and b['massnahmen'][0]['ueberfaellig'] is True
      and b['kennzahlen']['massnahmen'] == {'erfasst': 1, 'abgeschlossen': 0, 'offen': 1, 'ueberfaellig': 1}, b['massnahmen'])

# Längere Intervalle, „bei Bedarf“, inaktiv; Maßnahmen, Dokumente, Handbuch – Kalenderjahr 2026, Stichtag 31.12.2026
R = lambda i, t, iv, s, aktiv=True: {'id': i, 'kapitel': '1.1', 'titel': t, 'intervall': iv, 'start': s, 'aktiv': aktiv}  # noqa: E731
PLAN = [R('J', 'Jährlich', 'jaehrlich', '2025-03-01'), R('H', 'Halbjährlich', 'halbjahr', '2026-03-01'), R('Q', 'Vierteljährlich nie', 'quartal', ''),
        R('V', 'Vierteljährlich', 'quartal', '2026-06-30'), R('J2', 'Jährlich Stichtag', 'jaehrlich', '2026-12-31'), R('A', 'Bei Bedarf', 'anlass', ''),
        R('I', 'Inaktiv mit Nachweis', 'woechentlich', '2026-01-01', False), R('I2', 'Inaktiv ohne Nachweis', 'woechentlich', '2026-01-01', False)]
NR = [nw('J', '2025-02-20'), nw('J', '2026-04-10'), nw('V', '2026-06-30'), nw('V', '2026-09-15'), nw('A', '2026-03-03'), nw('A', '2026-03-20'), nw('I', '2026-05-05')]
MASSN = [{'id': 'm2', 'nr': '2025-001', 'quelle': 'intern', 'kapitel': '1.1', 'feststellung': 'alt', 'massnahme': '', 'zustaendig': '', 'frist': '', 'status': 'wirksam',
          'erfasst': {'t': '2025-05-01T08:00:00Z'}, 'wirksamkeit': {'t': '2025-06-01T08:00:00Z', 'notiz': 'ok'}},
         {'id': 'm6', 'nr': '2025-002', 'quelle': 'extern', 'kapitel': '1.4', 'bewertung': 'B', 'feststellung': 'Thermometer', 'massnahme': 'kalibrieren', 'zustaendig': 'QB',
          'frist': '2027-02-01', 'status': 'in_arbeit', 'erfasst': {'t': '2025-11-01T08:00:00Z'}},
         {'id': 'm3', 'nr': '2026-001', 'quelle': 'intern', 'kapitel': '1.1', 'feststellung': '=SUMME(A1:A9)', 'massnahme': '@Formel', 'zustaendig': 'QB', 'frist': '2026-02-20',
          'status': 'wirksam', 'erfasst': {'t': '2026-02-01T08:00:00Z'}, 'wirksamkeit': {'t': '2026-03-01T08:00:00Z', 'notiz': 'geprüft'}},
         {'id': 'm5', 'nr': '2026-002', 'quelle': 'kunde', 'kapitel': '', 'feststellung': 'Etikett', 'massnahme': '', 'zustaendig': '', 'frist': '', 'status': 'verworfen',
          'erfasst': {'t': '2026-04-01T08:00:00Z'}, 'geaendert': {'t': '2026-05-01T08:00:00Z'}},
         {'id': 'm4', 'nr': '2027-001', 'quelle': 'intern', 'kapitel': '', 'feststellung': 'neu', 'massnahme': '', 'zustaendig': '', 'frist': '', 'status': 'offen', 'erfasst': {'t': '2027-01-05T08:00:00Z'}}]
V_ = lambda nr, datum, status, **x: {'id': 'v' + str(nr) + status, 'nr': nr, 'datum': datum, 'status': status, **x}  # noqa: E731
DOKS = [{'id': 'd1', 'kapitel': '1.4', 'nr': '1.4-F01', 'titel': 'Probenahme', 'art': 'F', 'versionen': [V_(1, '2025-05-01', 'ersetzt'),
         V_(2, '2026-03-10', 'freigegeben', aenderungen='Spalte ergänzt', grund='Audit', kuerzel={'erstellt': 'AB', 'geprueft': 'CD', 'freigegeben': 'EF'})],
         'ueberprueft': {'t': '2026-11-01T09:00:00Z', 'von': 'Quinn', 'notiz': 'passt'}},
        {'id': 'd2', 'kapitel': '1.4', 'nr': '1.4-A01', 'titel': 'Alt', 'art': 'A', 'versionen': [V_(1, '2020-01-01', 'freigegeben')], 'eingezogen': {'t': '2026-06-01T08:00:00Z', 'von': 'Quinn', 'grund': 'veraltet'}},
        {'id': 'd3', 'kapitel': '1.3', 'nr': '1.3-L01', 'titel': 'HACCP-Analyse', 'art': 'L', 'versionen': [V_(1, '2020-01-01', 'freigegeben')]},
        {'id': 'd4', 'kapitel': '2.1', 'nr': '2.1-F01', 'titel': 'Wareneingang', 'art': 'F', 'versionen': [V_(1, '2025-12-01', 'freigegeben'), V_(None, '', 'entwurf')]},
        {'id': 'd5', 'kapitel': '2.1', 'nr': '2.1-F02', 'titel': 'Neu', 'art': 'F', 'versionen': [V_(None, '', 'entwurf')]},
        {'id': 'd7', 'kapitel': '2.1', 'nr': '2.1-E01', 'titel': 'Verworfen', 'art': 'E', 'versionen': [V_(None, '', 'verworfen')]},
        {'id': 'd6', 'kapitel': '1.4', 'nr': '', 'titel': 'Befunde', 'art': 'O'}]
HB = [{'id': 'h1', 'kapitel': '1.1', 'nr': '2.0.15', 'titel': 'Freigabe', 'aenderungen': [{'datum': '2025-12-31', 'was': 'alt'}, {'datum': '2026-05-05', 'was': 'Text ergänzt', 'grund': 'Audit', 'von': 'Quinn'}]}]
XL = os.path.join(SH, 'qs_bericht_kern.xlsx')
b = bericht(PLAN, NR, '2026-01-01', '2026-12-31', '2027-01-15', massnahmen=MASSN, dokumente=DOKS, handbuch=HB, xlsx=XL)
pr = {x['id']: x for x in b.get('pruefungen', [])}
F = lambda i: [(x['faellig'], x['erledigt'], x['status'], x['tage']) for x in pr[i]['fristen']]  # noqa: E731
check('jährlich: Frist 20.02.2026 erst am 10.04. erfüllt → verspätet 49 Tage; Frist vor dem Zeitraum erledigt bleibt weg', F('J') == [('2026-02-20', '2026-04-10', 'verspaetet', 49)], F('J'))
check('halbjährlich ohne Nachweis: seit 01.03. überfällig, 305 Tage zum Stichtag', F('H') == [('2026-03-01', '', 'ueberfaellig', 305)], F('H'))
check('ohne Nachweis und ohne erste Fälligkeit: noch nie geprüft', F('Q') == [('', '', 'nie', None)], F('Q'))
check('vierteljährlich: zweimal eingehalten (auch am Fälligkeitstag), dann 16 Tage überfällig',
      F('V') == [('2026-06-30', '2026-06-30', 'eingehalten', None), ('2026-09-30', '2026-09-15', 'eingehalten', None), ('2026-12-15', '', 'ueberfaellig', 16)], F('V'))
check('fällig am Stichtag: noch nicht überfällig', F('J2') == [('2026-12-31', '', 'offen', None)], F('J2'))
check('bei Bedarf: nur Nachweise je Monat; inaktive Prüfung nur mit Nachweis im Zeitraum', pr['A']['art'] == 'anlass' and pr['A']['monate'] == {'2026-03': {'n': 2}}
      and pr['I']['art'] == 'inaktiv' and pr['I']['nachweise'] == 1 and 'I2' not in pr and pr['A']['faellig'] is None, (pr.get('A'), pr.get('I')))
kz = b['kennzahlen']['pruefplan']
check('Kennzahlen der Fristen und Prüfungen (inaktive zählen nicht)', (kz['fristen'], kz['eingehalten'], kz['verspaetet'], kz['ueberfaellig'], kz['offen'], kz['nie'], kz['pruefungen'],
                                                                       kz['nachweise'], kz['faellig'], kz['quote']) == (7, 2, 1, 2, 1, 1, 6, 6, 0, None), kz)
check('Maßnahmen: vor dem Zeitraum abgeschlossen und nach dem Stichtag erfasst bleiben weg', [m['nr'] for m in b['massnahmen']] == ['2025-002', '2026-001', '2026-002'], b['massnahmen'])
check('Maßnahmen: 2 neu, 1 wirksam abgeschlossen (verworfen zählt nicht), 1 offen, keine überfällig', b['kennzahlen']['massnahmen'] == {'erfasst': 2, 'abgeschlossen': 1, 'offen': 1, 'ueberfaellig': 0}
      and [m['abgeschlossen'] for m in b['massnahmen']] == ['', '2026-03-01', '2026-05-01'], b['kennzahlen']['massnahmen'])
check('Dokumente: gültig (auch in Überarbeitung), in Arbeit, nie freigegeben, Überprüfung fällig, Vorgänge im Zeitraum',
      b['kennzahlen']['dokumente'] == {'gueltig': 3, 'inArbeit': 2, 'ohneFreigabe': 1, 'faellig': 1, 'freigaben': 1, 'ueberprueft': 1, 'eingezogen': 1}, b['kennzahlen']['dokumente'])
check('Vorgänge nach Datum: Freigabe Version 2 mit Kürzeln, Einziehung mit Grund, Überprüfung mit Notiz',
      [(x['datum'], x['vorgang'], x['nr']) for x in b['dokumente']['vorgaenge']] == [('2026-03-10', 'freigabe', '1.4-F01'), ('2026-06-01', 'eingezogen', '1.4-A01'), ('2026-11-01', 'ueberprueft', '1.4-F01')]
      and b['dokumente']['vorgaenge'][0]['freigegeben'] == 'EF' and b['dokumente']['vorgaenge'][1]['notiz'] == 'veraltet', b['dokumente']['vorgaenge'])
check('Überprüfung fällig: nur gültige, nicht eingezogene Dokumente (seit 2020, fällig ab 2023)', b['dokumente']['faellig'] == [{'kapitel': '1.3', 'nr': '1.3-L01', 'titel': 'HACCP-Analyse', 'art': 'L',
                                                                                            'seit': '2020-01-01', 'faelligAb': '2023-01-01'}], b['dokumente']['faellig'])
check('Handbuch: nur Änderungen im Zeitraum', [(x['datum'], x['was']) for x in b['handbuch']] == [('2026-05-05', 'Text ergänzt')] and b['kennzahlen']['handbuch'] == 1, b['handbuch'])

wb = openpyxl.load_workbook(XL)
check('Excel: acht Blätter', wb.sheetnames == ['Zusammenfassung', 'Prüfplan', 'Versäumnisse', 'Abweichungen', 'Maßnahmen', 'Dokumente', 'Überprüfung fällig', 'Handbuch'], wb.sheetnames)
zs = {r[1]: r[2] for r in wb['Zusammenfassung'].iter_rows(values_only=True) if r and len(r) > 2 and r[1]}
check('Excel: Zusammenfassung mit Zeitraum, Ersteller und Kennzahlen', wb['Zusammenfassung']['A2'].value == 'Zeitraum: 01.01.2026 – 31.12.2026'
      and 'Quinn Prüfer' in wb['Zusammenfassung']['A3'].value and zs.get('Fristen längerer Intervalle') == 7 and zs.get('davon überfällig bzw. noch nie geprüft') == 3, zs)
vs = [r for r in wb['Versäumnisse'].iter_rows(min_row=2, values_only=True)]
check('Excel: Versäumnisse und Fristen ohne eingehaltene', sorted((r[1], r[3]) for r in vs) == sorted([('Jährlich', 'verspätet'), ('Halbjährlich', 'überfällig'), ('Vierteljährlich nie', 'noch nie geprüft'),
                                                                                                     ('Vierteljährlich', 'überfällig'), ('Jährlich Stichtag', 'fällig am Stichtag')]), vs)
ms = {r[0]: r for r in wb['Maßnahmen'].iter_rows(min_row=2, values_only=True)}
xml = zipfile.ZipFile(XL).read('xl/worksheets/sheet5.xml').decode()
check('Excel: Texte mit = und @ bleiben Text (keine Formeln)', ms['2026-001'][4] == '=SUMME(A1:A9)' and ms['2026-001'][5] == '@Formel' and '<f>' not in xml and '<f ' not in xml, ms.get('2026-001'))
check('Excel: Prüfplan je Prüfung, Monatsspalten Jan–Dez', [c.value for c in wb['Prüfplan'][1]][14:] == ['Jan 26', 'Feb 26', 'Mär 26', 'Apr 26', 'Mai 26', 'Jun 26', 'Jul 26', 'Aug 26', 'Sep 26', 'Okt 26', 'Nov 26', 'Dez 26']
      and wb['Prüfplan'].max_row >= 8, [c.value for c in wb['Prüfplan'][1]])
z = kern({'zeitraum': [['2026-01-01', '2027-01-16', '2027-01-15'], ['2026-02-01', '2026-01-01', '2027-01-15'], ['2025-01-14', '2027-01-15', '2027-01-15'],
                       ['2025-01-15', '2027-01-15', '2027-01-15'], ['2026-02-30', '2026-03-01', '2027-01-15'], ['', '2026-03-01', '2027-01-15']]}).get('zeitraum', [])
check('Zeitraum: nicht in der Zukunft, Von ≤ Bis, höchstens zwei Jahre (730 Tage ok), gültiges Datum, Pflicht',
      len(z) == 6 and 'Zukunft' in z[0] and 'Von muss vor Bis' in z[1] and 'zwei Jahre' in z[2] and z[3] == 'ok' and 'gültiges Datum' in z[4] and 'bitte ein Datum' in z[5], z)

# ── Schnittstelle ──
A = Client('198.51.100.2')
check('Administrator angemeldet', A.call('login', {'username': 'admin', 'password': 'Sicheres-Passwort-1'})[0] == 200)
NUTZER = (('lisa', 'Lisa Lesen', 'viewer', 'Leserin-Konto-12', {'grant': ['qs.sehen']}),
          ('bernd', 'Bernd Bearbeiter', 'editor', 'Bearbeit-Konto-34', {'grant': ['qs.bearbeiten']}),
          ('quinn', 'Quinn QMB', 'editor', 'Freigabe-Konto-56', {'grant': ['qs.freigeben']}),
          ('vera', 'Vera Verwalten', 'editor', 'Verwalt-Konto-90', {'grant': ['qs.verwalten']}),
          ('otto', 'Otto Ohne', 'editor', 'Ohne-Konto-11', {}))
for u, n, r, p, extra in NUTZER:
    st, r_ = A.call('user_save', {'id': '', 'name': n, 'username': u, 'role': r, 'active': True, 'password': p, 'mustChange': False, **extra})
    assert st == 200, r_
K = {}
for i, (u, _, _, p, _) in enumerate(NUTZER):
    K[u] = Client(f'198.51.100.{10 + i}'); K[u].call('login', {'username': u, 'password': p})
L, Be, Q, V, O, N = K['lisa'], K['bernd'], K['quinn'], K['vera'], K['otto'], Client('198.51.100.30')
paket = json.load(open(PAKET_DATEI, encoding='utf-8'))
check('Datenpaket (erfunden) übernommen', V.call('qs_import', {'paket': paket, 'ersetzen': True})[0] == 200)


def zeitraum(von, bis):
    return f'&von={von}&bis={bis}'


VON, BIS = tag(-120), tag(0)
check('ohne Anmeldung: 401', N.call('qs_bericht' + zeitraum(VON, BIS))[0] == 401 and N.call('qs_bericht_excel' + zeitraum(VON, BIS))[0] == 401)
check('ohne QM-Recht: 403 (Bericht und Excel)', O.call('qs_bericht' + zeitraum(VON, BIS))[0] == 403 and O.call('qs_bericht_excel' + zeitraum(VON, BIS))[0] == 403)
for name, q, text in (('ohne Von', zeitraum('', BIS), 'Von: bitte ein Datum'), ('Bis in der Zukunft', zeitraum(VON, tag(1)), 'Zukunft'),
                      ('Von nach Bis', zeitraum(tag(-1), tag(-2)), 'Von muss vor Bis'), ('mehr als zwei Jahre', zeitraum(tag(-731), BIS), 'zwei Jahre'),
                      ('ungültiges Datum', zeitraum('2026-13-01', BIS), 'gültiges Datum')):
    st, r = L.call('qs_bericht' + q)
    check(f'Zeitraum {name}: 422 mit Hinweis', st == 422 and text in r.get('error', ''), (st, r))

# Prüfungen relativ zu heute, Nachweise in der Vergangenheit
P = {'kapitel': '2.1', 'beschreibung': 'Test', 'sollwert': '', 'massnahme': '', 'womit': '', 'zustaendig': 'QB', 'vertreter': 'QV'}
monat0 = (HEUTE.replace(day=1) - datetime.timedelta(days=100)).replace(day=1).isoformat()
P1 = Q.call('qs_pruefung_save', {**P, 'titel': 'Wareneingang Bericht', 'intervall': 'arbeitstaeglich', 'start': tag(-45)})[1]['pruefung']
P2 = Q.call('qs_pruefung_save', {**P, 'titel': 'Wöchentliche Begehung', 'intervall': 'woechentlich', 'start': tag(-70)})[1]['pruefung']
P3 = Q.call('qs_pruefung_save', {**P, 'titel': 'Monatliche Kontrolle', 'intervall': 'monatlich', 'start': monat0})[1]['pruefung']
P4 = Q.call('qs_pruefung_save', {**P, 'titel': 'Jährliche Bewertung', 'intervall': 'jaehrlich', 'start': tag(-20)})[1]['pruefung']
ok = True
for d in range(-45, 0, 3):
    ok = Be.call('qs_nachweis_save', {'pruefung': P1['id'], 'datum': tag(d), 'ergebnis': 'io', 'bemerkung': '', 'durch': ''})[0] == 200 and ok
for d in (-70, -56, -14):
    ok = Be.call('qs_nachweis_save', {'pruefung': P2['id'], 'datum': tag(d), 'ergebnis': 'io', 'bemerkung': '', 'durch': ''})[0] == 200 and ok
for d in (3, 65):
    ok = Be.call('qs_nachweis_save', {'pruefung': P3['id'], 'datum': (datetime.date.fromisoformat(monat0) + datetime.timedelta(days=d)).isoformat(), 'ergebnis': 'io', 'bemerkung': '', 'durch': ''})[0] == 200 and ok
st, r = Be.call('qs_nachweis_save', {'pruefung': P1['id'], 'datum': tag(-2), 'ergebnis': 'abweichung', 'bemerkung': 'Ware zu warm angeliefert', 'durch': '',
                                     'massnahme': {'korrektur': 'Ware gesperrt', 'massnahme': 'Spedition anschreiben', 'frist': tag(10)}})
ok = st == 200 and ok
NR_ABW = r.get('massnahme', {}).get('nr')
st, r = Be.call('qs_nachweis_save', {'pruefung': P2['id'], 'datum': tag(-30), 'ergebnis': 'io', 'bemerkung': '', 'durch': ''})
ok = ok and st == 200 and Be.call('qs_nachweis_storno', {'id': r['nachweis']['id'], 'grund': 'falsche Prüfung'})[0] == 200
ok = ok and Be.call('qs_massnahme_save', {'quelle': 'intern', 'feststellung': '=HYPERLINK("http://x.invalid";"klick")', 'massnahme': '+1', 'status': 'offen'})[0] == 200
check('Testdaten: Prüfungen, Nachweise, Abweichung mit Maßnahme, Storno, Maßnahme mit Formeltext', ok and P1 and P2 and P3 and P4 and NR_ABW)

st, s = L.call('qs_start')
st, r = L.call('qs_bericht' + zeitraum(VON, BIS))
check('Bericht mit „ansehen“: 200, Zeitraum, Stand, Ersteller', st == 200 and r['von'] == VON and r['bis'] == BIS and r['heute'] == BIS and r['erstellt']['von'] == 'Lisa Lesen'
      and r['monate'][0] == VON[:7] and r['monate'][-1] == BIS[:7], {k_: r.get(k_) for k_ in ('von', 'bis', 'heute', 'erstellt', 'monate')})
rp = {x['id']: x for x in r.get('pruefungen', [])}
sp = {x['id']: x for x in s.get('pruefungen', [])}
gleich = [(rp[p['id']]['versaeumt'], sp[p['id']]['faellig']['versaeumtAnzahl']) for p in (P1, P2, P3)]
check('gleiche Versäumnisse wie die Fälligkeit im Prüfplan (täglich, wöchentlich, monatlich)', all(a == b_ for a, b_ in gleich) and gleich[0][0] > 0, gleich)
check('erfüllt + versäumt = fällig, Summe der Monate = fällig', all(rp[p['id']]['erfuellt'] + rp[p['id']]['versaeumt'] == rp[p['id']]['faellig']
                                                                      == sum(m.get('f', 0) for m in rp[p['id']]['monate'].values()) for p in (P1, P2, P3)), [rp[p['id']] for p in (P1, P2)])
check('jährlich seit 20 Tagen fällig: überfällig, 20 Tage', rp[P4['id']]['fristen'] == [{'faellig': tag(-20), 'erledigt': '', 'status': 'ueberfaellig', 'tage': 20}], rp[P4['id']]['fristen'])
check('Abweichung mit der angelegten Maßnahme, storniert gezählt', [(a['datum'], a['massnahme']['nr']) for a in r['abweichungen'] if a['pruefung'] == 'Wareneingang Bericht'] == [(tag(-2), NR_ABW)]
      and r['kennzahlen']['pruefplan']['storniert'] == 1, r['abweichungen'])
gelenkt = [d for d in s['dokumente'] if d['art'] in s['gelenkt'] and not d.get('eingezogen')]
erwartet_frei = sorted((d['nr'], str(v['nr'])) for d in s['dokumente'] for v in d.get('versionen', []) if v['status'] in ('freigegeben', 'ersetzt') and VON <= (v.get('datum') or '') <= BIS)
check('Dokumente wie im QM-System: gültig, Überprüfung fällig, Freigaben im Zeitraum',
      r['kennzahlen']['dokumente']['gueltig'] == sum(1 for d in gelenkt if any(v['status'] == 'freigegeben' for v in d['versionen']))
      and r['kennzahlen']['dokumente']['faellig'] == sum(1 for d in s['dokumente'] if (d.get('ueberpruefung') or {}).get('faellig'))
      and sorted((x['nr'], x['version']) for x in r['dokumente']['vorgaenge'] if x['vorgang'] == 'freigabe') == erwartet_frei, r['kennzahlen']['dokumente'])
check('Handbuch-Änderungen wie im QM-System', r['kennzahlen']['handbuch'] == sum(1 for a in s['handbuch'] for c in a.get('aenderungen', []) if VON <= (c.get('datum') or '') <= BIS), r['kennzahlen']['handbuch'])
check('Maßnahmen: Datenpaket, Abweichung und neue Maßnahme (heute erfasst)', {NR_ABW, '2026-001'} <= {m['nr'] for m in r['massnahmen']}
      and any(m['feststellung'].startswith('=HYPERLINK') for m in r['massnahmen']), [m['nr'] for m in r['massnahmen']])

req = urllib.request.Request(API + 'qs_bericht_excel' + zeitraum(VON, BIS), headers={'X-Forwarded-For': L.ip})
ant = L.http.open(req)
daten = ant.read()
check('Excel: Typ, Dateiname mit Zeitraum, nicht zwischengespeichert, nosniff', 'spreadsheetml.sheet' in ant.headers.get('Content-Type', '')
      and f'QM-Audit-Bericht_{VON}_{BIS}.xlsx' in ant.headers.get('Content-Disposition', '') and 'no-store' in ant.headers.get('Cache-Control', '')
      and ant.headers.get('X-Content-Type-Options') == 'nosniff' and int(ant.headers.get('Content-Length', 0)) == len(daten), dict(ant.headers))
wb = openpyxl.load_workbook(io.BytesIO(daten))
titel = [row[1] for row in wb['Prüfplan'].iter_rows(min_row=2, values_only=True) if row and row[1]]
check('Excel: Prüfplan mit allen Prüfungen des Berichts', len(titel) == len(r['pruefungen']) and 'Wareneingang Bericht' in titel, titel)
mx = zipfile.ZipFile(io.BytesIO(daten)).read('xl/worksheets/sheet5.xml').decode()
check('Excel: Formeltext der Maßnahme bleibt Text', '=HYPERLINK' in mx.replace('&quot;', '"') and '<f>' not in mx, re.findall(r'.{40}HYPERLINK.{40}', mx)[:1])
log = json.dumps(A.call('log')[1], ensure_ascii=False)
check('Protokoll: Excel-Download mit Zeitraum vermerkt (Ansicht nicht)', 'Audit-Bericht' in log and 'als Excel heruntergeladen' in log
      and datetime.date.fromisoformat(VON).strftime('%d.%m.%Y') in log, log[-600:])

# ── Oberfläche ──
_j, _m = (HEUTE.year, HEUTE.month - 11) if HEUTE.month == 12 else (HEUTE.year - 1, HEUTE.month + 1)
V12 = datetime.date(_j, _m, 1)   # letzte 12 Monate einschließlich des laufenden
st, r12 = L.call('qs_bericht' + zeitraum(V12.isoformat(), BIS))
with sync_playwright() as pw:
    br = pw.chromium.launch()

    def anmelden(user, pw_, ip, breite=1280, thema='light', hoehe=900):
        ctx = br.new_context(viewport={'width': breite, 'height': hoehe}, extra_http_headers={'X-Forwarded-For': ip}, locale='de-DE', timezone_id='Europe/Berlin',
                             color_scheme=thema, accept_downloads=True)
        pg = ctx.new_page()
        pg.on('pageerror', lambda e: fehler.append(str(e)))
        pg.on('console', lambda m: fehler.append(m.text) if m.type == 'error' and 'Failed to load resource' not in m.text else None)
        pg.goto(B + '/admin/'); pg.wait_for_selector('#screen-auth:not([hidden])')
        pg.fill('#au-user', user); pg.fill('#pw1', pw_); pg.click('#auth-submit')
        pg.wait_for_selector('#screen-auth', state='hidden')
        pg.goto(B + '/module/qs/#bericht'); pg.wait_for_selector('#qm-app:not([hidden])')
        return ctx, pg

    def text(pg, sel):
        return pg.text_content(sel) or ''

    ctx, pg = anmelden('lisa', 'Leserin-Konto-12', '198.51.100.10')
    pg.wait_for_selector('#qm-b-bericht .qm-b-kopf')
    check('Reiter „Audit“: Bericht der letzten 12 Monate entsteht beim Öffnen', pg.is_visible('#qm-tabs [data-tab="bericht"][aria-selected="true"]')
          and pg.input_value('#qm-b-von') == r12['von'] and pg.input_value('#qm-b-bis') == BIS and V12.strftime('%d.%m.%Y') in text(pg, '.qm-b-kopf h2'),
          (pg.input_value('#qm-b-von'), r12['von'], text(pg, '.qm-b-kopf h2')))
    q = r12['kennzahlen']['pruefplan']['quote']
    check('Kennzahl „Prüfplan erfüllt“ wie vom Server, Matrix mit allen Prüfungen und Monaten',
          text(pg, '.qm-b-kpis .qm-kpi:first-child b') == (f"{math.floor(q * 1000 + 0.5) / 10:g} %".replace('.', ',') if q is not None else '—')
          and pg.locator('.qm-b-matrix tbody tr').count() == len(r12['pruefungen']) and pg.locator('.qm-b-matrix thead th.qm-b-mon').count() == len(r12['monate']),
          (text(pg, '.qm-b-kpis .qm-kpi:first-child b'), q))
    check('Versäumnisse, Abweichungen, Maßnahmen, Dokumente und Handbuch dargestellt', 'Wareneingang Bericht' in text(pg, '#qm-b-bericht')
          and 'Ware zu warm angeliefert' in text(pg, '#qm-b-bericht') and '=HYPERLINK' in text(pg, '#qm-b-bericht') and 'Überprüfung fällig' in text(pg, '#qm-b-bericht')
          and pg.locator('#qm-b-bericht .qm-b-abschnitt').count() == 6, text(pg, '#qm-b-bericht')[:300])
    zeilen = pg.evaluate('''() => { const s = [...document.querySelectorAll('.qm-b-abschnitt')].find((x) => x.querySelector('h3').textContent === 'Versäumnisse und Fristen');
                                     return [...s.querySelectorAll('tbody tr')].filter((r) => r.textContent.includes('Wareneingang Bericht')).map((r) => [r.cells[2].textContent, r.cells[3].textContent]); }''')
    vb = next(x for x in r12['pruefungen'] if x['titel'] == 'Wareneingang Bericht')
    check('Versäumnisse je Prüfung in einer Zeile mit allen Bereichen und der Anzahl', len(zeilen) == 1 and zeilen[0][0].count(' · ') == min(len(vb['bereiche']), 12) - 1
          and zeilen[0][1] == str(vb['versaeumt']), (zeilen, len(vb['bereiche'])))
    pg.screenshot(path=os.path.join(SH, 'qs_bericht_1280.png'), full_page=True)
    pg.select_option('#qm-b-zeitraum', 'vorjahr')
    pg.wait_for_function('y => document.querySelector(".qm-b-kopf h2").textContent.includes("01.01." + y)', arg=str(HEUTE.year - 1))
    check('Vorjahr: neuer Bericht 01.01.–31.12.', pg.input_value('#qm-b-von') == f'{HEUTE.year - 1}-01-01' and pg.input_value('#qm-b-bis') == f'{HEUTE.year - 1}-12-31')
    pg.fill('#qm-b-bis', tag(1)); pg.click('#qm-b-los')
    pg.wait_for_function('() => document.querySelector("#qm-b-meldung").textContent.length > 0')
    check('Zeitraum in der Zukunft: Hinweis, Felder markiert, Auswahl „Frei wählen“', 'Zukunft' in text(pg, '#qm-b-meldung') and pg.get_attribute('#qm-b-bis', 'aria-invalid') == 'true'
          and pg.input_value('#qm-b-zeitraum') == 'frei', text(pg, '#qm-b-meldung'))
    pg.select_option('#qm-b-zeitraum', '12')
    pg.wait_for_function('() => !document.querySelector("#qm-b-meldung").textContent && !document.querySelector("#qm-b-bis").hasAttribute("aria-invalid")')
    with pg.expect_download() as dl:
        pg.click('#qm-b-excel')
    check('Excel-Download aus der Oberfläche', dl.value.suggested_filename == f'QM-Audit-Bericht_{r12["von"]}_{BIS}.xlsx', dl.value.suggested_filename)
    pg.emulate_media(media='print')
    check('Druckansicht: nur der Bericht (ohne Kopf, Reiter, Zeitraumwahl)', pg.is_visible('#qm-b-bericht .qm-b-matrix') and not pg.is_visible('.qm-kopf')
          and not pg.is_visible('#qm-b-form') and not pg.is_visible('#qm-tabs'))
    pdf = pg.pdf(path=os.path.join(SH, 'qs_bericht.pdf'), prefer_css_page_size=True, print_background=True)
    check('PDF (A4 quer) entsteht', len(pdf) > 10000 and pdf.count(b'/Type /Page') - pdf.count(b'/Type /Pages') >= 1 and b'/MediaBox [0 0 841' in pdf, len(pdf))
    pg.emulate_media(media='screen')
    ctx.close()
    for breite, thema in ((390, 'light'), (390, 'dark'), (1280, 'dark')):
        ctx, pg = anmelden('lisa', 'Leserin-Konto-12', '198.51.100.10', breite, thema, 844 if breite < 500 else 900)
        pg.wait_for_selector('#qm-b-bericht .qm-b-kopf')
        breit = pg.evaluate('() => document.documentElement.scrollWidth - window.innerWidth')
        check(f'{breite} px {thema}: Bericht ohne waagerechtes Scrollen der Seite', breit <= 0, breit)
        pg.screenshot(path=os.path.join(SH, f'qs_bericht_{breite}_{thema}.png'), full_page=True)
        ctx.close()
    br.close()
check('keine JavaScript-Fehler', not fehler, fehler[:5])

print(f'{sum(res)}/{len(res)} passed')
