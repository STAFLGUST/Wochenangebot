"""Bestand: Marke je Partie (Spalte E der WinAB-Liste, Überschrift „Marke“ eine Spalte rechts davon), dezente MHD-Farben
(hell/dunkel) und PDF ohne Preise, erfassten Bestand, Gesamt verfügbar und erwartete Zugänge. Die Bestandsliste wird hier
synthetisch erzeugt (erfundene Artikel und Marken, Aufbau wie das WinAB-Original). Nur synthetische Testdaten."""
import datetime
import json
import os
import re
import sys
import urllib.error
import urllib.request
import zlib

import openpyxl
from playwright.sync_api import sync_playwright

from kunden_hilfe import Client
from umgebung import BASE, WORK

PW = 'Sicheres-Passwort-1'
res = []


def check(n, c, i=''):
    res.append(bool(c))
    print(('PASS ' if c else 'FAIL ') + n + (('  ' + str(i)) if i else ''))


heute = datetime.date.today()
tag = lambda d: int((heute + datetime.timedelta(days=d)).strftime('%d%m%Y'))   # noqa: E731 – TTMMJJJJ wie WinAB
kurz = lambda d: (heute + datetime.timedelta(days=d)).strftime('%d.%m.%y')     # noqa: E731

# ---------- synthetische WinAB-Bestandsliste ----------
# Spalten wie im Original: B Art.-Nr., C Bezeichnung, E Marke (Überschrift in F!), K Status (Überschrift in L), N Ankunft,
# P Partie, R MHD, T Kolli, W netto, X EKPreis
ZEILEN = [   # code, Bezeichnung, Marke, Status, Ankunft (Tage), Partie, MHD (Tage), Kolli, kg, EK
    (10913, 'Testfilet, Färse, frisch', 'Hofgut Nord', 'eingelagert', -40, 700001, 0, 1, 2.3, 22.4),
    (10913, 'Testfilet, Färse, frisch', 'Weidegut Süd', 'eingelagert', -30, 700002, 10, 3, 29.5, 21.9),
    (10913, 'Testfilet, Färse, frisch', 'Hofgut Nord', 'eingelagert', -10, 700003, 40, 2, 25.12, 20.4),
    (10601, 'Testroastbeef, 4xD, frisch', 'Zukauf Mitte', 'eingelagert', -20, 700004, 6, 4, 48.7, 23.3),
    (10601, 'Testroastbeef, 4xD, frisch', 'Lieferant Ost', 'erfasst', 1, 900001, None, 6, 142.85, 22.0),
    (10000, 'Testoberschale, 4 x D, frisch', 'Hofgut Nord', 'eingelagert', -5, 700005, 40, 10, 541.0, 11.2),
    (10002, 'Testoberschale, Färse, 4xD', 'Weidegut Süd', 'eingelagert', -60, 700006, -1, 2, 63.62, 11.8),
    (10006, 'Testoberschale JB, EG, frisch', '', 'eingelagert', -3, 700007, 20, 5, 379.87, 12.0),
    (10007, 'Testoberschale, Färse, EG, frisch', 'Weidegut Süd', 'eingelagert', -8, 700009, 10, 3, 34.4, 12.1),
    (99600, 'Test Knochen / Verschnitt', 'Eigene Zerl.', 'eingelagert', -2, 700008, None, 30, 1858.16, 0.5),
]
wb = openpyxl.Workbook(); ws = wb.active
ws['A1'] = 'Musterfirma Test'; ws['M1'] = 'Partieliste Verkäufer verfügbar'; ws['Z1'] = 'Datum:' + heute.strftime('%d.%m.%Y')
ws['B2'] = 'Art.-Nr.'; ws['D2'] = 'Bezeichnung'
for col, txt in {'F': 'Marke', 'G': 'EK Gruppe', 'I': 'VK Gruppe', 'L': 'Status', 'N': 'Ankunft', 'Q': 'Partie', 'S': 'MHD', 'U': 'Kolli', 'V': 'netto', 'X': 'EKPreis'}.items():
    ws[f'{col}3'] = txt
for r, (code, name, marke, st, ank, lot, mhd, kolli, kg, ek) in enumerate(ZEILEN, start=5):
    ws[f'B{r}'] = code; ws[f'C{r}'] = name; ws[f'H{r}'] = 120; ws[f'J{r}'] = 1350; ws[f'K{r}'] = st
    if marke: ws[f'E{r}'] = marke
    ws[f'N{r}'] = tag(ank); ws[f'P{r}'] = lot; ws[f'T{r}'] = kolli; ws[f'W{r}'] = kg; ws[f'X{r}'] = ek
    if mhd is not None: ws[f'R{r}'] = tag(mhd)
DATEI = os.path.join(WORK, 'Bestand_Marke_Test.xlsx')
wb.save(DATEI)


def pdf_text(raw):
    """Text aus den (komprimierten) Seiteninhalten des PDF-Erzeugers: (…) Tj, WinAnsi"""
    out = []; farben = []
    for m in re.finditer(rb'stream\n(.*?)\nendstream', raw, re.S):
        try:
            s = zlib.decompress(m.group(1))
        except zlib.error:
            continue
        for t in re.finditer(rb'\(((?:\\.|[^\\)])*)\) Tj', s):
            out.append(re.sub(rb'\\(.)', rb'\1', t.group(1)).decode('cp1252'))
        farben += re.findall(rb'([\d.]+ [\d.]+ [\d.]+) rg', s)
    return out, {f.decode() for f in farben}


def pdf_roh(c, data):
    """stock_pdf als Bytes (Client.call liefert Text)"""
    req = urllib.request.Request(c_api + 'stock_pdf', data=json.dumps(data).encode(), headers={'X-Forwarded-For': c.ip, 'Content-Type': 'application/json', 'X-CSRF-Token': c.csrf})
    try:
        r = c.http.open(req)
    except urllib.error.HTTPError as e:
        r = e
    return r.code, r.read()


c_api = BASE + '/api/?action='
A = Client('198.51.100.2')
assert A.call('login', {'username': 'admin', 'password': PW})[0] == 200
A.call('user_save', {'username': 'lea', 'name': 'Lea Leser', 'role': 'viewer', 'active': True, 'password': 'Leserin-Zugang-1', 'mustChange': False})
A.call('user_save', {'username': 'otto', 'name': 'Otto Ohne', 'role': 'editor', 'active': True, 'password': 'Editor-Zugang-77', 'mustChange': False, 'deny': ['bestand.sehen']})

errors = []
with sync_playwright() as p:
    br = p.chromium.launch()
    ctx = br.new_context(viewport={'width': 1440, 'height': 950}, accept_downloads=True)
    a = ctx.new_page(); a.on('pageerror', lambda e: errors.append(str(e)))
    a.goto(BASE + '/admin/'); a.wait_for_selector('#screen-auth:not([hidden])')
    a.fill('#au-user', 'admin'); a.fill('#pw1', PW); a.click('#auth-submit'); a.wait_for_selector('#screen-editor:not([hidden])')
    a.click('#vt-stock'); a.wait_for_selector('#view-stock [data-st=import]')
    with a.expect_file_chooser() as fc: a.click('#view-stock [data-st=import]')
    fc.value.set_files(DATEI)
    a.wait_for_selector('#dlg-stock-import [data-si=apply]', timeout=20000)
    a.click('#dlg-stock-import [data-si=apply]'); a.wait_for_selector('.stab tr.srow', timeout=10000)

    # ---------- Marke ----------
    st = a.evaluate("fetch('../api/?action=stock').then(r => r.json()).then(j => j.stock)")
    f913 = next(x for x in st['articles'] if x['code'] == '10913')
    check('Import: Marke je Partie aus Spalte E (Überschrift steht in F)', [l['marke'] for l in f913['lots']] == ['Hofgut Nord', 'Weidegut Süd', 'Hofgut Nord'], f913['lots'])
    check('… Bezeichnung bleibt richtig (nicht die Marke)', f913['name'] == 'Testfilet, Färse, frisch', f913['name'])
    check('… leere Marke bleibt leer', next(x for x in st['articles'] if x['code'] == '10006')['lots'][0]['marke'] == '')
    check('… auch bei erwarteten Zugängen', [z['marke'] for z in st['incoming']] == ['Lieferant Ost'], st['incoming'])
    a.locator('.stab tr.srow:has(td.s-code:text-is("10913"))').click(); a.wait_for_selector('#dlg-stock-item[open]')
    kopf = a.locator('#dlg-stock-item .lottab thead th').all_inner_texts()
    zeile = a.locator('#dlg-stock-item .lottab tbody tr').first.locator('td').all_inner_texts()
    check('Pop-up: Spalte „Marke“ vor der Partie, Wert in der Partiezeile', kopf[:2] == ['Marke', 'Partie'] and zeile[:2] == ['Hofgut Nord', '700001'], (kopf, zeile))
    check('… Summenzeile passt (eine Zelle mehr, „Summe“ unter der Partie)', a.locator('#dlg-stock-item .lottab tfoot td').count() == len(kopf)
          and a.locator('#dlg-stock-item .lottab tfoot td').all_inner_texts()[1] == 'Summe', a.locator('#dlg-stock-item .lottab tfoot td').all_inner_texts())
    a.wait_for_function("document.getAnimations().every((x) => x.playState !== 'running')")
    a.locator('#dlg-stock-item').screenshot(path=os.path.join(WORK, 'bestand-marke-1440.png'))
    a.set_viewport_size({'width': 390, 'height': 844})
    a.wait_for_function("getComputedStyle(document.querySelector('#dlg-stock-item .lottab th.lot-ank')).display === 'none'")
    passt = a.evaluate("(() => { const w = document.querySelector('#dlg-stock-item .lot-scroll'); return [w.scrollWidth <= w.clientWidth + 1, document.documentElement.scrollWidth <= 390]; })()")
    a.screenshot(path=os.path.join(WORK, 'bestand-marke-390.png'))
    sub = a.locator('#dlg-stock-item .lottab tbody tr').first.locator('.lot-marke-sub')
    check('Handy: Partien mit Marke passen in den Dialog (Marke klein unter der Partie)', all(passt) and sub.is_visible() and sub.inner_text() == 'Hofgut Nord'
          and a.locator('#dlg-stock-item .lottab th.lot-marke').is_hidden(), passt)
    a.set_viewport_size({'width': 1440, 'height': 950}); a.keyboard.press('Escape')
    a.locator('.stab tr.srow:has(td.s-code:text-is("10601"))').click(); a.wait_for_selector('#dlg-stock-item[open]')
    zt = a.locator('#dlg-stock-item .zutab tbody tr').first.locator('td').all_inner_texts()
    check('Pop-up: erwarteter Zugang nennt Marke und Partie (Spalten wie bei den Partien)', zt[:2] == ['Lieferant Ost', '900001'], zt)
    a.keyboard.press('Escape')

    # ---------- dezente MHD-Farben ----------
    farbe = lambda sel: a.evaluate(f"(() => {{ const e = document.querySelector('{sel}'); const s = getComputedStyle(e); return [s.backgroundColor, s.color]; }})()")   # noqa: E731
    a.evaluate("document.documentElement.dataset.theme = 'light'")
    hell = {k: farbe(f'.stab td.s-mhd.mhd-{k}') for k in ['0-4', '5-7', '8-14', 'past']}
    check('hell: getönte Hintergründe statt Signalfarben', hell == {'0-4': ['rgb(251, 224, 221)', 'rgb(143, 29, 20)'], '5-7': ['rgb(253, 231, 211)', 'rgb(122, 50, 9)'],
                                                        '8-14': ['rgb(251, 240, 201)', 'rgb(107, 74, 5)'], 'past': ['rgb(217, 214, 210)', 'rgb(31, 29, 26)']}, hell)
    a.screenshot(path=os.path.join(WORK, 'bestand-mhd-hell.png'))
    a.evaluate("document.documentElement.dataset.theme = 'dark'")
    dunkel = farbe('.stab td.s-mhd.mhd-0-4')
    check('dunkel: zurückhaltender Ton mit heller Schrift (kein grelles Rot)', dunkel[0].startswith('rgba(229, 83, 75, 0.2') and dunkel[1] == 'rgb(246, 184, 178)', dunkel)
    legende = a.evaluate("[...document.querySelectorAll('.mhd-legend > span')].map((e) => getComputedStyle(e).backgroundColor)")
    check('Legende in denselben Tönen', len(legende) == 4 and legende[1].startswith('rgba(229, 83, 75'), legende)
    a.screenshot(path=os.path.join(WORK, 'bestand-mhd-dunkel.png'))
    a.evaluate("document.documentElement.dataset.theme = 'light'")

    # ---------- PDF ----------
    check('Knopf „PDF“ nach „Drucken“', a.locator('.st-actions [data-st=print] ~ [data-st=pdf]').count() == 1)
    with a.expect_download() as dl: a.click('[data-st=pdf]')
    raw = open(dl.value.path(), 'rb').read()
    text, farben = pdf_text(raw)
    voll = '\n'.join(text)
    check('PDF erstellt (Dateiname mit Bestandsdatum)', raw.startswith(b'%PDF-1.4') and dl.value.suggested_filename == f'Bestand_{heute.isoformat()}.pdf', dl.value.suggested_filename)
    check('PDF: Spalten Art.-Nr., Bezeichnung, Eingelagerter Bestand, Kürz. MHD, Kg kürz. MHD', all(x in text for x in ['Art.-Nr.', 'Art. Bezeichnung', 'Eingelagerter Bestand', 'Kürz. MHD', 'Kg kürz. MHD']), text[:20])
    check('PDF: keine Preise, kein erfasster Bestand, kein Gesamt verfügbar, keine Zugänge', '€' not in voll and 'Aktueller Preis' not in voll and 'Vorwoche' not in voll
          and 'Erfasster Bestand' not in voll and 'Gesamt verfügbar' not in voll and 'Erwartete Zugänge' not in voll and '900001' not in voll and '142,85' not in voll, voll[:500])
    check('PDF: Artikel, Mengen und MHD aus dem Bestand', 'Testfilet, Färse, frisch' in text and '56,92 kg' in text and kurz(0) in text and '2,30 kg' in text, [t for t in text if 'Testfilet' in t or 'kg' in t][:8])
    check('PDF: Warengruppen mit Summen und Gesamtsumme', 'Rindfleisch, frisch' in text and 'Summe Rindfleisch, frisch' in text and 'Gesamtsumme' in text
          and '1.124,51 kg' in text and '2.982,67 kg' in text, [t for t in text if 'umme' in t or 'kg' in t][-8:])
    check('PDF: MHD in den hellen Tönen markiert (heute – 4 Tage, abgelaufen)', '0.984 0.878 0.867' in farben and '0.851 0.839 0.824' in farben, sorted(farben))
    check('PDF: Seitenzahl', 'Seite 1 von 1' in text, [t for t in text if 'Seite' in t])
    # Auswahl wie angezeigt
    a.fill('#st-q', 'Testfilet'); a.wait_for_function("document.querySelectorAll('.stab tr.srow').length === 1")
    with a.expect_download() as dl: a.click('[data-st=pdf]')
    t2, _ = pdf_text(open(dl.value.path(), 'rb').read())
    check('PDF der Suche: nur die angezeigten Artikel, Suche im Titel', 'Testfilet, Färse, frisch' in t2 and 'Testroastbeef, 4xD, frisch' not in t2 and any('Suche „Testfilet“' in t for t in t2)
          and dl.value.suggested_filename.endswith('_Auswahl.pdf'), (dl.value.suggested_filename, t2[:6]))
    a.fill('#st-q', 'MHD'); a.wait_for_selector('.stab.flat')
    with a.expect_download() as dl: a.click('[data-st=pdf]')
    t3, _ = pdf_text(open(dl.value.path(), 'rb').read())
    codes = [t for t in t3 if re.fullmatch(r'\d{4,5}', t)]
    check('PDF der MHD-Auswahl: Reihenfolge wie angezeigt (früheste zuerst), ohne Gruppen', codes == ['10002', '10913', '10601', '10007'] and 'Rindfleisch, frisch' not in t3
          and any('MHD bis' in t for t in t3), (codes, [t for t in t3 if 'MHD' in t]))
    a.fill('#st-q', '')
    ctx.close()
    br.close()

# ---------- Schnittstelle: Rechte und Prüfungen ----------
L = Client('198.51.100.3'); assert L.call('login', {'username': 'lea', 'password': 'Leserin-Zugang-1'})[0] == 200
st, body = pdf_roh(L, {'rows': [{'g': 'Rindfleisch, frisch'}, {'c': '10913'}]})
check('Leser (Bestand ansehen) darf die PDF erstellen', st == 200 and body.startswith(b'%PDF'), st)
O = Client('198.51.100.4'); assert O.call('login', {'username': 'otto', 'password': 'Editor-Zugang-77'})[0] == 200
check('ohne Recht „Bestand ansehen“: 403', O.call('stock_pdf', {'rows': []})[0] == 403)
check('ohne Anmeldung: 401', Client('198.51.100.5').call('stock_pdf', {'rows': []})[0] == 401)
check('ungültige Auswahl: 422', A.call('stock_pdf', {'rows': [{'c': '<script>'}]})[0] == 422 and A.call('stock_pdf', {'rows': 'x'})[0] == 422)
st, body = pdf_roh(A, {'rows': [{'g': 'Leere Gruppe'}, {'c': '424242'}, {'g': 'Rindfleisch, frisch'}, {'c': '10913'}]})
tx, _ = pdf_text(body)
check('unbekannte Artikel und leere Gruppen werden ausgelassen', st == 200 and 'Leere Gruppe' not in tx and 'Testfilet, Färse, frisch' in tx and '424242' not in tx, tx[:12])
# Bezeichnung aus den Artikelstammdaten (Bestandsliste bleibt gespeichert wie importiert), Filter in der Überschrift
check('Stammdaten mit eigener Bezeichnung gespeichert', A.call('master_save', {'rev': A.call('master')[1]['master']['rev'], 'articles': [
      {'code': '10913', 'name': 'Filetköpfe aus Stammdaten', 'category': 'Rind', 'origin': 'AR', 'condition': 'frisch'}]})[0] == 200)
st, body = pdf_roh(A, {'rows': [{'g': 'Rindfleisch, frisch'}, {'c': '10913'}], 'q': 'Filet', 'filter': 'Gesamt verfügbar ab 10 kg'})
open(os.path.join(WORK, 'bestand-pdf-merkmale.pdf'), 'wb').write(body)
tx, _ = pdf_text(body)
check('PDF: Bezeichnung aus den Stammdaten, Filter in der Überschrift', st == 200 and 'Filetköpfe aus Stammdaten' in tx and 'Testfilet, Färse, frisch' not in tx
      and any('Filter: Gesamt verfügbar ab 10 kg' in x for x in tx), [x for x in tx if 'Filter' in x or 'Stammdaten' in x][:4])
check('PDF: Merkmale aus den Stammdaten grau hinter der Bezeichnung (Fleischart · Herkunft · Zustand)', 'Rind · AR · frisch' in tx, [x for x in tx if 'AR' in x][:4])
G = Client('198.51.100.6'); G.call('login', {'username': 'admin', 'password': PW}); G.csrf = ''
check('ohne CSRF-Schutzwert: abgewiesen', G.call('stock_pdf', {'rows': []})[0] == 403)
check('keine JavaScript-Fehler', not errors, errors)

print(f'{sum(res)}/{len(res)} passed')
sys.exit(0 if all(res) else 1)
