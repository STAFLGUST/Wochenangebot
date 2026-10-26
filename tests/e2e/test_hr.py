"""HR · Urlaubsverwaltung (Modul hr, Oberfläche hr.intern unter /module/hr/).

Geprüft: Freigabe (produktiv seit 23.10.2026: jeder Mitarbeiter mit hr.antrag, keine Testfreigabe), Mitarbeiter = Nutzer
(jeder Mitarbeiter ist sofort in der Urlaubsverwaltung; Mitarbeiter ohne Anmeldung; HR-Angaben in der Verwaltung;
Übernahme alter Einträge ohne Mitarbeiter), Rechte (hr.antrag für alle, hr.verwalten nur Personalstelle, hr.genehmigen als eigenes
zuweisbares Recht, nie der eigene Antrag und nie eigene Abwesenheiten – auch nicht als Administrator; serverseitig – GR-103), Antrag → Genehmigung/Ablehnung/Zurückziehen, Grenzfälle (Vergangenheit,
Wochenende, Überschneidung, Urlaubssperre), Urlaubskonto mit Übertrag bis 31.03. und Verfall, 4-Tage-Woche,
Datenschutz im Kalender (Kollegen sehen nur „abwesend“, nie Art, Notiz oder offene Anträge; Protokoll ohne Art),
Excel-Export für den Steuerberater (Aufbau der bisherigen Planung, nur genehmigte Einträge), Übernahme der bisherigen
Excel-Planung über die Importzentrale (synthetische Datei tests/fixtures/hr/Urlaub_Muster_2025.xlsx), Abwesenheiten aus der
Zeiterfassung (Importvorlage, tests/fixtures/hr/Abwesenheitsbericht_Muster.xlsx: Personalnummer, Schreibweisen, vom Urlaub
abgezogene Krankheit, Ersetzen, veralteter Stand) und die Oberflächen
(HR-Seite, Verwaltung → Mitarbeiter und Urlaub; Desktop, 390 px).
"""
import datetime as dt
import io
import json
import os
import sys
import zipfile

import openpyxl
from playwright.sync_api import sync_playwright

from kunden_hilfe import Client, B
from umgebung import SH, FIX, storage_raw, storage_write

PW = 'Sicheres-Passwort-1'
res = []


def check(n, c, i=''):
    res.append(bool(c)); print(('PASS ' if c else 'FAIL ') + n + (('  ' + str(i)[:400]) if i else ''))


def montag(jahr, monat, woche=1):
    """Montag der n-ten vollen Woche eines Monats (ohne Feiertage im Februar/August)."""
    d = dt.date(jahr, monat, 1)
    while d.weekday() != 0: d += dt.timedelta(days=1)
    return d + dt.timedelta(weeks=woche - 1)


def iso(d): return d.isoformat()


def dokument(name):
    """Gespeichertes Dokument lesen (Dateien: ohne Schutzzeile)."""
    raw = storage_raw(name)
    if not raw:
        return {}
    if raw.startswith('<?php'):
        raw = raw.split('\n', 1)[1]
    return json.loads(raw)


Y = dt.date.today().year + 1           # Anträge nur für die Zukunft: alles im nächsten Jahr
FEB = montag(Y, 2)                      # Montag der ersten vollen Februarwoche (vor dem 31.03.: verbraucht Übertrag)
AUG = montag(Y, 8, 2)

A = Client('198.51.100.2')
check('Administrator angemeldet', A.call('login', {'username': 'admin', 'password': PW})[0] == 200)
for u, n, r, p, extra in (('erika', 'Erika Editor', 'editor', 'Editor-Zugang-77', {}), ('lena', 'Lena Leser', 'viewer', 'Leser-Zugang-77', {}),
                          ('paul', 'Paul Personal', 'viewer', 'Personal-Zugang-77', {'grant': ['hr.verwalten', 'hr.genehmigen']}), ('otto', 'Otto Ohne', 'editor', 'Ohne-Zugang-77', {})):
    st, r_ = A.call('user_save', {'id': '', 'name': n, 'username': u, 'role': r, 'active': True, 'password': p, 'mustChange': False, **extra})
    assert st == 200, r_
st, r_ = A.call('user_save', {'id': '', 'vorname': 'Andre', 'nachname': 'Vierzig', 'role': 'editor', 'active': True, 'zugang': False})
check('Mitarbeiter ohne Anmeldung: ohne Passwort angelegt, deaktiviert, Rolle Leser', st == 200 and r_['user']['zugang'] is False and r_['user']['active'] is False
      and r_['user']['role'] == 'viewer' and r_['user']['username'] == 'andre-vierzig', r_)
ids = {u['username']: u['id'] for u in A.call('users')[1]['users']}
check('… kann sich nicht anmelden', Client('198.51.100.30').call('login', {'username': 'andre-vierzig', 'password': 'Urlaub-Konto-4711'})[0] in (401, 403))
st, r_ = A.call('user_save', {'id': ids['andre-vierzig'], 'vorname': 'Andre', 'nachname': 'Vierzig', 'role': 'viewer', 'active': True, 'zugang': True})
check('Anmeldung einrichten nur mit Passwort (422)', st == 422 and 'Passwort' in r_.get('error', ''), r_)
check('eigenen Zugang nicht auf „ohne Anmeldung“ (422)', A.call('user_save', {'id': ids['admin'], 'name': 'Administrator', 'role': 'admin', 'active': True, 'zugang': False})[0] == 422)
E = Client('198.51.100.3'); E.call('login', {'username': 'erika', 'password': 'Editor-Zugang-77'})
L = Client('198.51.100.4'); L.call('login', {'username': 'lena', 'password': 'Leser-Zugang-77'})
P = Client('198.51.100.5'); P.call('login', {'username': 'paul', 'password': 'Personal-Zugang-77'})
O = Client('198.51.100.6'); O.call('login', {'username': 'otto', 'password': 'Ohne-Zugang-77'})
N = Client('198.51.100.7')

# ── Freigabe ──
# Produktiv seit der Freigabe des Inhabers (23.10.2026): kein Testbetrieb, keine Tester – jeder Mitarbeiter mit hr.antrag
mod = {m['id']: m for m in A.call('module_status')[1]['module']}['hr']
check('Register: HR fertig und produktiv, Adresse hr.dashboard-stafl.de/module/hr/', (lambda o: o['entwicklung'] == 'fertig' and o['freigabe'] == 'produktiv'
      and o['zugang'] and o['adresse'] == 'https://hr.dashboard-stafl.de/module/hr/')(mod['oberflaechen'][0]), mod)
check('Kachel verfügbar', mod['kachel']['status'] == 'verfuegbar' and mod['kachel']['zugang'] is True, mod['kachel'])
st, r = A.call('modul_freigabe_save', {'bereich': 'hr.intern', 'freigabe': 'test', 'grund': 'Test'})
check('Produktivbereich: keine Testfreigabe über die Modulverwaltung (nur über ein Release)', st == 422, (st, r))
st, r = O.call('hr_start')
check('jeder Mitarbeiter mit hr.antrag: Start ok (otto, nie als Tester eingetragen)', st == 200 and r.get('ok'), (st, r))
check('Administrator und Personalstelle: Start ok', A.call('hr_start')[0] == 200 and P.call('hr_start')[0] == 200)
check('ohne Anmeldung: 401', N.call('hr_start')[0] == 401 and N.call('hr_verwaltung')[0] == 401 and N.call('mitarbeiter_hr')[0] == 401)

# ── Rechte ──
check('Mitarbeiter ohne Personalrecht: hr_verwaltung 403', L.call('hr_verwaltung')[0] == 403 and E.call('hr_verwaltung')[0] == 403 and L.call('mitarbeiter_hr')[0] == 403)
for a, b in (('hr_entscheiden', {'id': 'x', 'entscheidung': 'genehmigt'}), ('hr_eintrag_save', {'person': 'x', 'art': 'K', 'von': iso(FEB), 'bis': iso(FEB)}),
             ('mitarbeiter_hr_save', {'nutzer': 'x', 'rev': 0}), ('hr_export', {'jahr': Y, 'monate': [2]}), ('hr_einstellungen_save', {'sperren': []}),
             ('hr_personal_uebernehmen', {}), ('hr_import_pruefen', {'daten': {'blaetter': []}})):
    check(f'{a}: ohne Personalrecht 403', L.call(a, b)[0] == 403)
check('hr_offen: ohne „Urlaub genehmigen“ 403', L.call('hr_offen')[0] == 403 and E.call('hr_offen')[0] == 403 and N.call('hr_offen')[0] == 401)
check('hr_start: Rechte-Hinweise für die Oberfläche', E.call('hr_start')[1]['genehmigen'] is False and P.call('hr_start')[1]['genehmigen'] is True)
kat = {x['key']: x for x in A.call('rechte')[1]['catalog']}
check('Recht „Urlaub genehmigen“ im Katalog: zuweisbar, keine Standardrollen, setzt hr.antrag voraus', kat['hr.genehmigen']['default'] == [] and kat['hr.genehmigen']['requires'] == ['hr.antrag']
      and kat['hr.genehmigen']['group'] == 'HR' and 'hr.verwalten' in kat, kat.get('hr.genehmigen'))

# ── Mitarbeiter = Nutzer ──
st, r = E.call('hr_start')
check('jeder Mitarbeiter ist sofort in der Urlaubsverwaltung (ohne Zuordnung)', st == 200 and r['ich'] and r['ich']['id'] == ids['erika'] and r['konto']['anspruch'] == 0, r.get('ich'))
v = P.call('mitarbeiter_hr')[1]
pv = {p_['nutzer']: p_ for p_ in v['personen'] if p_['nutzer']}
check('Verwaltung: alle Mitarbeiter mit HR-Angaben (Standard: Mo–Fr, im Urlaubsplan)', set(ids.values()) <= set(pv) and pv[ids['otto']]['arbeitstage'] == [1, 2, 3, 4, 5]
      and pv[ids['otto']]['aktiv'] is True and pv[ids['andre-vierzig']]['zugang'] is False, sorted(pv))
rev = v['rev']


def hr(uid, **werte):
    global rev
    st, r = P.call('mitarbeiter_hr_save', {'nutzer': uid, 'rev': rev, **werte})
    assert st == 200, r
    rev = r['rev']
    return r['person']['id']


PE = hr(ids['erika'], abteilung='Verwaltung', pos=1, jahr=Y, anspruch=30, uebertrag=5)
PL = hr(ids['lena'], abteilung='Verwaltung', pos=2, jahr=2025, anspruch=30, uebertrag=4)
PA = hr(ids['andre-vierzig'], abteilung='Lager', pos=3, arbeitstage=[1, 2, 3, 4], jahr=Y, anspruch=24)
check('Person = Mitarbeiter-Kennung', (PE, PL, PA) == (ids['erika'], ids['lena'], ids['andre-vierzig']))
check('veralteter Stand: 409', P.call('mitarbeiter_hr_save', {'nutzer': ids['erika'], 'rev': rev - 1, 'pos': 1})[0] == 409)
check('ungültiger Anspruch (Viertel-Tage): 422', P.call('mitarbeiter_hr_save', {'nutzer': ids['erika'], 'rev': rev, 'jahr': Y, 'anspruch': 10.25})[0] == 422)
check('ohne Arbeitstage: 422', P.call('mitarbeiter_hr_save', {'nutzer': ids['erika'], 'rev': rev, 'arbeitstage': []})[0] == 422)
check('unbekannter Mitarbeiter: 404', P.call('mitarbeiter_hr_save', {'nutzer': 'u-gibt-es-nicht', 'rev': rev})[0] == 404)
hr(ids['otto'], aktiv=False)
check('nicht im Urlaubsplan: kein eigenes Konto, Antrag 409', O.call('hr_start')[1]['ich'] is None and O.call('hr_antrag', {'von': iso(AUG), 'bis': iso(AUG)})[0] == 409)
check('… und nicht im Kalender', ids['otto'] not in [p_['id'] for p_ in P.call(f'hr_monat&jahr={Y}&monat=2')[1]['personen']])
hr(ids['admin'], aktiv=False); hr(ids['paul'], aktiv=False)
check('Namen kommen aus dem Mitarbeiter (Verwaltung → Mitarbeiter)', {p_['id']: p_['name'] for p_ in P.call('hr_verwaltung')[1]['personen']}[PA] == 'Andre Vierzig')

# ── Konto und Antrag ──
st, r = E.call('hr_start&jahr=' + str(Y))
k = r['konto']
check('Konto: Anspruch 30 + Übertrag 5 = 35 Resturlaub', st == 200 and k['anspruch'] == 30 and k['uebertrag'] == 5 and k['rest'] == 35, k)
st, r = E.call('hr_antrag', {'von': iso(FEB), 'bis': iso(FEB + dt.timedelta(days=2)), 'notiz': 'Skiurlaub', 'vertretung': PL})
check('Antrag gestellt (3 Arbeitstage, beantragt)', st == 200 and r['eintrag']['status'] == 'beantragt' and r['eintrag']['tage'] == 3 and r['konto']['beantragt'] == 3, r)
ANTRAG1 = r['eintrag']['id']
check('Überschneidung: 409', E.call('hr_antrag', {'von': iso(FEB + dt.timedelta(days=2)), 'bis': iso(FEB + dt.timedelta(days=4))})[0] == 409)
gestern = dt.date.today() - dt.timedelta(days=1)
check('Vergangenheit: 422', E.call('hr_antrag', {'von': iso(gestern), 'bis': iso(gestern)})[0] == 422)
samstag = FEB + dt.timedelta(days=5)
check('nur Wochenende: 422 mit Hinweis', (lambda s: s[0] == 422 and 'Arbeitstage' in s[1]['error'])(E.call('hr_antrag', {'von': iso(samstag), 'bis': iso(samstag + dt.timedelta(days=1))})))
check('Ende vor Beginn: 422', E.call('hr_antrag', {'von': iso(AUG), 'bis': iso(AUG - dt.timedelta(days=1))})[0] == 422)
check('ungültige Vertretung: 422', E.call('hr_antrag', {'von': iso(AUG), 'bis': iso(AUG), 'vertretung': PE})[0] == 422)

# Urlaubssperre
st, r = P.call('hr_einstellungen_save', {'sperren': [{'von': f'{Y}-12-14', 'bis': f'{Y}-12-31', 'text': 'Weihnachtsgeschäft'}], 'ferien': [{'von': iso(AUG), 'bis': iso(AUG + dt.timedelta(days=11)), 'text': 'Sommerferien'}]})
check('Urlaubssperre und Ferien gespeichert', st == 200 and len(r['sperren']) == 1, r)
st, r = E.call('hr_antrag', {'von': f'{Y}-12-15', 'bis': f'{Y}-12-16'})
check('Antrag in der Urlaubssperre: 422 mit Bezeichnung', st == 422 and 'Weihnachtsgeschäft' in r['error'], r)
check('Sperre: Ende vor Beginn abgelehnt', P.call('hr_einstellungen_save', {'sperren': [{'von': f'{Y}-12-14', 'bis': f'{Y}-12-01'}]})[0] == 422)

# Krank für Lena (Personalstelle) – Gesundheitsdatum
st, r = P.call('hr_eintrag_save', {'person': PL, 'art': 'K', 'von': iso(FEB), 'bis': iso(FEB + dt.timedelta(days=1)), 'notiz': 'Arztbesuch-Notiz'})
check('Krank eingetragen (genehmigt)', st == 200 and r['eintrag']['status'] == 'genehmigt' and r['eintrag']['tage'] == 2, r)
KRANK = r['eintrag']['id']
check('Überschneidung bei der Personalstelle: 409', P.call('hr_eintrag_save', {'person': PL, 'art': 'U', 'von': iso(FEB + dt.timedelta(days=1)), 'bis': iso(FEB + dt.timedelta(days=1))})[0] == 409)
check('unbekannte Art: 422', P.call('hr_eintrag_save', {'person': PL, 'art': 'Z', 'von': iso(AUG), 'bis': iso(AUG)})[0] == 422)

# ── Datenschutz im Kalender ──
st, r, raw = E.call(f'hr_monat&jahr={Y}&monat=2', raw=True)
zeilen = {p['id']: p['tage'] for p in r['personen']}
tag = str(FEB.day)
check('Kollege sieht bei Lena nur „abwesend“', st == 200 and zeilen[PL].get(tag) == 'A' and r['voll'] is False, zeilen[PL])
check('… nie Art oder Notiz', '"K"' not in raw and 'Arzt' not in raw and 'Krank' not in raw, raw[:300])
check('eigene Zeile: beantragt mit „U?“', zeilen[PE].get(tag) == 'U?', zeilen[PE])
st, r, raw = L.call(f'hr_monat&jahr={Y}&monat=2', raw=True)
zl = {p['id']: p['tage'] for p in r['personen']}
check('offener Antrag eines Kollegen ist unsichtbar', tag not in zl[PE] and 'Skiurlaub' not in raw, zl[PE])
check('eigene Krankheit sieht Lena selbst', zl[PL].get(tag) == 'K', zl[PL])
st, r = P.call(f'hr_monat&jahr={Y}&monat=2')
zp = {p['id']: p['tage'] for p in r['personen']}
check('Personalstelle sieht Art und offene Anträge', r['voll'] is True and zp[PL].get(tag) == 'K' and zp[PE].get(tag) == 'U?', (zp[PL], zp[PE]))

# ── Anträge entscheiden ──
st, r = P.call('hr_verwaltung&jahr=' + str(Y))
offen = {a['id']: a for a in r['offen']}
a = offen.get(ANTRAG1, {})
check('offener Antrag mit Konto, Vertretung und gleichzeitig Abwesenden', a.get('tage') == 3 and a.get('vertretungName') == 'Lena Leser'
      and any('Lena Leser' in x for x in a.get('kollegen', [])) and a.get('notiz') == 'Skiurlaub', a)
check('Ablehnen ohne Grund: 422', P.call('hr_entscheiden', {'id': ANTRAG1, 'entscheidung': 'abgelehnt'})[0] == 422)
check('Genehmigen', P.call('hr_entscheiden', {'id': ANTRAG1, 'entscheidung': 'genehmigt'})[0] == 200)
check('zweimal entscheiden: 409', P.call('hr_entscheiden', {'id': ANTRAG1, 'entscheidung': 'abgelehnt', 'grund': 'x'})[0] == 409)
k = E.call('hr_start&jahr=' + str(Y))[1]['konto']
check('Konto nach Genehmigung: geplant 3, Rest 32', k['geplant'] == 3 and k['rest'] == 32 and k['beantragt'] == 0, k)
check('Übertrag: 2 Tage verfallen nach dem 31.03., falls nicht genommen', k['verfaellt'] == 2 and k['frist'] == f'{Y}-03-31', k)

# Ablehnung mit Grund
st, r = E.call('hr_antrag', {'von': iso(AUG), 'bis': iso(AUG + dt.timedelta(days=4))})
A2 = r['eintrag']['id']
check('Ablehnen mit Grund', P.call('hr_entscheiden', {'id': A2, 'entscheidung': 'abgelehnt', 'grund': 'Inventur'})[0] == 200)
e = {x['id']: x for x in E.call('hr_start&jahr=' + str(Y))[1]['eintraege']}[A2]
check('Mitarbeiter sieht Ablehnung mit Grund', e['status'] == 'abgelehnt' and e['entscheidung']['grund'] == 'Inventur', e)

# Zurückziehen
st, r = E.call('hr_antrag', {'von': iso(AUG + dt.timedelta(days=7)), 'bis': iso(AUG + dt.timedelta(days=8))})
A3 = r['eintrag']['id']
check('fremden Antrag zurückziehen: 404', L.call('hr_antrag_zurueckziehen', {'id': A3})[0] == 404)
check('eigenen Antrag zurückziehen', E.call('hr_antrag_zurueckziehen', {'id': A3})[0] == 200)
check('zweimal zurückziehen: 409', E.call('hr_antrag_zurueckziehen', {'id': A3})[0] == 409)
check('genehmigten künftigen Urlaub zurückziehen', E.call('hr_antrag_zurueckziehen', {'id': ANTRAG1})[0] == 200)
k = E.call('hr_start&jahr=' + str(Y))[1]['konto']
check('… Konto wieder 35', k['rest'] == 35 and k['geplant'] == 0, k)
st, r = E.call('hr_antrag', {'von': iso(FEB), 'bis': iso(FEB + dt.timedelta(days=2))})
P.call('hr_entscheiden', {'id': r['eintrag']['id'], 'entscheidung': 'genehmigt'})
OFFEN = E.call('hr_antrag', {'von': iso(montag(Y, 2, 3)), 'bis': iso(montag(Y, 2, 3) + dt.timedelta(days=1))})[1]['eintrag']['id']

# ── Urlaub genehmigen: eigenes, zuweisbares Recht (Inhaber 25.10.2026); nie der eigene Antrag, nie eigene Abwesenheiten ──
OKT = montag(Y, 10, 2)
for u, n, extra in (('gina', 'Gina Genehmigt', ['hr.genehmigen']), ('vera', 'Vera Verwalten', ['hr.verwalten'])):
    assert A.call('user_save', {'id': '', 'name': n, 'username': u, 'role': 'viewer', 'active': True, 'password': 'Genehmigen-Zugang-77', 'mustChange': False, 'grant': extra})[0] == 200
ids = {u['username']: u['id'] for u in A.call('users')[1]['users']}
G = Client('198.51.100.8'); G.call('login', {'username': 'gina', 'password': 'Genehmigen-Zugang-77'})
V = Client('198.51.100.9'); V.call('login', {'username': 'vera', 'password': 'Genehmigen-Zugang-77'})
st, r = G.call('hr_offen')
og = {a['id']: a for a in r.get('offen', [])}
check('nur „Urlaub genehmigen“: offene Anträge mit Konto, ohne Personaldaten', st == 200 and OFFEN in og and 'konto' in og[OFFEN] and og[OFFEN]['eigen'] is False
      and 'personen' not in r and 'eintraege' not in r, (st, list(r)))
check('… nur Urlaubsanträge, keine Krankheiten', all(a['art'] == 'U' for a in r['offen']) and 'Arzt' not in json.dumps(r), [a['art'] for a in r['offen']])
check('… keine Personalstelle: hr_verwaltung, Einträge und Export 403', G.call('hr_verwaltung')[0] == 403
      and G.call('hr_eintrag_save', {'person': PL, 'art': 'K', 'von': iso(OKT), 'bis': iso(OKT)})[0] == 403 and G.call('hr_export', {'jahr': Y, 'monate': [2]})[0] == 403)
st, r = E.call('hr_antrag', {'von': iso(OKT), 'bis': iso(OKT + dt.timedelta(days=1))})
A_OKT = r['eintrag']['id']
check('fremden Antrag genehmigen mit „Urlaub genehmigen“', G.call('hr_entscheiden', {'id': A_OKT, 'entscheidung': 'genehmigt'})[0] == 200)
check('… Mitarbeiterin zieht den künftigen Urlaub wieder zurück (Konto unverändert)', E.call('hr_antrag_zurueckziehen', {'id': A_OKT})[0] == 200)
st, r = V.call('hr_verwaltung&jahr=' + str(Y))
check('Personalstelle ohne Genehmigungsrecht: sieht offene Anträge', st == 200 and OFFEN in [a['id'] for a in r['offen']], st)
check('… darf aber nicht entscheiden (403) und hr_offen nicht (403)', V.call('hr_entscheiden', {'id': OFFEN, 'entscheidung': 'genehmigt'})[0] == 403 and V.call('hr_offen')[0] == 403)
check('… Start: Zahl offener Anträge trotzdem', (lambda x: x['verwalten'] is True and x['genehmigen'] is False and x['offen'] >= 1)(V.call('hr_start')[1]))
# Vier-Augen-Prinzip
st, r = G.call('hr_antrag', {'von': iso(OKT + dt.timedelta(days=7)), 'bis': iso(OKT + dt.timedelta(days=8))})
A_GINA = r['eintrag']['id']
og = {a['id']: a for a in G.call('hr_offen')[1]['offen']}
check('eigener Antrag als „eigen“ markiert, zählt nicht am Reiter', og[A_GINA]['eigen'] is True and G.call('hr_start')[1]['offen'] == len(og) - 1, (og.get(A_GINA), G.call('hr_start')[1]['offen']))
st, r = G.call('hr_entscheiden', {'id': A_GINA, 'entscheidung': 'genehmigt'})
check('eigenen Antrag genehmigen: 403', st == 403 and 'eigenen Antrag' in r.get('error', ''), (st, r))
check('eigenen Antrag ablehnen: 403', G.call('hr_entscheiden', {'id': A_GINA, 'entscheidung': 'abgelehnt', 'grund': 'x'})[0] == 403)
check('… Antrag bleibt offen', {x['id']: x for x in G.call('hr_start&jahr=' + str(Y))[1]['eintraege']}[A_GINA]['status'] == 'beantragt')
check('Administrator darf ihn entscheiden (andere Person)', A.call('hr_entscheiden', {'id': A_GINA, 'entscheidung': 'abgelehnt', 'grund': 'Test'})[0] == 200)
hr(ids['admin'], aktiv=True)
st, r = A.call('hr_antrag', {'von': iso(OKT + dt.timedelta(days=14)), 'bis': iso(OKT + dt.timedelta(days=14))})
A_ADMIN = r['eintrag']['id']
check('Administrator: eigenen Antrag nicht genehmigen (403, alle Rechte reichen nicht)', A.call('hr_entscheiden', {'id': A_ADMIN, 'entscheidung': 'genehmigt'})[0] == 403)
check('… Paul genehmigt ihn', P.call('hr_entscheiden', {'id': A_ADMIN, 'entscheidung': 'genehmigt'})[0] == 200)
check('Administrator: eigenen Urlaub nicht direkt eintragen (403)', A.call('hr_eintrag_save', {'person': ids['admin'], 'art': 'U', 'von': iso(OKT + dt.timedelta(days=21)), 'bis': iso(OKT + dt.timedelta(days=21))})[0] == 403)
check('… und den genehmigten nicht ändern oder stornieren (403)', A.call('hr_eintrag_save', {'id': A_ADMIN, 'person': ids['admin'], 'art': 'U', 'von': iso(OKT + dt.timedelta(days=14)), 'bis': iso(OKT + dt.timedelta(days=16))})[0] == 403
      and A.call('hr_eintrag_save', {'id': A_ADMIN, 'stornieren': True})[0] == 403)
check('… Paul storniert ihn (andere Person)', P.call('hr_eintrag_save', {'id': A_ADMIN, 'stornieren': True})[0] == 200)
hr(ids['admin'], aktiv=False)
st, r = P.call('hr_eintrag_save', {'person': ids['paul'], 'art': 'U', 'von': iso(OKT), 'bis': iso(OKT)})
check('Personalstelle: sich selbst Urlaub eintragen 403 (auch nicht im Urlaubsplan)', st == 403 and 'andere Person' in r.get('error', ''), (st, r))
for art in ('K', 'G', 'X'):
    check(f'… auch {art} für sich selbst 403', P.call('hr_eintrag_save', {'person': ids['paul'], 'art': art, 'von': iso(OKT), 'bis': iso(OKT)})[0] == 403)
st, r = A.call('hr_eintrag_save', {'person': ids['paul'], 'art': 'K', 'von': iso(OKT), 'bis': iso(OKT)})
EP = r.get('eintrag', {}).get('id')
check('Eintrag einer anderen Person für Paul: Paul darf ihn nicht ändern oder umhängen (403)', st == 200
      and P.call('hr_eintrag_save', {'id': EP, 'person': PL, 'art': 'K', 'von': iso(OKT), 'bis': iso(OKT)})[0] == 403
      and P.call('hr_eintrag_save', {'id': EP, 'stornieren': True})[0] == 403, (st, r))
r = P.call('hr_verwaltung&jahr=' + str(Y))[1]
check('… in „Alle Einträge“ als eigen markiert', {e['id']: e for e in r['eintraege']}.get(EP, {}).get('eigen') is True)
check('… Administrator storniert ihn', A.call('hr_eintrag_save', {'id': EP, 'stornieren': True})[0] == 200)
hr(ids['gina'], aktiv=False); hr(ids['vera'], aktiv=False)

# 4-Tage-Woche
st, r = P.call('hr_eintrag_save', {'person': PA, 'art': 'U', 'von': iso(AUG), 'bis': iso(AUG + dt.timedelta(days=4))})
check('4-Tage-Woche: Mo–Fr zählt 4 Urlaubstage', st == 200 and r['eintrag']['tage'] == 4, r)
zp = {p['id']: p['tage'] for p in P.call(f'hr_monat&jahr={Y}&monat=8')[1]['personen']}
fr = str((AUG + dt.timedelta(days=4)).day)
check('… Freitag als freier Tag „x“', zp[PA].get(fr) == 'x' and zp[PA].get(str(AUG.day)) == 'U', zp[PA])
zl = {p['id']: p['tage'] for p in L.call(f'hr_monat&jahr={Y}&monat=8')[1]['personen']}
check('… für Kollegen nur „abwesend“', zl[PA].get(fr) == 'A' and zl[PA].get(str(AUG.day)) == 'A', zl[PA])
check('Eintrag nur an freiem Tag (U): 422', P.call('hr_eintrag_save', {'person': PA, 'art': 'U', 'von': iso(AUG + dt.timedelta(days=11)), 'bis': iso(AUG + dt.timedelta(days=11))})[0] == 422)

# Eingetragener freier Tag („X“) sperrt keine Abwesenheit: er zählt nur nicht mit
SEP = montag(Y, 9, 2)
st, r = P.call('hr_eintrag_save', {'person': PE, 'art': 'X', 'von': iso(SEP + dt.timedelta(days=4)), 'bis': iso(SEP + dt.timedelta(days=4))})
check('freier Tag (X) eingetragen', st == 200, r)
FREI_X = r['eintrag']['id']
check('zweiter freier Tag am selben Tag: 409', P.call('hr_eintrag_save', {'person': PE, 'art': 'X', 'von': iso(SEP + dt.timedelta(days=4)), 'bis': iso(SEP + dt.timedelta(days=4))})[0] == 409)
st, r = E.call('hr_antrag', {'von': iso(SEP), 'bis': iso(SEP + dt.timedelta(days=4))})
check('Antrag über eine Woche mit freiem Tag: 200, zählt 4 Tage', st == 200 and r['eintrag']['tage'] == 4, (st, r))
ANTRAG_X = r['eintrag']['id'] if st == 200 else None
st, r = P.call('hr_eintrag_save', {'person': PE, 'art': 'K', 'von': iso(SEP + dt.timedelta(days=4)), 'bis': iso(SEP + dt.timedelta(days=4))})
check('… Krank am freien Tag: weiter 409 wegen des Antrags, nicht wegen „X“', st == 409 and 'Eintrag' in r.get('error', ''), (st, r))
check('… aufräumen', E.call('hr_antrag_zurueckziehen', {'id': ANTRAG_X})[0] == 200 and P.call('hr_eintrag_save', {'id': FREI_X, 'stornieren': True})[0] == 200)

# Antrag über den Jahreswechsel: „Resturlaub danach“ rechnet nur die Tage im Jahr des Kontos
st, r = E.call('hr_antrag', {'von': f'{Y + 1}-12-28', 'bis': f'{Y + 2}-01-08'})
check('Antrag über den Jahreswechsel', st == 200, (st, r))
JW = r['eintrag']['id'] if st == 200 else None
jw = next((a for a in P.call('hr_verwaltung')[1]['offen'] if a['id'] == JW), {})
check('… Tage im Jahr des Kontos getrennt ausgewiesen', jw.get('tageImJahr', 0) > 0 and jw['tageImJahr'] < jw['tage'] and jw['konto']['jahr'] == Y + 1, jw)
check('… zurückgezogen', E.call('hr_antrag_zurueckziehen', {'id': JW})[0] == 200)

# Übertrag aus einem vergangenen Jahr: verfallen und automatisch übertragen
for von, bis in (('2025-02-03', '2025-02-04'), ('2025-07-07', '2025-07-11')):
    assert P.call('hr_eintrag_save', {'person': PL, 'art': 'U', 'von': von, 'bis': bis})[0] == 200
pv = {p['id']: p for p in P.call('hr_verwaltung&jahr=2025')[1]['personen']}
k = pv[PL]['konto']
check('2025: 4 Tage Übertrag, 2 bis 31.03. genommen → 2 verfallen; Rest 25', k['uebertrag'] == 4 and k['verfallen'] == 2 and k['genommen'] == 7 and k['rest'] == 25, k)
pv = {p['id']: p for p in P.call('hr_verwaltung&jahr=2026')[1]['personen']}
k = pv[PL]['konto']
check('2026: Übertrag automatisch 25, Anspruch fortgeschrieben 30', k['uebertrag'] == 25 and k['anspruch'] == 30 and not k['uebertrag_manuell'], k)

# Stornieren
check('Krank stornieren', P.call('hr_eintrag_save', {'id': KRANK, 'stornieren': True})[0] == 200)
check('zweimal stornieren: 409', P.call('hr_eintrag_save', {'id': KRANK, 'stornieren': True})[0] == 409)
assert P.call('hr_eintrag_save', {'person': PL, 'art': 'K', 'von': iso(FEB), 'bis': iso(FEB + dt.timedelta(days=1))})[0] == 200

# Protokoll ohne Art der Abwesenheit
log = json.dumps(A.call('log')[1], ensure_ascii=False)
check('Protokoll nennt HR-Vorgänge, aber nie „Krank“', 'HR: Abwesenheit Lena Leser' in log and 'Krank' not in log and 'Arzt' not in log, log[:200])
raw = storage_raw('hr-abwesenheiten') or ''
check('Einträge bleiben erhalten (Status statt Löschen)', '"storniert"' in raw and '"zurueckgezogen"' in raw and '"abgelehnt"' in raw)

# ── Altbestand: Einträge der Urlaubsliste ohne Mitarbeiter (vor „Mitarbeiter = Nutzer“) ──
for u, n in (('gerda-gleich', 'Gerda Gleich'), ('hans', 'Hans Doppelt')):
    st, r_ = A.call('user_save', {'id': '', 'name': n, 'username': u, 'role': 'viewer', 'active': True, 'password': 'Urlaub-Konto-4711', 'mustChange': False})
    assert st == 200, r_
ids = {u['username']: u['id'] for u in A.call('users')[1]['users']}
alt = dokument('hr-personal')
alt['personen'] += [{'id': 'pgerda', 'vorname': 'Gerda', 'nachname': 'Gleich', 'abteilung': 'Lager', 'pos': 4, 'aktiv': True, 'jahre': {str(Y): {'anspruch': 28}}},
                    {'id': 'phans1', 'vorname': 'Hans', 'nachname': 'Doppelt', 'abteilung': 'Lager', 'pos': 5, 'aktiv': True},
                    {'id': 'phans2', 'vorname': 'Hans', 'nachname': 'Doppelt', 'abteilung': 'GF', 'pos': 6, 'aktiv': True},
                    {'id': 'pkarl', 'vorname': 'Karl', 'nachname': 'Alt', 'abteilung': 'Lager', 'pos': 7, 'aktiv': True}]
alt['rev'] = alt.get('rev', 0) + 1
storage_write('hr-personal', alt)
ab = dokument('hr-abwesenheiten')
ab['eintraege'] += [{'id': 'aalt1', 'person': 'pgerda', 'art': 'U', 'von': iso(montag(Y, 6)), 'bis': iso(montag(Y, 6)), 'status': 'genehmigt', 'notiz': '', 'quelle': 'personalstelle'},
                    {'id': 'aalt2', 'person': 'pkarl', 'art': 'U', 'von': iso(montag(Y, 6)), 'bis': iso(montag(Y, 6)), 'status': 'genehmigt', 'notiz': '', 'quelle': 'personalstelle'}]
ab['rev'] = ab.get('rev', 0) + 1
storage_write('hr-abwesenheiten', ab)
v = P.call('mitarbeiter_hr')[1]
check('Verwaltung zeigt den Altbestand (ohne Mitarbeiter)', sorted(p_['id'] for p_ in v['personen'] if p_['altbestand']) == ['pgerda', 'phans1', 'phans2', 'pkarl'])
G = Client('198.51.100.8'); G.call('login', {'username': 'gerda-gleich', 'password': 'Urlaub-Konto-4711'})
st, r = G.call('hr_start&jahr=' + str(Y))
check('erster Aufruf: gleichnamiger Eintrag automatisch übernommen (mit Urlaub und Anspruch)', st == 200 and r['neuVerknuepft'] is True and r['ich']['id'] == 'pgerda'
      and r['konto']['anspruch'] == 28 and any(e['id'] == 'aalt1' for e in r['eintraege']), r.get('ich'))
check('… danach nicht erneut gemeldet', G.call('hr_start')[1]['neuVerknuepft'] is False)
log = json.dumps(A.call('log')[1], ensure_ascii=False)
check('automatische Übernahme im Protokoll', 'automatisch dem Mitarbeiter Gerda Gleich zugeordnet' in log)
H = Client('198.51.100.9'); H.call('login', {'username': 'hans', 'password': 'Urlaub-Konto-4711'})
st, r = H.call('hr_start')
check('gleicher Name zweimal im Altbestand: keine automatische Übernahme', st == 200 and r['neuVerknuepft'] is False and r['ich']['id'] == ids['hans'], r.get('ich'))
# Ohne Mitarbeiter = ehemaliger Mitarbeiter (Inhaber 26.10.2026): nicht im Urlaubsplan, Einträge bleiben, keine neuen Mitarbeiter
v = P.call('mitarbeiter_hr')[1]
pv = {p_['id']: p_ for p_ in v['personen']}
check('ehemalige Mitarbeiter: markiert, nicht im Urlaubsplan', pv['pkarl']['ehemalig'] is True and pv['pkarl']['aktiv'] is False, pv['pkarl'])
check('… nicht im Kalender, nicht als Vertretung', 'pkarl' not in [p_['id'] for p_ in P.call(f'hr_monat&jahr={Y}&monat=6')[1]['personen']]
      and 'pkarl' not in [k['id'] for k in E.call('hr_start')[1]['kollegen']])
check('… ihre Einträge bleiben in „Alle Einträge“', any(e['person'] == 'pkarl' for e in P.call('hr_verwaltung&jahr=' + str(Y))[1]['eintraege']))
anzahl_nutzer = len(A.call('users')[1]['users'])
st, r = P.call('hr_personal_uebernehmen', {})
check('Gleichnamige zuordnen: nur eindeutig gleichnamige Mitarbeiter, keine neuen Mitarbeiter', st == 200 and r['anzahl'] == 1 and len(A.call('users')[1]['users']) == anzahl_nutzer, r)
v = P.call('mitarbeiter_hr')[1]
rev = v['rev']
pv = {p_['id']: p_ for p_ in v['personen']}
check('… Hans Doppelt: einer beim Mitarbeiter, der zweite bleibt ehemalig; Karl Alt bleibt ehemalig', pv['phans1']['nutzer'] == ids['hans']
      and sorted(p_['id'] for p_ in v['personen'] if p_['altbestand']) == ['phans2', 'pkarl'], sorted(p_['id'] for p_ in v['personen'] if p_['altbestand']))
hr(ids['hans'], aktiv=False)

# ── Altbestand einzeln zuordnen (abweichender Name, Mitarbeiter mit HR-Angaben: zusammenführen) ──
for u, n in (('mara', 'Mara Konto'), ('nina', 'Nina Neu')):
    assert A.call('user_save', {'id': '', 'name': n, 'username': u, 'role': 'viewer', 'active': True, 'password': 'Urlaub-Konto-4711', 'mustChange': False})[0] == 200
ids = {u['username']: u['id'] for u in A.call('users')[1]['users']}
rev = P.call('mitarbeiter_hr')[1]['rev']
PM = hr(ids['mara'], abteilung='Vertrieb', pos=40, jahr=Y, anspruch=20)
JUL = montag(Y, 7)
alt = dokument('hr-personal')
alt['personen'] += [{'id': 'pmarlene', 'vorname': 'Marlene', 'nachname': 'Konto-Alt', 'abteilung': 'Lager', 'pos': 41, 'aktiv': True, 'personalnummer': '9901',
                     'jahre': {str(Y): {'anspruch': 25, 'uebertrag': 3}, str(Y - 1): {'anspruch': 24}}},
                    {'id': 'pmarlene2', 'vorname': 'Marlene', 'nachname': 'Zwei', 'abteilung': 'Lager', 'pos': 42, 'aktiv': True, 'personalnummer': '9902'}]
alt['rev'] = alt.get('rev', 0) + 1
storage_write('hr-personal', alt)
ab = dokument('hr-abwesenheiten')
ab['eintraege'] += [{'id': 'am1', 'person': 'pmarlene', 'art': 'U', 'von': iso(JUL), 'bis': iso(JUL + dt.timedelta(days=1)), 'status': 'genehmigt', 'notiz': '', 'quelle': 'zeiterfassung'},
                    {'id': 'am2', 'person': 'pmarlene', 'art': 'K', 'von': iso(JUL + dt.timedelta(days=7)), 'bis': iso(JUL + dt.timedelta(days=7)), 'status': 'genehmigt', 'notiz': '', 'quelle': 'zeiterfassung'},
                    {'id': 'am3', 'person': 'pmarlene', 'art': 'U', 'von': iso(JUL + dt.timedelta(days=14)), 'bis': iso(JUL + dt.timedelta(days=15)), 'status': 'genehmigt', 'notiz': '', 'quelle': 'zeiterfassung'}]
ab['rev'] = ab.get('rev', 0) + 1
storage_write('hr-abwesenheiten', ab)
# Mara hat schon denselben Urlaub (doppelt) und einen überschneidenden
assert P.call('hr_eintrag_save', {'person': PM, 'art': 'U', 'von': iso(JUL), 'bis': iso(JUL + dt.timedelta(days=1))})[0] == 200
st, r = P.call('hr_eintrag_save', {'person': PM, 'art': 'U', 'von': iso(JUL + dt.timedelta(days=15)), 'bis': iso(JUL + dt.timedelta(days=16))})
UEBER = r['eintrag']['id']
v = P.call('mitarbeiter_hr')[1]
rev = v['rev']
st, r = P.call('mitarbeiter_hr_save', {'nutzer': ids['mara'], 'rev': rev, 'personalnummer': '9901'})
check('Personalnummer schon beim Eintrag ohne Mitarbeiter: 409 mit Hinweis auf „Einzeln zuordnen“', st == 409 and 'Einzeln zuordnen' in r.get('error', ''), r)
check('Zuordnen ohne Personalrecht: 403', L.call('hr_altbestand_zuordnen', {'eintrag': 'pmarlene', 'nutzer': ids['mara'], 'rev': rev})[0] == 403)
check('Zuordnen: unbekannter Mitarbeiter 422, veralteter Stand 409', P.call('hr_altbestand_zuordnen', {'eintrag': 'pmarlene', 'nutzer': 'u-gibt-es-nicht', 'rev': rev})[0] == 422
      and P.call('hr_altbestand_zuordnen', {'eintrag': 'pmarlene', 'nutzer': ids['mara'], 'rev': rev - 1})[0] == 409)
st, r = P.call('hr_altbestand_zuordnen', {'eintrag': 'pmarlene', 'nutzer': ids['mara'], 'rev': rev})
check('Überschneidung mit vorhandener Abwesenheit: 409, nichts geändert', st == 409 and 'Überschneidende' in r.get('error', '')
      and 'pmarlene' in [p_['id'] for p_ in P.call('mitarbeiter_hr')[1]['personen'] if p_['altbestand']], r)
assert P.call('hr_eintrag_save', {'id': UEBER, 'stornieren': True})[0] == 200
anzahl_vorher = len(dokument('hr-abwesenheiten')['eintraege'])
st, r = P.call('hr_altbestand_zuordnen', {'eintrag': 'pmarlene', 'nutzer': ids['mara'], 'rev': rev})
check('zusammengeführt: 2 Abwesenheiten übernommen, 1 doppelte storniert', st == 200 and r['abwesenheiten'] == 2 and r['doppelt'] == 1, (st, r))
v = P.call('mitarbeiter_hr')[1]
pm = {p_['id']: p_ for p_ in v['personen']}[PM]
check('… Angaben: Mitarbeiter gilt, leere Felder aus dem Eintrag (Personalnummer, Übertrag, Vorjahr)', pm['personalnummer'] == '9901' and pm['abteilung'] == 'Vertrieb'
      and pm['jahre'][str(Y)] == {'anspruch': 20, 'uebertrag': 3} and pm['jahre'][str(Y - 1)]['anspruch'] == 24 and pm['nutzer'] == ids['mara'], pm)
check('… Eintrag ohne Mitarbeiter verschwunden', 'pmarlene' not in [p_['id'] for p_ in v['personen']])
abw = dokument('hr-abwesenheiten')['eintraege']
am = {e['id']: e for e in abw if e['id'] in ('am1', 'am2', 'am3')}
check('… Abwesenheiten beim Mitarbeiter, doppelte storniert mit Grund, nichts gelöscht', len(abw) == anzahl_vorher and all(e['person'] == PM for e in am.values())
      and am['am1']['status'] == 'storniert' and 'doppelt' in am['am1']['entscheidung']['grund'] and am['am2']['status'] == 'genehmigt', am)
check('… im Protokoll', 'Eintrag der Urlaubsliste „Marlene Konto-Alt“ dem Mitarbeiter Mara Konto zugeordnet' in json.dumps(A.call('log')[1], ensure_ascii=False))
rev = v['rev']
check('schon zugeordneter Eintrag: 404', P.call('hr_altbestand_zuordnen', {'eintrag': 'pmarlene', 'nutzer': ids['mara'], 'rev': rev})[0] == 404)
st, r = P.call('hr_altbestand_zuordnen', {'eintrag': 'pmarlene2', 'nutzer': ids['mara'], 'rev': rev})
check('unterschiedliche Personalnummern: 409', st == 409 and 'Personalnummern' in r.get('error', ''), r)
st, r = P.call('hr_altbestand_zuordnen', {'eintrag': 'pmarlene2', 'nutzer': ids['nina'], 'rev': rev})
v = P.call('mitarbeiter_hr')[1]
pn = [p_ for p_ in v['personen'] if p_['nutzer'] == ids['nina']]
check('Mitarbeiter ohne HR-Angaben: Eintrag wird seine Person (Personalnummer bleibt)', st == 200 and len(pn) == 1 and pn[0]['id'] == 'pmarlene2' and pn[0]['personalnummer'] == '9902', (st, r, pn))
rev = v['rev']
hr(ids['mara'], aktiv=False); hr(ids['nina'], aktiv=False)

# ── Verwaiste Einträge: verknüpfter Mitarbeiter gelöscht bzw. zweiter Eintrag desselben Mitarbeiters (Fehler aus dem Betrieb
#    26.10.2026: „schon einem Mitarbeiter zugeordnet“, obwohl der Eintrag als „ohne Mitarbeiter“ angezeigt wurde) ──
for u, n in (('paula', 'Paula Platz'), ('quentin', 'Quentin Quer')):
    assert A.call('user_save', {'id': '', 'name': n, 'username': u, 'role': 'viewer', 'active': True, 'password': 'Urlaub-Konto-4711', 'mustChange': False})[0] == 200
ids = {u['username']: u['id'] for u in A.call('users')[1]['users']}
alt = dokument('hr-personal')
alt['personen'] += [{'id': 'pverwaist', 'nutzer': 'u-geloescht-1', 'vorname': 'Paula', 'nachname': 'Alt', 'abteilung': 'Lager', 'pos': 44, 'aktiv': True, 'personalnummer': '9910'},
                    {'id': 'pdoppel', 'nutzer': ids['erika'], 'vorname': 'Erika', 'nachname': 'Doppelt', 'abteilung': '', 'pos': 45, 'aktiv': True},
                    {'id': 'pquentin', 'nutzer': 'u-geloescht-2', 'vorname': 'Quentin', 'nachname': 'Quer', 'abteilung': 'Lager', 'pos': 46, 'aktiv': True}]
alt['rev'] = alt.get('rev', 0) + 1
storage_write('hr-personal', alt)
ab = dokument('hr-abwesenheiten')
ab['eintraege'].append({'id': 'av1', 'person': 'pverwaist', 'art': 'U', 'von': iso(montag(Y, 9, 2)), 'bis': iso(montag(Y, 9, 2)), 'status': 'genehmigt', 'notiz': '', 'quelle': 'zeiterfassung'})
ab['rev'] = ab.get('rev', 0) + 1
storage_write('hr-abwesenheiten', ab)
v = P.call('mitarbeiter_hr')[1]
rev = v['rev']
altb = sorted(p_['id'] for p_ in v['personen'] if p_['altbestand'])
check('verwaiste Einträge erscheinen als „ohne Mitarbeiter“ (gelöschter Mitarbeiter, zweiter Eintrag)', altb == ['pdoppel', 'phans2', 'pkarl', 'pquentin', 'pverwaist'], altb)
pe_vorher = {p_['id']: p_ for p_ in v['personen']}[PE]
st, r = P.call('hr_altbestand_zuordnen', {'eintrag': 'pverwaist', 'nutzer': ids['paula'], 'rev': rev})
v = P.call('mitarbeiter_hr')[1]
pp = [p_ for p_ in v['personen'] if p_['nutzer'] == ids['paula']]
check('Eintrag eines gelöschten Mitarbeiters zuordnen (Mitarbeiter ohne Angaben): ok, Personalnummer und Urlaub kommen mit', st == 200
      and len(pp) == 1 and pp[0]['id'] == 'pverwaist' and pp[0]['personalnummer'] == '9910' and r['abwesenheiten'] == 1, (st, r, pp))
rev = v['rev']
st, r = P.call('hr_altbestand_zuordnen', {'eintrag': 'pdoppel', 'nutzer': ids['erika'], 'rev': rev})
v = P.call('mitarbeiter_hr')[1]
pe = {p_['id']: p_ for p_ in v['personen']}.get(PE, {})
check('zweiter Eintrag desselben Mitarbeiters: in seine Angaben zusammengeführt, Angaben unverändert', st == 200 and 'pdoppel' not in [p_['id'] for p_ in v['personen']]
      and pe.get('abteilung') == pe_vorher['abteilung'] and pe.get('jahre') == pe_vorher['jahre'], (st, r))
st, r = P.call('hr_personal_uebernehmen', {})
v = P.call('mitarbeiter_hr')[1]
pq = [p_ for p_ in v['personen'] if p_['nutzer'] == ids['quentin']]
check('„Gleichnamige zuordnen“ erfasst auch verwaiste Einträge', st == 200 and r['anzahl'] == 1 and len(pq) == 1 and pq[0]['id'] == 'pquentin'
      and sorted(p_['id'] for p_ in v['personen'] if p_['altbestand']) == ['phans2', 'pkarl'], (r, pq))
rev = v['rev']
hr(ids['paula'], aktiv=False); hr(ids['quentin'], aktiv=False)

# ── Excel-Export ──
# Lena nicht mehr im Urlaubsplan (z. B. ausgeschieden): im Februar (genehmigte Krankheit) weiter im Export, im August nicht
hr(ids['lena'], aktiv=False)
st, r, body = P.call('hr_export', {'jahr': Y, 'monate': [2]}, raw=True)
check('Export: 200, xlsx', st == 200 and body.startswith('PK'), (st, body[:80]))
import urllib.request
req = urllib.request.Request(B + '/api/?action=hr_export', data=json.dumps({'jahr': Y, 'monate': [2, 8]}).encode(),
                             headers={'Content-Type': 'application/json', 'X-CSRF-Token': P.csrf, 'X-Forwarded-For': P.ip})
antwort = P.http.open(req)
daten = antwort.read()
check('Dateiname und Typ', 'Urlaub_' in antwort.headers.get('Content-Disposition', '') and 'spreadsheetml' in antwort.headers.get('Content-Type', ''),
      antwort.headers.get('Content-Disposition'))
check('gültiges ZIP', zipfile.ZipFile(io.BytesIO(daten)).testzip() is None)
wb = openpyxl.load_workbook(io.BytesIO(daten))
check('je Monat ein Blatt', wb.sheetnames == [f'Februar {Y}', f'August {Y}'], wb.sheetnames)
ws = wb[f'Februar {Y}']
check('Aufbau wie die Planung: Legende, Titel K2, Pos. A8, Monat C8, Tag 1 in D8', ws['C2'].value == 'Urlaub' and ws['D2'].value == 'U' and ws['K2'].value == f'Februar {Y}'
      and ws['A8'].value == 'Pos.' and ws['C8'].value == 'Februar' and ws['D8'].value == 1 and ws.freeze_panes == 'D9',
      [ws['C2'].value, ws['K2'].value, ws['A8'].value, ws['D8'].value, ws.freeze_panes])
zeile = {ws.cell(z, 3).value: z for z in range(9, 20) if ws.cell(z, 3).value}
nt = 28 + (1 if Y % 4 == 0 and (Y % 100 or Y % 400 == 0) else 0)
spU = 3 + nt + 1
werte = lambda z: [ws.cell(z, c).value for c in range(4, 4 + nt)]
check('Erika: 3 × U (nur genehmigt, offener Antrag fehlt)', werte(zeile['Erika Editor']).count('U') == 3, werte(zeile['Erika Editor']))
check('Lena: 2 × K', werte(zeile['Lena Leser']).count('K') == 2)
check('Summenformeln wie in der Planung', str(ws.cell(zeile['Erika Editor'], spU).value).startswith('=COUNTIF(D') and ws.cell(7, spU).value == 'U', ws.cell(zeile['Erika Editor'], spU).value)
check('Abteilungen als senkrechte, verbundene Zellen', any(str(m).startswith('B') for m in ws.merged_cells.ranges) and ws.cell(zeile['Erika Editor'], 2).alignment.textRotation == 90)
ws8 = wb[f'August {Y}']
z8 = {ws8.cell(z, 3).value: z for z in range(9, 20) if ws8.cell(z, 3).value}
check('August: 4-Tage-Woche mit „x“', 'x' in [ws8.cell(z8['Andre Vierzig'], c).value for c in range(4, 35)])
check('nicht mehr im Urlaubsplan: nur in Monaten mit genehmigten Abwesenheiten im Export', 'Lena Leser' in zeile and 'Lena Leser' not in z8, (sorted(zeile), sorted(z8)))
hr(ids['lena'], aktiv=True)
check('Export ohne Monat: 422', P.call('hr_export', {'jahr': Y, 'monate': []})[0] == 422)

# ── Verfallstag des Resturlaubs einstellbar (Verwaltung → Urlaub; Standard 31.03.) ──
assert A.call('user_save', {'id': '', 'name': 'Fritz Frist', 'username': 'fritz', 'role': 'viewer', 'active': True, 'password': 'Urlaub-Konto-4711', 'mustChange': False})[0] == 200
ids = {u['username']: u['id'] for u in A.call('users')[1]['users']}
rev = P.call('mitarbeiter_hr')[1]['rev']
PF = hr(ids['fritz'], abteilung='Lager', pos=50, jahr=Y, anspruch=20, uebertrag=5)
APR = montag(Y, 4, 2)
st, r = P.call('hr_eintrag_save', {'person': PF, 'art': 'U', 'von': iso(APR), 'bis': iso(APR + dt.timedelta(days=4))})
NAPR = r['eintrag']['tage']
konto_f = lambda: next(p_ for p_ in P.call(f'mitarbeiter_hr&jahr={Y}')[1]['personen'] if p_['id'] == PF)['konto']
v = P.call('mitarbeiter_hr')[1]
check('Standard: Verfallstag 31.03.', v['uebertragBis'] == '03-31' and E.call('hr_start')[1]['uebertragBis'] == '03-31', v.get('uebertragBis'))
k = konto_f()
check('… Urlaub im April verbraucht den Übertrag nicht (verfällt ganz)', k['frist'] == f'{Y}-03-31' and k['verfaellt'] == 5, k)
einst = {'sperren': v['sperren'], 'ferien': v['ferien']}
for falsch in ('02-29', f'{Y}-13-01', '31.04.', 'x'):
    st, r = P.call('hr_einstellungen_save', {**einst, 'uebertragBis': falsch})
    check(f'ungültiger Verfallstag {falsch}: 422', st == 422 and 'Verfallstag' in r.get('error', ''), (st, r))
check('ohne Personalrecht: 403', L.call('hr_einstellungen_save', {**einst, 'uebertragBis': '04-30'})[0] == 403)
st, r = P.call('hr_einstellungen_save', {**einst, 'uebertragBis': f'{Y}-04-30'})
check('Verfallstag 30.04. gespeichert (aus dem Datumsfeld nur Tag und Monat)', st == 200 and r['uebertragBis'] == '04-30', r)
k = konto_f()
check('… Konto: Frist 30.04., Urlaub im April verbraucht den Übertrag zuerst', k['frist'] == f'{Y}-04-30' and k['verfaellt'] == max(0, 5 - NAPR), (k, NAPR))
check('… gilt überall (Start, Verwaltung, Übersicht)', E.call('hr_start')[1]['uebertragBis'] == '04-30'
      and P.call(f'hr_uebersicht&jahr={Y}')[1]['uebertragBis'] == '04-30')
st, r = P.call('hr_einstellungen_save', einst)
check('Speichern ohne Verfallstag lässt ihn unverändert', st == 200 and r['uebertragBis'] == '04-30' and P.call('mitarbeiter_hr')[1]['uebertragBis'] == '04-30', r)
check('Änderung im Protokoll', 'Verfallstag Resturlaub auf 30.04. geändert' in json.dumps(A.call('log')[1], ensure_ascii=False))
st, r = P.call('hr_einstellungen_save', {**einst, 'uebertragBis': '03-31'})
check('Rückweg auf 31.03.', st == 200 and r['uebertragBis'] == '03-31' and konto_f()['verfaellt'] == 5, r)
rev = P.call('mitarbeiter_hr')[1]['rev']
hr(ids['fritz'], aktiv=False)

# ── Übernahme der Excel-Planung (synthetische Datei) im Browser ──
fehler = []
with sync_playwright() as p:
    br = p.chromium.launch()

    def anmelden(user, pw, ip, breite=1280):
        ctx = br.new_context(viewport={'width': breite, 'height': 900}, extra_http_headers={'X-Forwarded-For': ip}, locale='de-DE', timezone_id='Europe/Berlin')
        pg = ctx.new_page(); pg.on('pageerror', lambda e: fehler.append(str(e)))
        pg.goto(B + '/admin/'); pg.wait_for_selector('#screen-auth:not([hidden])')
        pg.fill('#au-user', user); pg.fill('#pw1', pw); pg.click('#auth-submit')
        pg.wait_for_selector('#screen-auth', state='hidden')
        return ctx, pg

    # ── Verwaltung → Import: bisherige Excel-Planung (Importzentrale, Personalstelle ohne Administratorrechte) ──
    ctx, pg = anmelden('paul', 'Personal-Zugang-77', '198.51.100.5')
    pg.click('#vt-importe'); pg.wait_for_selector('#view-importe .imp-art[data-art="urlaubsplanung"]')
    check('Importzentrale: Importart „Urlaubsplanung“ für die Personalstelle', 'HR' in pg.text_content('#view-importe .imp-art[data-art="urlaubsplanung"]'))
    pg.set_input_files('#view-importe input[data-imp-datei="urlaubsplanung"]', os.path.join(FIX, 'hr', 'Urlaub_Muster_2025.xlsx'))
    pg.wait_for_selector('#imp-neu', timeout=20000)
    txt = pg.inner_text('#view-importe')
    check('Vorschau: neue Mitarbeiter ohne Anmeldung, unbekanntes Kürzel, regelmäßige freie Tage', 'Anna Beispiel' in txt and 'ohne Anmeldung' in txt and '„S“ (1×)' in txt
          and 'Dieter Test (meist Fr)' in txt, txt[:600])
    pg.screenshot(path=os.path.join(SH, 'hr_import_1280.png'), full_page=True)
    pg.click('#imp-ok'); pg.wait_for_selector('#imp-ergebnis')
    st, r = P.call('hr_verwaltung&jahr=2025')
    namen = {p_['name']: p_ for p_ in r['personen']}
    anzahl_personen = len(r['personen'])
    imp = [e for e in r['eintraege'] if e['quelle'] == 'import']
    check('Übernommen: Mitarbeiter mit Abteilung (verbundene Zellen fortgeführt)', namen.get('Emil Lager', {}).get('abteilung') == 'Lager' and namen.get('Bernd Muster', {}).get('abteilung') == 'Verwaltung'
          and namen.get('Carla Probe', {}).get('abteilung') == 'Verwaltung', {k: v['abteilung'] for k, v in namen.items()})
    nutzer = {u['name']: u for u in A.call('users')[1]['users']}
    check('Übernommen: unbekannte Namen sind jetzt Mitarbeiter ohne Anmeldung', all(nutzer.get(n, {}).get('zugang') is False for n in ('Anna Beispiel', 'Emil Lager', 'Dieter Test')), list(nutzer))
    arten = sorted((e['name'], e['art'], e['von'], e['bis']) for e in imp if e['art'] != 'X')
    check('Übernommen: Zeiträume zusammengefasst', arten == [('Anna Beispiel', 'U', '2025-01-13', '2025-01-17'), ('Bernd Muster', 'K', '2025-01-20', '2025-01-21'),
                                                          ('Carla Probe', 'G', '2025-01-31', '2025-01-31'), ('Carla Probe', 'U', '2025-02-03', '2025-02-04')], arten)
    check('Übernommen: freie Tage „x“ als einzelne Tage', len([e for e in imp if e['art'] == 'X']) == 9 and all(e['status'] == 'genehmigt' for e in imp))
    # erneute Übernahme ersetzt (keine Doppelten, Namen jetzt bekannt)
    pg.click('[data-imp="liste"]'); pg.wait_for_selector('#view-importe input[data-imp-datei="urlaubsplanung"]', state='attached')
    pg.set_input_files('#view-importe input[data-imp-datei="urlaubsplanung"]', os.path.join(FIX, 'hr', 'Urlaub_Muster_2025.xlsx'))
    pg.wait_for_selector('#imp-neu', timeout=20000)
    check('erneute Übernahme: frühere Einträge werden ersetzt, keine neuen Mitarbeiter', '13 Einträge der letzten Übernahme werden ersetzt' in pg.inner_text('#view-importe')
          and 'ohne Anmeldung' not in pg.inner_text('#imp-zeilen'), pg.inner_text('#view-importe')[:400])
    pg.click('#imp-ok'); pg.wait_for_selector('#imp-ergebnis')
    r = P.call('hr_verwaltung&jahr=2025')[1]
    check('… keine doppelten Einträge oder Mitarbeiter', len([e for e in r['eintraege'] if e['quelle'] == 'import']) == 13 and len(r['personen']) == anzahl_personen, (len(r['eintraege']), len(r['personen'])))
    k26 = {p_['name']: p_['konto'] for p_ in P.call('hr_verwaltung&jahr=2026')[1]['personen']}
    check('ohne bekannten Anspruch im Vorjahr kein (negativer) automatischer Übertrag', k26['Anna Beispiel']['uebertrag'] == 0, k26['Anna Beispiel'])
    abt = [p_['abteilung'] for p_ in P.call(f'hr_monat&jahr={Y}&monat=1')[1]['personen']]
    bloecke = [x for i, x in enumerate(abt) if i == 0 or abt[i - 1] != x]
    check('Kalender: jede Abteilung ein Block (Reihenfolge nach Position)', len(bloecke) == len(set(bloecke)) and bloecke[0] == 'GF', abt)

    # ── Abwesenheiten aus der Zeiterfassung (Importvorlage, synthetischer Abwesenheitsbericht passend zur Planung 2025) ──
    BERICHT = os.path.join(FIX, 'hr', 'Abwesenheitsbericht_Muster.xlsx')
    wsb = openpyxl.load_workbook(BERICHT).active
    D = {'blaetter': [{'name': wsb.title, 'grid': [[(v - dt.datetime(1899, 12, 30)).days if isinstance(v, dt.datetime) else v for v in z]
                                                   for z in wsb.iter_rows(values_only=True)]}]}

    def zeit(cl, uebernehmen=False, stand=None, **opt):
        b = {'art': 'zeiterfassung', 'daten': D, 'quelle': 'Abwesenheitsbericht_Muster.xlsx', 'optionen': opt}
        if uebernehmen: b['stand'] = stand
        return cl.call('import_uebernehmen' if uebernehmen else 'import_pruefen', b)

    check('Zeiterfassung: ohne Personalrecht 403', zeit(L)[0] == 403 and zeit(E)[0] == 403)
    st, r = P.call('import_pruefen', {'art': 'zeiterfassung', 'daten': {'blaetter': [{'name': 'x', 'grid': [['Artikel', 'Preis'], ['A', 1]]}]}})
    check('Zeiterfassung: andere Datei 422 mit Hinweis', st == 422 and 'Abwesenheitsbericht' in r.get('error', ''), (st, r))
    st, v = zeit(P)
    zl = {(z['art'], z['name']): z for z in v['zeilen'] if z['schluessel'] != 'Abwesenheiten'}
    hw = ' '.join(h['text'] for h in v['hinweise'])
    check('Vorschau: unbekannte Mitarbeiter neu (mit/ohne Personalnummer), Feiertage nicht übernommen', st == 200 and ('neu', 'Hans Neumann') in zl
          and 'Nr. 105' == zl[('neu', 'Hans Neumann')]['schluessel'] and ('neu', 'Ina Ohnenummer') in zl, sorted(zl))
    check('Vorschau: ähnliche Schreibweise als Hinweis, nie automatisch zugeordnet', 'ähnlich: Carla Probe' in zl.get(('neu', 'Karla Probe'), {}).get('text', ''), zl.get(('neu', 'Karla Probe')))
    check('Vorschau: schon gleich eingetragene Tage bleiben (5 Tage Anna)', v['zahlen']['unveraendert'] == 5, v['zahlen'])
    check('Vorschau: im System anders eingetragen → Fehlerzeile, bleibt unverändert', '20.01.2025 (K statt U)' in zl.get(('fehler', 'Bernd Muster'), {}).get('text', ''), zl.get(('fehler', 'Bernd Muster')))
    check('Vorschau: ungültiges Datum und unbekannte Art gemeldet', any('Hans Neumann: Start- oder Enddatum' in z['text'] for z in v['zeilen']) and '„Elternzeit“ (1×)' in hw, hw)
    # gleiche Person, andere Schreibweise: Personalnummer am Mitarbeiter eintragen → Zuordnung über die Nummer
    rev = P.call('mitarbeiter_hr')[1]['rev']
    hr(nutzer['Carla Probe']['id'], personalnummer='103')
    check('Personalnummer eindeutig (409)', P.call('mitarbeiter_hr_save', {'nutzer': ids['erika'], 'rev': rev, 'personalnummer': '103'})[0] == 409)
    check('Personalnummer: ungültige Zeichen 422', P.call('mitarbeiter_hr_save', {'nutzer': ids['erika'], 'rev': rev, 'personalnummer': 'ab cd'})[0] == 422)
    st, v2 = zeit(P, abteilungen=True)
    check('Option Abteilungen: Abteilung aus der Zeiterfassung als Änderung', any(z['art'] == 'geaendert' and z['name'] == 'Anna Beispiel' and 'Abteilung Geschäftsführung' in z['text'] for z in v2['zeilen']), v2['zeilen'][:6])
    # Übernahme in der Oberfläche (Importzentrale)
    pg.click('[data-imp="liste"]'); pg.wait_for_selector('#view-importe input[data-imp-datei="zeiterfassung"]', state='attached')
    check('Importzentrale: Importart „Abwesenheiten aus der Zeiterfassung“', 'Zeiterfassung' in pg.text_content('#view-importe .imp-art[data-art="zeiterfassung"]'))
    pg.set_input_files('#view-importe input[data-imp-datei="zeiterfassung"]', BERICHT)
    pg.wait_for_selector('#imp-neu', timeout=20000)
    txt = pg.inner_text('#view-importe')
    check('Oberfläche: Vorschau über die Personalnummer (Carla statt neuer Karla)', 'Hans Neumann' in txt and 'Karla' not in txt and 'Carla Probe' in txt, txt[:700])
    pg.screenshot(path=os.path.join(SH, 'hr_zeiterfassung_1280.png'), full_page=True)
    pg.set_viewport_size({'width': 390, 'height': 844})
    check('390px: Vorschau ohne waagerechtes Scrollen', pg.evaluate('document.documentElement.scrollWidth <= window.innerWidth + 1'))
    pg.screenshot(path=os.path.join(SH, 'hr_zeiterfassung_390.png'), full_page=True)
    pg.set_viewport_size({'width': 1280, 'height': 900})
    pg.click('#imp-ok'); pg.wait_for_selector('#imp-ergebnis')
    r = P.call('hr_verwaltung&jahr=2025')[1]
    ze = sorted((e['name'], e['art'], e['von'], e['bis'], e['urlaubAbzug'], e['status']) for e in r['eintraege'] if e['quelle'] == 'zeiterfassung')
    check('Übernommen: Zeiträume (über Wochenenden zusammengefasst), genehmigt', ze == [
        ('Bernd Muster', 'K', '2025-03-05', '2025-03-05', True, 'genehmigt'), ('Carla Probe', 'U', '2025-03-17', '2025-03-18', False, 'genehmigt'),
        ('Erika Editor', 'K', '2025-02-10', '2025-02-14', False, 'genehmigt'), ('Erika Editor', 'U', '2025-03-24', '2025-03-28', False, 'genehmigt'),
        ('Hans Neumann', 'U', '2025-02-27', '2025-03-04', False, 'genehmigt')], ze)
    kk = {p_['name']: p_['konto'] for p_ in r['personen']}
    check('Konto: „Krank ohne AU / vom Urlaub abgezogen“ zählt als Urlaub, Krank mit AU nicht', kk['Bernd Muster']['genommen'] == 1 and kk['Erika Editor']['genommen'] == 5
          and kk['Hans Neumann']['genommen'] == 4, {k: kk[k]['genommen'] for k in ('Bernd Muster', 'Erika Editor', 'Hans Neumann')})
    check('Übernommen: Planung (K 20.–21.01.) unverändert', any(e['name'] == 'Bernd Muster' and e['art'] == 'K' and e['von'] == '2025-01-20' and e['quelle'] == 'import' for e in r['eintraege']))
    pm = {p_['name']: p_ for p_ in P.call('mitarbeiter_hr')[1]['personen']}
    check('Übernommen: Personalnummern gespeichert, neue Mitarbeiter mit Abteilung, bestehende Abteilung unverändert (Option aus)',
          [pm[n]['personalnummer'] for n in ('Anna Beispiel', 'Bernd Muster', 'Carla Probe', 'Erika Editor', 'Hans Neumann', 'Ina Ohnenummer')] == ['101', '102', '103', '104', '105', '']
          and pm['Hans Neumann']['abteilung'] == 'Produktion' and pm['Anna Beispiel']['abteilung'] == 'GF' and pm['Hans Neumann']['zugang'] is False,
          {n: (pm[n]['personalnummer'], pm[n]['abteilung']) for n in pm if n in ('Anna Beispiel', 'Hans Neumann', 'Ina Ohnenummer')})
    # erneute Übernahme ersetzt die frühere, veralteter Stand wird abgewiesen
    st, v3 = zeit(P)
    check('erneut: 5 frühere Einträge ersetzt, keine neuen Mitarbeiter, 7 Tage unverändert', v3['zahlen']['entfaellt'] == 5 and not any(z['art'] == 'neu' and z['schluessel'] != 'Abwesenheiten' for z in v3['zeilen'])
          and v3['zahlen']['unveraendert'] == 7, v3['zahlen'])
    rev = P.call('mitarbeiter_hr')[1]['rev']
    hr(nutzer['Carla Probe']['id'], pos=3)
    check('Übernahme mit veraltetem Stand: 409', zeit(P, True, v3['stand'])[0] == 409)
    st, v3 = zeit(P)
    check('Übernahme mit aktuellem Stand: 200', zeit(P, True, v3['stand'])[0] == 200)
    r = P.call('hr_verwaltung&jahr=2025')[1]
    check('… keine doppelten Einträge oder Mitarbeiter', len([e for e in r['eintraege'] if e['quelle'] == 'zeiterfassung']) == 5 and len(r['personen']) == anzahl_personen + 2, (len(r['personen']), anzahl_personen))
    # gleicher Name, andere Personalnummer: andere Person – nie automatisch zuordnen, Personalnummer nicht überschreiben
    D2 = json.loads(json.dumps(D))
    for z in D2['blaetter'][0]['grid']:
        if z[0] == 'Anna' and z[1] == 'Beispiel': z[7] = '199'
    b2 = {'art': 'zeiterfassung', 'daten': D2, 'quelle': 'Abwesenheitsbericht_Muster.xlsx', 'optionen': {}}
    st, v4 = P.call('import_pruefen', b2)
    z4 = {(z['art'], z['name']): z for z in v4.get('zeilen', []) if z['schluessel'] != 'Abwesenheiten'}
    check('gleicher Name, andere Personalnummer: als neu mit Hinweis, nicht der bestehenden Anna zugeordnet',
          st == 200 and ('neu', 'Anna Beispiel') in z4 and 'Nr. 199' == z4[('neu', 'Anna Beispiel')]['schluessel']
          and 'ähnlich: Anna Beispiel' in z4[('neu', 'Anna Beispiel')]['text']
          and not any(z['art'] == 'geaendert' and z['name'] == 'Anna Beispiel' for z in v4['zeilen']), (st, sorted(z4)))

    # Personalstelle ändert die Art: „vom Urlaub abgezogen“ entfällt
    eb = next(e for e in r['eintraege'] if e['quelle'] == 'zeiterfassung' and e['name'] == 'Bernd Muster')
    st, x = P.call('hr_eintrag_save', {'id': eb['id'], 'person': eb['person'], 'art': 'G', 'von': eb['von'], 'bis': eb['bis']})
    r = P.call('hr_verwaltung&jahr=2025')[1]
    eb2 = next(e for e in r['eintraege'] if e['id'] == eb['id'])
    check('Art geändert: zählt nicht mehr gegen den Urlaub', st == 200 and eb2['urlaubAbzug'] is False and {p_['name']: p_['konto'] for p_ in r['personen']}['Bernd Muster']['genommen'] == 0, (st, x, eb2))

    # ── Verwaltung → Mitarbeiter (Personalstelle ohne Administratorrechte): nur HR-Angaben ──
    pg.click('#vt-users'); pg.wait_for_selector('#view-users tr[data-hr-uid]')
    check('Personalstelle: Mitarbeiterliste mit HR-Angaben, ohne Rollen/Rechte', 'Rest' in pg.text_content('#view-users') and pg.locator('#us-zu-rollen').count() == 0
          and pg.locator('#user-add').count() == 0)
    pg.click(f'#view-users tr[data-hr-uid="{ids["erika"]}"]'); pg.wait_for_selector('#dlg-hr-person[open]')
    pg.fill('#dlg-hr-person [name="hr_abteilung"]', 'Vertrieb'); pg.fill('#dlg-hr-person [name="hr_anspruch"]', '31')
    pg.click('#dlg-hr-person [type="submit"]'); pg.wait_for_selector('#dlg-hr-person[open]', state='hidden')
    pe = {p_['nutzer']: p_ for p_ in P.call('mitarbeiter_hr')[1]['personen'] if p_['nutzer']}[ids['erika']]
    check('Oberfläche: Abteilung und Anspruch gespeichert', pe['abteilung'] == 'Vertrieb' and pe['jahre'][str(dt.date.today().year)]['anspruch'] == 31, pe)
    pg.screenshot(path=os.path.join(SH, 'hr_verwaltung_personal_1280.png'), full_page=True)
    pg.click('#vt-urlaub'); pg.wait_for_selector('#hr-einst')
    check('Verwaltung → Urlaub: Sperren und Ferien aus der HR', pg.locator('#hr-sperren .hr-zr').count() == 1 and pg.locator('#hr-ferien .hr-zr').count() == 1)
    pg.click('[data-plus="hr-ferien"]')
    zr = pg.locator('#hr-ferien .hr-zr').nth(1).locator('input')
    zr.nth(0).fill(f'{Y}-10-05'); zr.nth(1).fill(f'{Y}-10-16'); zr.nth(2).fill('Herbstferien')
    pg.click('#hr-einst [type="submit"]'); pg.wait_for_function('document.querySelectorAll("#hr-ferien .hr-zr").length === 2 && !document.querySelector("#hr-einst-err").textContent')
    check('Verwaltung → Urlaub: Ferien gespeichert', len(P.call('mitarbeiter_hr')[1]['ferien']) == 2)
    check('Verwaltung → Urlaub: Verfallstag vorbelegt (31.03.)', pg.input_value('#hr-uebertrag-bis').endswith('-03-31') and '31.03.' in pg.text_content('#hr-uebertrag-anzeige'))
    pg.fill('#hr-uebertrag-bis', f'{Y}-04-15')
    pg.click('#hr-einst [type="submit"]'); pg.wait_for_function('document.querySelector("#hr-uebertrag-anzeige") && document.querySelector("#hr-uebertrag-anzeige").textContent.includes("15.04.")')
    check('Oberfläche: Verfallstag 15.04. gespeichert, Ferien bleiben', P.call('mitarbeiter_hr')[1]['uebertragBis'] == '04-15' and len(P.call('mitarbeiter_hr')[1]['ferien']) == 2)
    pg.screenshot(path=os.path.join(SH, 'hr_verwaltung_urlaub_1280.png'), full_page=True)
    v_ = P.call('mitarbeiter_hr')[1]
    assert P.call('hr_einstellungen_save', {'sperren': v_['sperren'], 'ferien': v_['ferien'], 'uebertragBis': '03-31'})[0] == 200
    ctx.close()
    ctx, pg = anmelden('paul', 'Personal-Zugang-77', '198.51.100.5', 390)
    pg.goto(B + '/admin/#urlaub'); pg.wait_for_selector('#hr-uebertrag-bis')
    check('390px Verwaltung → Urlaub: Verfallstag ohne waagerechtes Scrollen', pg.evaluate('document.documentElement.scrollWidth <= window.innerWidth + 1'))
    pg.screenshot(path=os.path.join(SH, 'hr_verwaltung_urlaub_390.png'), full_page=True)
    ctx.close()

    # ── Verwaltung → Mitarbeiter: Eintrag ohne Mitarbeiter einzeln zuordnen (abweichender Name) ──
    alt = dokument('hr-personal')
    alt['personen'].append({'id': 'pute', 'vorname': 'Ute', 'nachname': 'Neu', 'abteilung': 'Lager', 'pos': 43, 'aktiv': True, 'personalnummer': '9903'})
    alt['rev'] = alt.get('rev', 0) + 1
    storage_write('hr-personal', alt)
    for breite in (390, 1280):
        ctx, pg = anmelden('admin', PW, '198.51.100.2', breite)
        pg.goto(B + '/admin/#users'); pg.wait_for_selector('#hr-alt [data-hr-alt-einzeln]')
        check(f'{breite}px: Hinweis nennt ehemalige Mitarbeiter, „Einzeln zuordnen“ vorne', 'ehemalige Mitarbeiter' in pg.text_content('#hr-alt')
              and pg.locator('#hr-alt .btn.primary[data-hr-alt-einzeln]').count() == 1)
        pg.locator('#hr-alt').screenshot(path=os.path.join(SH, f'hr_ehemalige_hinweis_{breite}.png'))
        pg.click('#hr-alt [data-hr-alt-einzeln]'); pg.wait_for_selector('#dlg-hr-alt[open] li[data-eintrag="pute"]')
        breit = pg.evaluate("document.querySelector('#dlg-hr-alt').scrollWidth <= document.querySelector('#dlg-hr-alt').clientWidth + 1")
        check(f'{breite}px: Dialog „Einzeln zuordnen“ mit Eintrag und Mitarbeiterauswahl, ohne waagerechtes Scrollen', breit and pg.locator('#dlg-hr-alt li[data-eintrag="pute"] select option').count() > 3)
        pg.screenshot(path=os.path.join(SH, f'hr_alt_zuordnen_{breite}.png'))
        if breite == 1280:
            check('Vorschlag: eindeutig gleicher Nachname vorausgewählt (Nina Neu)', pg.input_value('#dlg-hr-alt li[data-eintrag="pute"] select') == ids['nina'])
            pg.click('#dlg-hr-alt li[data-eintrag="pute"] [data-zuordnen]'); pg.wait_for_selector('#dlg-hr-alt li[data-eintrag="pute"] .form-err:not(:empty)')
            check('Fehler des Servers im Dialog (verschiedene Personalnummern)', 'Personalnummern' in pg.text_content('#dlg-hr-alt li[data-eintrag="pute"] .form-err'))
            pg.select_option('#dlg-hr-alt li[data-eintrag="pute"] select', '')
            pg.click('#dlg-hr-alt li[data-eintrag="pute"] [data-zuordnen]'); pg.wait_for_function("document.querySelector('#dlg-hr-alt li[data-eintrag=\"pute\"] .form-err').textContent.includes('wählen')")
            check('ohne Auswahl: Hinweis im Dialog', 'Mitarbeiter wählen' in pg.text_content('#dlg-hr-alt li[data-eintrag="pute"] .form-err'))
            pg.select_option('#dlg-hr-alt li[data-eintrag="pute"] select', ids['otto'])
            pg.click('#dlg-hr-alt li[data-eintrag="pute"] [data-zuordnen]')
            pg.wait_for_selector('#dlg-hr-alt li[data-eintrag="pute"]', state='detached')
            pn = [p_ for p_ in P.call('mitarbeiter_hr')[1]['personen'] if p_['nutzer'] == ids['otto']]
            check('Oberfläche: Eintrag dem gewählten Mitarbeiter zugeordnet (zusammengeführt, Personalnummer übernommen)', len(pn) == 1 and pn[0]['personalnummer'] == '9903'
                  and 'pute' not in [p_['id'] for p_ in P.call('mitarbeiter_hr')[1]['personen']], pn)
        ctx.close()

    # ── Verwaltung → Mitarbeiter (Administrator): anlegen ohne Anmeldung, HR-Teil im Dialog ──
    ctx, pg = anmelden('admin', PW, '198.51.100.2')
    pg.click('#vt-users'); pg.wait_for_selector('#view-users tr[data-uid]')
    check('Navigation heißt „Mitarbeiter“, Spalten Abteilung und Urlaub', pg.text_content('#vt-users').strip() == 'Mitarbeiter' and 'Abteilung' in pg.text_content('#view-users thead')
          and 'Urlaub' in pg.text_content('#view-users thead'))
    pg.click('#user-add'); pg.wait_for_selector('#dlg-user[open]')
    pg.check('#user-body [name="zugang"][value="0"]')
    check('ohne Anmeldung: keine Felder für Passwort, Rolle, Rechte', pg.is_hidden('#user-body [name="password"]') and pg.is_hidden('#user-body .role-pick:not(.us-zugang)') and pg.is_visible('#hr-teil'))
    pg.fill('#user-body [name="vorname"]', 'Lisa'); pg.fill('#user-body [name="nachname"]', 'Lager')
    pg.fill('#user-body [name="hr_abteilung"]', 'Lager'); pg.uncheck('#user-body [name="hr_tag"][value="5"]'); pg.fill('#user-body [name="hr_anspruch"]', '25')
    pg.screenshot(path=os.path.join(SH, 'hr_verwaltung_mitarbeiter_neu_1280.png'), full_page=True)
    pg.click('#user-body [type="submit"]'); pg.wait_for_selector('#dlg-user[open]', state='hidden')
    lisa = next(u for u in A.call('users')[1]['users'] if u['name'] == 'Lisa Lager')
    pl = {p_['nutzer']: p_ for p_ in P.call('mitarbeiter_hr')[1]['personen'] if p_['nutzer']}[lisa['id']]
    check('Oberfläche: Mitarbeiter ohne Anmeldung mit HR-Angaben angelegt', lisa['zugang'] is False and pl['abteilung'] == 'Lager' and pl['arbeitstage'] == [1, 2, 3, 4]
          and pl['jahre'][str(dt.date.today().year)]['anspruch'] == 25, (lisa, pl))
    pg.wait_for_selector(f'#view-users tr[data-uid="{lisa["id"]}"]')
    check('Liste: „ohne Anmeldung“ statt Benutzername und Rolle', 'ohne Anmeldung' in pg.text_content(f'#view-users tr[data-uid="{lisa["id"]}"]'))
    pg.screenshot(path=os.path.join(SH, 'hr_verwaltung_mitarbeiter_1280.png'), full_page=True)
    ctx.close()
    # Nur „Urlaub genehmigen“ (Gina): Reiter Anträge mit Knöpfen, ohne Personalstelle (Alle Einträge, Export, Verwaltungs-Links)
    for breite in (1280, 390):
        ctx, pg = anmelden('gina', 'Genehmigen-Zugang-77', '198.51.100.8', breite)
        pg.goto(B + '/module/hr/#antraege'); pg.wait_for_selector('.hr-antrag [data-genehmigen]')
        check(f'{breite}px Genehmigende: Anträge mit Knöpfen, ohne Export und „Alle Einträge“', pg.is_visible('[data-tab="antraege"]') and not pg.is_visible('[data-tab="export"]')
              and not pg.is_visible('#eintraege-tabelle') and not pg.is_visible('[data-verwaltung-link="users"]'))
        check(f'{breite}px Genehmigende: kein waagerechtes Scrollen', pg.evaluate('document.documentElement.scrollWidth <= window.innerWidth + 1'))
        pg.screenshot(path=os.path.join(SH, f'hr_genehmigen_{breite}.png'), full_page=True)
        ctx.close()
    # Personalstelle ohne Genehmigungsrecht (Vera): Anträge sichtbar, aber ohne Knöpfe, mit Hinweis
    ctx, pg = anmelden('vera', 'Genehmigen-Zugang-77', '198.51.100.9')
    pg.goto(B + '/module/hr/#antraege'); pg.wait_for_selector('.hr-antrag .hr-hinweis')
    check('Personalstelle ohne „Urlaub genehmigen“: Hinweis statt Knöpfe, „Alle Einträge“ sichtbar', pg.locator('.hr-antrag [data-genehmigen]').count() == 0
          and 'Urlaub genehmigen' in pg.text_content('.hr-antrag .hr-hinweis') and pg.is_visible('#eintraege-tabelle'))
    pg.screenshot(path=os.path.join(SH, 'hr_antraege_ohne_genehmigen_1280.png'), full_page=True)
    ctx.close()
    ctx, pg = anmelden('paul', 'Personal-Zugang-77', '198.51.100.5')
    pg.goto(B + '/module/hr/')
    pg.wait_for_selector('#hr-app:not([hidden])')
    # Antrag genehmigen in der Oberfläche
    pg.click('[data-tab="antraege"]'); pg.wait_for_selector('.hr-antrag')
    pg.screenshot(path=os.path.join(SH, 'hr_antraege_1280.png'), full_page=True)
    pg.click('.hr-antrag [data-genehmigen]'); pg.wait_for_selector('.hr-leer')
    check('Oberfläche: Antrag genehmigt, keine offenen mehr', P.call('hr_verwaltung&jahr=' + str(Y))[1]['offen'] == [])

    # Kalender der Personalstelle: Eintrag per Klick
    pg.click('[data-tab="kalender"]'); pg.wait_for_selector('#kal-plan tbody tr')
    pg.screenshot(path=os.path.join(SH, 'hr_kalender_1280.png'), full_page=True)
    pg.emulate_media(color_scheme='dark')
    pg.screenshot(path=os.path.join(SH, 'hr_kalender_1280_dunkel.png'), full_page=True)
    check('dunkles Design: dunkler Grund', pg.evaluate('getComputedStyle(document.body).backgroundColor') == 'rgb(18, 16, 17)', pg.evaluate('getComputedStyle(document.body).backgroundColor'))
    pg.emulate_media(color_scheme='light')
    pg.click('#kal-eintragen'); pg.wait_for_selector('#dlg-eintrag[open]')
    pg.select_option('#eintrag-person', PA); pg.select_option('#eintrag-art', 'G')
    gtag = montag(Y, 3, 2) + dt.timedelta(days=1)   # Dienstag im März (keine Feiertage)
    pg.fill('#eintrag-von', iso(gtag)); pg.dispatch_event('#eintrag-von', 'change')
    pg.click('#eintrag-speichern')
    pg.wait_for_selector('#dlg-eintrag[open]', state='hidden')
    gz = [e for e in P.call('hr_verwaltung&jahr=' + str(Y))[1]['eintraege'] if e['person'] == PA and e['art'] == 'G']
    check('Oberfläche: Gleittag eingetragen', len(gz) == 1 and gz[0]['von'] == iso(gtag) == gz[0]['bis'], gz)
    pg.click('#kal-eintragen'); pg.wait_for_selector('#dlg-eintrag[open]')
    pg.select_option('#eintrag-person', PA); pg.fill('#eintrag-von', iso(gtag + dt.timedelta(days=3))); pg.dispatch_event('#eintrag-von', 'change')
    pg.click('#eintrag-speichern'); pg.wait_for_selector('#eintrag-meldung.fehler')
    check('Oberfläche: Fehler im Dialog (freier Tag), Dialog bleibt offen', 'Arbeitstage' in pg.text_content('#eintrag-meldung') and pg.is_visible('#dlg-eintrag'))
    pg.click('#dlg-eintrag [data-schliessen]')

    # Export in der Oberfläche
    pg.click('[data-tab="export"]'); pg.wait_for_selector('#export-jahr option', state='attached')
    with pg.expect_download() as dl:
        pg.click('#export button[type="submit"]')
    check('Oberfläche: Excel-Download', dl.value.suggested_filename.startswith('Urlaub_') and dl.value.suggested_filename.endswith('.xlsx'), dl.value.suggested_filename)
    pg.screenshot(path=os.path.join(SH, 'hr_export_1280.png'))
    ctx.close()

    # Mitarbeiterin: Mein Urlaub, Antrag in der Oberfläche, Desktop und Handy
    for breite in (1280, 390):
        ctx, pg = anmelden('erika', 'Editor-Zugang-77', '198.51.100.3', breite)
        pg.goto(B + '/module/hr/'); pg.wait_for_selector('#mein-konto .hr-wert')
        check(f'{breite}px: Personalbereiche ausgeblendet', not pg.is_visible('[data-tab="antraege"]') and not pg.is_visible('[data-tab="export"]'))
        breit = pg.evaluate('document.documentElement.scrollWidth <= window.innerWidth + 1')
        check(f'{breite}px: kein waagerechtes Scrollen der Seite', breit)
        pg.screenshot(path=os.path.join(SH, f'hr_mein_{breite}.png'), full_page=True)
        if breite == 1280:
            pg.click('[data-jahr="1"]')
            pg.wait_for_function(f'document.querySelector("#mein-titel").textContent.includes("{Y}")')
            tag_ = montag(Y, 9)
            pg.evaluate("document.querySelectorAll('#antrag input[required]').forEach((i) => i.removeAttribute('required'))")   # eigene Meldung statt Browserhinweis
            pg.click('#antrag-senden'); pg.wait_for_selector('#antrag-meldung.fehler')
            pg.fill('#antrag-von', iso(tag_)); pg.dispatch_event('#antrag-von', 'change')
            check('Meldung zum leeren Antrag verschwindet nach Eingabe der Daten', pg.text_content('#antrag-meldung') == '', pg.text_content('#antrag-meldung'))
            pg.fill('#antrag-bis', iso(tag_ - dt.timedelta(days=3))); pg.dispatch_event('#antrag-bis', 'change')
            pg.click('#antrag-senden')
            pg.wait_for_function("document.activeElement && document.activeElement.id === 'antrag-bis'")
            check('Ende vor Beginn: Hinweis nur einmal, kein Antrag gesendet', pg.text_content('#antrag-tage') == 'Das Ende liegt vor dem Beginn.' and pg.text_content('#antrag-meldung') == '',
                  (pg.text_content('#antrag-tage'), pg.text_content('#antrag-meldung')))
            pg.fill('#antrag-bis', iso(tag_ + dt.timedelta(days=4))); pg.dispatch_event('#antrag-bis', 'change')
            check('Vorschau: 5 Tage Urlaub', '5 Tage' in pg.text_content('#antrag-tage'), pg.text_content('#antrag-tage'))
            pg.click('#antrag-senden'); pg.wait_for_selector('#antrag-meldung.ok')
            check('Oberfläche: Antrag gestellt, in der Liste', 'beantragt' in pg.text_content('#mein-liste'))
        pg.click('[data-tab="kalender"]'); pg.wait_for_selector('#kal-plan tbody tr')
        check(f'{breite}px: Kalender ohne Gründe (nur „abwesend“ in der Legende)', 'Krank' not in pg.text_content('#kal-legende') and 'abwesend' in pg.text_content('#kal-legende'))
        pg.screenshot(path=os.path.join(SH, f'hr_kalender_mitarbeiter_{breite}.png'), full_page=True)
        ctx.close()

    # produktiv: jeder Mitarbeiter sieht die normale Seite (otto war nie Tester)
    ctx, pg = anmelden('otto', 'Ohne-Zugang-77', '198.51.100.6')
    pg.goto(B + '/module/hr/'); pg.wait_for_selector('#hr-app:not([hidden]), #hr-err:not([hidden])')
    check('Mitarbeiter ohne Tester-Eintrag: normale Seite (Mein Urlaub), kein Hinweis „nicht freigegeben“',
          pg.is_visible('#hr-app') and pg.is_hidden('#hr-err') and pg.is_hidden('[data-tab="antraege"]'), pg.text_content('#hr-err'))
    ctx.close()
    ctx = br.new_context(); pg = ctx.new_page(); pg.on('pageerror', lambda e: fehler.append(str(e)))
    pg.goto(B + '/module/hr/'); pg.wait_for_selector('#hr-login:not([hidden])')
    check('ohne Anmeldung: Anmeldehinweis', pg.is_visible('#hr-login a'))
    ctx.close()
    br.close()
check('keine JavaScript-Fehler', not fehler, fehler)

print(f'{sum(res)}/{len(res)} passed')
sys.exit(0 if all(res) else 1)
