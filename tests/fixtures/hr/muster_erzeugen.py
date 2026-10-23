"""Erzeugt die synthetische Testdatei Urlaub_Muster_2025.xlsx im Aufbau der bisherigen Urlaubsplanung („Urlaub 2026.xlsx“):
ein Blatt je Monat, Legende oben, Monatsdatum in K2, Wochentage in Zeile 7, „Pos.“ und Tageszahlen in Zeile 8, ab Zeile 9
Pos. / Abteilung (senkrecht verbunden) / Name / Kürzel je Tag, darunter „Summe“. Nur erfundene Namen (GR-007).

    python3 tests/fixtures/hr/muster_erzeugen.py
"""
import calendar
import datetime as dt
import os

import openpyxl
from openpyxl.styles import Alignment, Font

WT = ['Mo', 'Di', 'Mi', 'Do', 'Fr', 'Sa', 'So']
PERSONEN = [('GF', 'Anna Beispiel'), ('Verwaltung', 'Bernd Muster'), ('Verwaltung', 'Carla Probe'), ('Lager', 'Dieter Test'), ('Lager', 'Emil Lager ')]
# (Name, Monat, Tage, Kürzel)
MARKEN = [('Anna Beispiel', 1, [13, 14, 15, 16, 17], 'U'), ('Bernd Muster', 1, [20, 21], 'K'), ('Carla Probe', 1, [31], 'G'),
          ('Carla Probe', 2, [3, 4], 'U'), ('Emil Lager ', 2, [10], 'S')]


def blatt(wb, monat, name):
    ws = wb.create_sheet(name)
    for z, t, k in ((2, 'Urlaub', 'U'), (3, 'Krank', 'K'), (4, 'Gleittag', 'G')):
        ws.cell(z, 3, t); ws.cell(z, 4, k)
    ws.cell(5, 3, 'Urlaubssperre')
    ws['K2'] = dt.datetime(2025, monat, 1); ws['K2'].number_format = 'mmm-yy'; ws.merge_cells('K2:U3')
    n = calendar.monthrange(2025, monat)[1]
    ws.cell(8, 1, 'Pos.'); ws.cell(8, 2, ' '); ws.cell(8, 3, ['Januar', 'Februar'][monat - 1])
    for t in range(1, n + 1):
        ws.cell(7, 3 + t, WT[dt.date(2025, monat, t).weekday()]); ws.cell(8, 3 + t, t)
    for i, k in enumerate('UKG'):
        ws.cell(7, 4 + n + i, k)
    z = 9; abt = None
    for pos, (a, name) in enumerate(PERSONEN, 1):
        if a != abt:
            ws.cell(z, 2, a).alignment = Alignment(textRotation=90); abt = a
        ws.cell(z, 1, pos); ws.cell(z, 3, name).font = Font(bold=True)
        if name == 'Dieter Test':   # 4-Tage-Woche: jeden Freitag „x“
            for t in range(1, n + 1):
                if dt.date(2025, monat, t).weekday() == 4: ws.cell(z, 3 + t, 'x')
        for nm, m, tage, k in MARKEN:
            if nm == name and m == monat:
                for t in tage: ws.cell(z, 3 + t, k)
        last = openpyxl.utils.get_column_letter(3 + n)
        for i, k in enumerate('UKG'):
            ws.cell(z, 4 + n + i, f'=COUNTIF(D{z}:{last}{z},"{k}")')
        z += 1
    ws.merge_cells('B10:B11'); ws.merge_cells('B12:B13')
    ws.cell(z + 1, 3, 'Summe')


wb = openpyxl.Workbook(); wb.remove(wb.active)
blatt(wb, 1, 'Januar ´25'); blatt(wb, 2, "Feb. '25")
wb.save(os.path.join(os.path.dirname(os.path.abspath(__file__)), 'Urlaub_Muster_2025.xlsx'))
